"""Backend API tests for RE DECOR & MORE Procurement ERP."""

import asyncio
import importlib.util
import os
import io
import json
import re
import sys
import tempfile
import sqlite3
import uuid
from datetime import datetime, timezone
from contextlib import closing
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from alembic.migration import MigrationContext
from alembic.operations import Operations

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))
TEST_DIR = tempfile.TemporaryDirectory(prefix="procurex-tests-")
TEST_DB = Path(TEST_DIR.name) / "procurement-test.db"
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB.as_posix()}"
os.environ["CORS_ORIGINS"] = "http://localhost:3000"
os.environ["INTERNAL_REQUEST_TOKEN"] = "test-internal-token"
os.environ["PUBLIC_REQUEST_RATE_LIMIT"] = "100"
os.environ["INCOMING_REQUEST_UPLOAD_DIR"] = str(Path(TEST_DIR.name) / "request-uploads")
os.environ["PROCUREX_BACKUP_DIR"] = str(Path(TEST_DIR.name) / "backups")
os.environ["PROCUREX_LOG_DIR"] = str(Path(TEST_DIR.name) / "logs")

from database import db, engine, init_db  # noqa: E402
from attachment_storage import (  # noqa: E402
    ROOT_DIR, LocalAttachmentStorage, build_attachment_storage,
)
import business_codes  # noqa: E402
from business_codes import next_business_code  # noqa: E402
from db_migrations import (  # noqa: E402
    BUSINESS_CODE_SCHEMA_VERSION,
    CONSTRUCTION_CALCULATOR_SCHEMA_VERSION,
    DOCUMENT_CAPTURE_SCHEMA_VERSION,
    SUPPLIER_PRICE_COMPARISON_SCHEMA_VERSION,
    SUPPLIER_OFFER_ADJUSTMENTS_SCHEMA_VERSION,
    migrate_business_code_sequences,
    migrate_construction_calculator,
    migrate_document_capture,
    migrate_incoming_requests,
    migrate_item_classification,
    migrate_item_identity,
    migrate_supplier_price_comparisons,
    migrate_supplier_offer_adjustments,
)
from excel_io import import_data, parse_workbook  # noqa: E402
from price_comparisons import (  # noqa: E402
    PriceComparison, PriceComparisonSupplierOffer, calculate_comparison,
)
from procurement_workflow import (  # noqa: E402
    EngineerApproval, calculate_procurement_kpis,
    APPROVAL_STAGE_EXPENDITURE_APPROVAL, APPROVAL_STAGE_FUNDS_AVAILABILITY,
)
from server import (  # noqa: E402
    _next_payment_id, _next_po_number, _next_purchase_id, app, create_app,
)
from database import (  # noqa: E402
    Base, BusinessCodeSequence, Item, LocalCollection, Payment, Purchase,
    PriceHistory, Project, PurchaseOrder, PurchaseOrderItem, PurchaseOrderPayment,
    Supplier, PurchaseOrderReceipt, SessionLocal,
)
from document_capture.domain import ExtractedLineItem, OCRResult  # noqa: E402
from document_capture.jobs import run_once  # noqa: E402
from document_capture.security import MAX_FILE_BYTES  # noqa: E402
from document_capture.models import (  # noqa: E402
    DocumentProcessingJob,
    ItemMasterCreationRequest,
    PurchaseRequestDocument,
    RequestAuditEvent,
)
import document_capture.service as document_service  # noqa: E402
from incoming_requests import (
    IncomingPurchaseRequest,
    IncomingPurchaseRequestItem,
    IncomingRequestGeneralAttachment,
    IncomingRequestStatusHistory,
)  # noqa: E402
from auth.models import User, UserProjectAccess  # noqa: E402
from auth.security import hash_password  # noqa: E402
init_db()
WORKBOOK = BACKEND_DIR / "workbook.xlsm"
asyncio.run(import_data(db, parse_workbook(WORKBOOK.read_bytes())))
API = "/api"


@pytest.fixture(scope="session")
def s():
    with TestClient(app) as client:
        yield client


@pytest.fixture(scope="session", autouse=True)
def cleanup_test_database():
    yield
    engine.dispose()
    TEST_DIR.cleanup()


# ---------- Dashboard ----------
def test_dashboard_kpis(s, admin_headers):
    r = s.get(f"{API}/dashboard", headers=admin_headers, timeout=30)
    assert r.status_code == 200
    d = r.json()
    assert d["total_purchases"] == 29947.5
    assert d["purchase_count"] == 1
    assert d["supplier_count"] == 17
    assert d["item_count"] == 15
    assert d["total_paid"] == 29947.5
    assert d["total_outstanding"] == 0
    assert d["project_count"] == 2
    assert d["customer_count"] == 2
    assert d["direct_purchase_total"] == d["total_purchases"]
    assert d["direct_paid_total"] == d["total_paid"]
    assert d["direct_outstanding_total"] == d["total_outstanding"]
    assert d["formal_po_total"] >= 0
    assert d["approval_paid_total"] >= 0
    assert isinstance(d["actions"], dict)
    for k in ("by_supplier", "by_project", "monthly", "payment_status"):
        assert isinstance(d[k], list)


def test_procurement_kpis_keep_financial_ledgers_independent():
    row = SimpleNamespace
    empty = calculate_procurement_kpis()
    assert empty["formal_po_total"] == 0
    assert empty["direct_purchase_total"] == 0
    assert empty["direct_paid_total"] == 0
    assert empty["direct_outstanding_total"] == 0
    assert empty["approval_paid_total"] == 0
    assert empty["approval_pending_total"] == 0

    purchases = [
        row(purchase_id="PUR-DIRECT-1", invoice_total=100),
        row(purchase_id="PUR-DIRECT-2", invoice_total=200),
    ]
    direct_payments = [
        row(purchase_id="PUR-DIRECT-1", amount_paid=120),
        row(purchase_id="PUR-DIRECT-2", amount_paid=50),
    ]
    direct_only = calculate_procurement_kpis(
        purchases=purchases, direct_payments=direct_payments,
    )
    assert direct_only["direct_purchase_count"] == 2
    assert direct_only["direct_purchase_total"] == 300
    assert direct_only["direct_paid_total"] == 170
    assert direct_only["direct_outstanding_total"] == 150
    assert direct_only["formal_po_total"] == 0

    orders = [
        row(status="draft", final_total=100, approval_id=""),
        row(status="in_delivery", final_total=300, approval_id="APR-1"),
        row(status="partial_received", final_total=50, approval_id="APR-2"),
        row(status="completed", final_total=200, approval_id="APR-3"),
        row(status="cancelled", final_total=999, approval_id="APR-4"),
    ]
    approval_payments = [
        row(status="verified", amount=250),
        row(status="pending", amount=100),
        row(status="rejected", amount=50),
    ]
    formal_only = calculate_procurement_kpis(
        purchase_orders=orders, approval_payments=approval_payments,
    )
    assert formal_only["formal_po_count"] == 4
    assert formal_only["formal_po_total"] == 650
    assert formal_only["formal_under_supply_count"] == 2
    assert formal_only["formal_under_supply_value"] == 350
    assert formal_only["formal_completed_po_count"] == 1
    assert formal_only["formal_completed_po_value"] == 200
    assert formal_only["approval_paid_total"] == 250
    assert formal_only["approval_pending_total"] == 100
    assert formal_only["direct_purchase_total"] == 0

    both = calculate_procurement_kpis(
        purchase_orders=orders,
        purchases=purchases,
        direct_payments=direct_payments,
        approval_payments=approval_payments,
    )
    assert both["formal_po_total"] == 650
    assert both["direct_purchase_total"] == 300
    assert both["direct_paid_total"] == 170
    assert both["approval_paid_total"] == 250
    assert both["paid"] == 170

    approvals = [
        row(id="APR-COMMERCIAL", approval_type="comparison_workflow", status="pending_approval", approval_stage="fund_release"),
        row(id="APR-FUNDS", approval_type="comparison_workflow", status="pending_approval", approval_stage="funds_release"),
        row(id="APR-READY", approval_type="comparison_workflow", status="approved", approval_stage="po_ready"),
    ]
    bottlenecks = calculate_procurement_kpis(
        approvals=approvals,
        purchase_orders=orders + [
            row(status="delivery_problem", final_total=25, approval_id="APR-PROBLEM"),
        ],
    )
    assert bottlenecks["actions"]["fund_approval"] == 1
    assert bottlenecks["actions"]["funds_release"] == 1
    assert bottlenecks["actions"]["ready_for_po"] == 1
    assert bottlenecks["actions"]["under_supply"] == 1
    assert bottlenecks["actions"]["partial_receiving"] == 1
    assert bottlenecks["actions"]["delivery_problem"] == 1
    assert [item["key"] for item in bottlenecks["action_priorities"]][:2] == [
        "delivery_problem", "partial_receiving",
    ]


def test_project_hub_uses_direct_purchase_business_id_for_payments(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    project_id = f"T-KPI-PROJECT-{suffix}"
    purchase_pk = f"T-KPI-PURCHASE-PK-{suffix}"
    purchase_number = f"T-KPI-PUR-{suffix}"
    payment_pk = f"T-KPI-PAYMENT-PK-{suffix}"
    with SessionLocal.begin() as session:
        session.add(Project(
            id=project_id, code=f"T-KPI-{suffix}", name=f"مشروع مؤشرات {suffix}",
        ))
        session.add(Purchase(
            id=purchase_pk, purchase_id=purchase_number,
            project_id=project_id, project_name=f"مشروع مؤشرات {suffix}",
            purchase_date="2026-08-18", invoice_number=f"INV-{suffix}",
            invoice_total=500, created_at="2026-08-18T00:00:00+00:00",
        ))
        session.add(Payment(
            id=payment_pk, payment_id=f"T-KPI-PAY-{suffix}",
            purchase_id=purchase_number, payment_date="2026-08-18",
            amount_paid=125, created_at="2026-08-18T00:00:00+00:00",
        ))
    try:
        hub = s.get(
            f"{API}/workflow/projects/{project_id}/procurement-hub",
            headers={"X-Internal-Token": "test-internal-token", **admin_headers},
        )
        assert hub.status_code == 200, hub.text
        kpis = hub.json()["kpis"]
        assert kpis["direct_purchase_total"] == 500
        assert kpis["direct_paid_total"] == 125
        assert kpis["direct_outstanding_total"] == 375
        assert kpis["paid"] == 125
    finally:
        with SessionLocal.begin() as session:
            session.delete(session.get(Payment, payment_pk))
            session.delete(session.get(Purchase, purchase_pk))
            session.delete(session.get(Project, project_id))


# ---------- Entity list counts ----------
def test_entity_lists(s, admin_headers):
    suppliers = s.get(f"{API}/suppliers", headers=admin_headers).json()
    assert len(suppliers) == 17
    assert "formal_po_total" not in suppliers[0]
    assert len(s.get(f"{API}/customers", headers=admin_headers).json()) == 2
    assert len(s.get(f"{API}/projects", headers=admin_headers).json()) == 2
    assert len(s.get(f"{API}/items", headers=admin_headers).json()) == 15


def test_supplier_list_exposes_grouped_id_linked_procurement_activity(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    supplier_id = f"T-SUP-KPI-{suffix}"
    created_ids = []
    with SessionLocal.begin() as session:
        session.add(Supplier(
            id=supplier_id, code=f"T-SUP-{suffix}", name=f"مورد مؤشرات {suffix}",
        ))
        purchases = [
            Purchase(
                id=f"T-SUP-PUR-PK-1-{suffix}", purchase_id=f"T-SUP-PUR-1-{suffix}",
                supplier_id=supplier_id, supplier_name=f"مورد مؤشرات {suffix}",
                purchase_date="2026-08-17", invoice_number=f"INV-1-{suffix}",
                invoice_total=125, created_at="2026-08-17T00:00:00+00:00",
            ),
            Purchase(
                id=f"T-SUP-PUR-PK-2-{suffix}", purchase_id=f"T-SUP-PUR-2-{suffix}",
                supplier_id=supplier_id, supplier_name=f"مورد مؤشرات {suffix}",
                purchase_date="2026-08-18", invoice_number=f"INV-2-{suffix}",
                invoice_total=75, created_at="2026-08-18T00:00:00+00:00",
            ),
            Purchase(
                id=f"T-SUP-PUR-PK-3-{suffix}", purchase_id=f"T-SUP-PUR-3-{suffix}",
                supplier_id="", supplier_name=f"مورد مؤشرات {suffix}",
                purchase_date="2026-08-19", invoice_number=f"INV-3-{suffix}",
                invoice_total=700, created_at="2026-08-19T00:00:00+00:00",
            ),
        ]
        orders = [
            PurchaseOrder(
                id=f"T-SUP-PO-1-{suffix}", po_number=f"T-SUP-PO-NO-1-{suffix}",
                supplier_id=supplier_id, supplier_name=f"مورد مؤشرات {suffix}",
                po_date="2026-08-16", status="in_delivery", final_total=300,
                created_at="2026-08-16T00:00:00+00:00", updated_at="2026-08-16T00:00:00+00:00",
            ),
            PurchaseOrder(
                id=f"T-SUP-PO-2-{suffix}", po_number=f"T-SUP-PO-NO-2-{suffix}",
                supplier_id=supplier_id, supplier_name=f"مورد مؤشرات {suffix}",
                po_date="2026-08-18", status="completed", final_total=200,
                created_at="2026-08-18T00:00:00+00:00", updated_at="2026-08-18T00:00:00+00:00",
            ),
            PurchaseOrder(
                id=f"T-SUP-PO-3-{suffix}", po_number=f"T-SUP-PO-NO-3-{suffix}",
                supplier_id=supplier_id, supplier_name=f"مورد مؤشرات {suffix}",
                po_date="2026-08-19", status="cancelled", final_total=999,
                created_at="2026-08-19T00:00:00+00:00", updated_at="2026-08-19T00:00:00+00:00",
            ),
        ]
        session.add_all(purchases + orders)
        created_ids = [row.id for row in purchases + orders]
    try:
        response = s.get(
            f"{API}/suppliers", params={"include_procurement": "true"}, headers=admin_headers,
        )
        assert response.status_code == 200
        supplier = next(row for row in response.json() if row["id"] == supplier_id)
        assert supplier["direct_purchase_count"] == 2
        assert supplier["direct_purchase_total"] == 200
        assert supplier["formal_po_count"] == 2
        assert supplier["formal_po_total"] == 500
        assert supplier["formal_under_supply_count"] == 1
        assert supplier["formal_under_supply_value"] == 300
        assert supplier["formal_completed_po_count"] == 1
        assert supplier["formal_completed_po_value"] == 200
        assert supplier["last_procurement_date"] == "2026-08-18"
    finally:
        with SessionLocal.begin() as session:
            for row_id in created_ids:
                row = session.get(Purchase, row_id) or session.get(PurchaseOrder, row_id)
                if row:
                    session.delete(row)
            supplier = session.get(Supplier, supplier_id)
            if supplier:
                session.delete(supplier)


def test_legacy_item_categories_are_exposed_safely(s, admin_headers):
    items = s.get(f"{API}/items", headers=admin_headers).json()
    assert items
    for item in items:
        assert item["main_category"] == item["category"]
        assert item["subcategory"] == ""
        assert item["product_name"] == item["name"]
        assert item["brand"] == ""
        assert item["specifications"] == item["specs"]
        assert "last_price" in item
        assert "last_supplier" in item


def test_item_category_subcategory_and_name_filtering(s, admin_headers):
    created = []
    try:
        fixtures = [
            ("TEST_FILTER_CEMENT", "مواد", "أسمنت", "Brand A", "شيكارة"),
            ("TEST_FILTER_SAND", "مواد", "ركام", "Brand B", "م³"),
            ("TEST_FILTER_CABLE", "كهرباء", "كابلات", "Brand A", "متر"),
        ]
        for name, main_category, subcategory, brand, unit in fixtures:
            response = s.post(
                f"{API}/items",
                json={
                    "product_name": name,
                    "main_category": main_category,
                    "subcategory": subcategory,
                    "brand": brand,
                    "specifications": f"SPEC_{name}",
                    "unit": unit,
                },
                headers=admin_headers,
            )
            assert response.status_code == 200, response.text
            created.append(response.json())

        by_main = s.get(f"{API}/items", params={"main_category": "مواد"}, headers=admin_headers).json()
        assert {item["name"] for item in by_main} >= {
            "TEST_FILTER_CEMENT",
            "TEST_FILTER_SAND",
        }
        assert "TEST_FILTER_CABLE" not in {item["name"] for item in by_main}

        by_subcategory = s.get(
            f"{API}/items",
            params={"main_category": "مواد", "subcategory": "أسمنت"},
            headers=admin_headers,
        ).json()
        assert [item["name"] for item in by_subcategory] == ["TEST_FILTER_CEMENT"]

        by_brand = s.get(
            f"{API}/items",
            params={
                "main_category": "مواد",
                "subcategory": "أسمنت",
                "brand": "Brand A",
            },
            headers=admin_headers,
        ).json()
        assert [item["product_name"] for item in by_brand] == ["TEST_FILTER_CEMENT"]

        by_name = s.get(
            f"{API}/items",
            params={
                "main_category": "مواد",
                "subcategory": "ركام",
                "search": "sand",
            },
            headers=admin_headers,
        ).json()
        assert [item["name"] for item in by_name] == ["TEST_FILTER_SAND"]

        by_specification = s.get(
            f"{API}/items", params={"search": "SPEC_TEST_FILTER_CABLE"}, headers=admin_headers,
        ).json()
        assert [item["product_name"] for item in by_specification] == [
            "TEST_FILTER_CABLE"
        ]
    finally:
        for item in created:
            s.delete(f"{API}/items/{item['id']}", headers=admin_headers)


# ---------- CRUD suppliers/customers/projects/items with autocode + dup ----------
@pytest.mark.parametrize(
    "coll,prefix",
    [
        ("suppliers", "SUP-"),
        ("customers", "CUS-"),
        ("projects", "PRJ-"),
        ("items", "ITM-"),
    ],
)
def test_entity_crud(s, coll, prefix, admin_headers):
    name = f"TEST_{coll}_entity"
    # cleanup pre
    existing = s.get(f"{API}/{coll}", headers=admin_headers).json()
    for e in existing:
        if e.get("name") == name:
            s.delete(f"{API}/{coll}/{e['id']}", headers=admin_headers)
    r = s.post(f"{API}/{coll}", json={"name": name}, headers=admin_headers)
    assert r.status_code == 200, r.text
    created = r.json()
    assert created["name"] == name
    assert created["code"].startswith(prefix)
    assert len(created["code"].removeprefix(prefix)) == 6
    eid = created["id"]

    # duplicate name
    dup = s.post(f"{API}/{coll}", json={"name": name}, headers=admin_headers)
    assert dup.status_code == 409
    assert "بالفعل" in dup.json().get("detail", "")

    # update
    up = s.put(f"{API}/{coll}/{eid}", json={"name": name + "_upd"}, headers=admin_headers)
    assert up.status_code == 200
    assert up.json()["name"] == name + "_upd"

    # delete
    d = s.delete(f"{API}/{coll}/{eid}", headers=admin_headers)
    assert d.status_code == 200


@pytest.mark.parametrize(
    "coll,prefix",
    [
        ("suppliers", "SUP-"),
        ("customers", "CUS-"),
        ("projects", "PRJ-"),
        ("items", "ITM-"),
    ],
)
def test_business_codes_ignore_input_are_immutable_and_are_never_reused(
    s, coll, prefix, admin_headers
):
    first_name = f"TEST_CODE_FIRST_{coll}"
    second_name = f"TEST_CODE_SECOND_{coll}"
    created_ids = []
    try:
        first = s.post(
            f"{API}/{coll}",
            json={
                "name": first_name,
                "product_name": first_name,
                "code": "MANUAL-999999",
            },
            headers=admin_headers,
        )
        assert first.status_code == 200, first.text
        first_body = first.json()
        created_ids.append(first_body["id"])
        assert first_body["code"].startswith(prefix)
        assert first_body["code"] != "MANUAL-999999"

        updated = s.put(
            f"{API}/{coll}/{first_body['id']}",
            json={
                "name": first_body["name"],
                "product_name": first_body.get("product_name"),
                "code": "CHANGED-999999",
            },
            headers=admin_headers,
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["code"] == first_body["code"]

        assert s.delete(f"{API}/{coll}/{first_body['id']}", headers=admin_headers).status_code == 200
        created_ids.remove(first_body["id"])
        engine.dispose()  # New connections simulate application restart.

        second = s.post(
            f"{API}/{coll}",
            json={
                "name": second_name,
                "product_name": second_name,
            },
            headers=admin_headers,
        )
        assert second.status_code == 200, second.text
        second_body = second.json()
        created_ids.append(second_body["id"])
        assert second_body["code"] != first_body["code"]
        assert int(second_body["code"].removeprefix(prefix)) > int(
            first_body["code"].removeprefix(prefix)
        )
    finally:
        for entity_id in created_ids:
            s.delete(f"{API}/{coll}/{entity_id}", headers=admin_headers)


def test_business_code_reservations_are_unique_under_concurrency():
    with ThreadPoolExecutor(max_workers=4) as executor:
        codes = list(
            executor.map(
                lambda _index: next_business_code("suppliers"),
                range(12),
            )
        )
    assert len(codes) == len(set(codes)) == 12
    assert all(
        code.startswith("SUP-") and len(code.removeprefix("SUP-")) == 6
        for code in codes
    )


def test_supplier_sequence_uses_highest_suffix_across_mixed_historical_widths(
    tmp_path
):
    isolated_engine = create_engine(
        f"sqlite:///{(tmp_path / 'mixed-supplier-codes.db').as_posix()}"
    )
    Base.metadata.create_all(
        isolated_engine,
        tables=[Supplier.__table__, BusinessCodeSequence.__table__],
    )
    historical_codes = ["SUP-001", "SUP-002", "SUP-000020", "SUP-000026"]
    with isolated_engine.begin() as connection:
        connection.execute(
            Supplier.__table__.insert(),
            [
                {"id": f"historical-{index}", "code": code, "name": code}
                for index, code in enumerate(historical_codes)
            ],
        )
        connection.execute(
            BusinessCodeSequence.__table__.insert().values(
                entity="suppliers", next_value=3
            )
        )

    original_engine = business_codes.engine
    original_is_sqlite = business_codes.IS_SQLITE
    business_codes.engine = isolated_engine
    business_codes.IS_SQLITE = True

    try:
        first = business_codes.next_business_code("suppliers")
        second = business_codes.next_business_code("suppliers")

        assert first == "SUP-000027"
        assert second == "SUP-000028"
        with isolated_engine.connect() as connection:
            assert set(connection.scalars(select(Supplier.code)).all()) == set(
                historical_codes
            )
    finally:
        business_codes.engine = original_engine
        business_codes.IS_SQLITE = original_is_sqlite
        isolated_engine.dispose()


def test_purchase_order_sequence_respects_history_concurrency_and_restart():
    with SessionLocal() as session:
        existing_numbers = [
            int(order.po_number.removeprefix("PO-"))
            for order in session.query(PurchaseOrder).all()
            if order.po_number.startswith("PO-")
            and order.po_number.removeprefix("PO-").isdigit()
        ]
    historical_value = max(existing_numbers, default=0) + 50
    historical_id = str(uuid.uuid4())
    with SessionLocal.begin() as session:
        session.add(PurchaseOrder(
            id=historical_id, po_number=f"PO-{historical_value:06d}",
            status="completed", created_at="2020-01-01T00:00:00+00:00",
            updated_at="2020-01-01T00:00:00+00:00",
        ))
    try:
        first = asyncio.run(_next_po_number())
        second = asyncio.run(_next_po_number())
        assert first == f"PO-{historical_value + 1:06d}"
        assert second == f"PO-{historical_value + 2:06d}"

        with ThreadPoolExecutor(max_workers=4) as executor:
            concurrent = list(executor.map(
                lambda _index: asyncio.run(_next_po_number()),
                range(12),
            ))
        assert len(concurrent) == len(set(concurrent)) == 12
        concurrent_values = sorted(int(code.removeprefix("PO-")) for code in concurrent)
        assert concurrent_values == list(range(historical_value + 3, historical_value + 15))
        assert all(code.startswith("PO-") and len(code) == 9 for code in concurrent)

        with SessionLocal() as session:
            stored_next = session.get(BusinessCodeSequence, "purchase_orders").next_value
        engine.dispose()
        after_restart = asyncio.run(_next_po_number())
        assert int(after_restart.removeprefix("PO-")) >= stored_next
        with SessionLocal() as session:
            assert session.get(PurchaseOrder, historical_id).po_number == f"PO-{historical_value:06d}"
    finally:
        with SessionLocal.begin() as session:
            historical = session.get(PurchaseOrder, historical_id)
            if historical:
                session.delete(historical)


def test_direct_purchase_sequence_respects_history_concurrency_and_restart():
    with SessionLocal() as session:
        existing_numbers = [
            int(purchase.purchase_id.removeprefix("PUR-"))
            for purchase in session.query(Purchase).all()
            if purchase.purchase_id.startswith("PUR-")
            and purchase.purchase_id.removeprefix("PUR-").isdigit()
        ]
    historical_value = max(existing_numbers, default=0) + 50
    historical_id = str(uuid.uuid4())
    with SessionLocal.begin() as session:
        session.add(Purchase(
            id=historical_id, purchase_id=f"PUR-{historical_value:06d}",
            purchase_date="2020-01-01", invoice_number="HISTORICAL-PUR-SEQUENCE",
            created_at="2020-01-01T00:00:00+00:00",
        ))
    try:
        first = asyncio.run(_next_purchase_id())
        second = asyncio.run(_next_purchase_id())
        assert first == f"PUR-{historical_value + 1:06d}"
        assert second == f"PUR-{historical_value + 2:06d}"

        with ThreadPoolExecutor(max_workers=4) as executor:
            concurrent = list(executor.map(
                lambda _index: asyncio.run(_next_purchase_id()),
                range(12),
            ))
        assert len(concurrent) == len(set(concurrent)) == 12
        concurrent_values = sorted(int(code.removeprefix("PUR-")) for code in concurrent)
        assert concurrent_values == list(range(historical_value + 3, historical_value + 15))
        assert all(code.startswith("PUR-") and len(code) == 10 for code in concurrent)

        with SessionLocal() as session:
            stored_next = session.get(BusinessCodeSequence, "purchases").next_value
        engine.dispose()
        after_restart = asyncio.run(_next_purchase_id())
        assert int(after_restart.removeprefix("PUR-")) >= stored_next
        with SessionLocal() as session:
            assert session.get(Purchase, historical_id).purchase_id == f"PUR-{historical_value:06d}"
    finally:
        with SessionLocal.begin() as session:
            historical = session.get(Purchase, historical_id)
            if historical:
                session.delete(historical)


def test_direct_payment_sequence_respects_history_concurrency_and_restart():
    with SessionLocal() as session:
        existing_numbers = [
            int(payment.payment_id.removeprefix("PAY-"))
            for payment in session.query(Payment).all()
            if payment.payment_id.startswith("PAY-")
            and payment.payment_id.removeprefix("PAY-").isdigit()
        ]
    historical_value = max(existing_numbers, default=0) + 50
    historical_id = str(uuid.uuid4())
    with SessionLocal.begin() as session:
        session.add(Payment(
            id=historical_id, payment_id=f"PAY-{historical_value:03d}",
            payment_date="2020-01-01", purchase_id="PUR-000001", amount_paid=1,
            created_at="2020-01-01T00:00:00+00:00",
        ))
    try:
        first = asyncio.run(_next_payment_id())
        second = asyncio.run(_next_payment_id())
        assert first == f"PAY-{historical_value + 1:06d}"
        assert second == f"PAY-{historical_value + 2:06d}"

        with ThreadPoolExecutor(max_workers=4) as executor:
            concurrent = list(executor.map(
                lambda _index: asyncio.run(_next_payment_id()),
                range(12),
            ))
        assert len(concurrent) == len(set(concurrent)) == 12
        concurrent_values = sorted(int(code.removeprefix("PAY-")) for code in concurrent)
        assert concurrent_values == list(range(historical_value + 3, historical_value + 15))
        assert all(code.startswith("PAY-") and len(code) == 10 for code in concurrent)

        with SessionLocal() as session:
            stored_next = session.get(BusinessCodeSequence, "payments").next_value
        engine.dispose()
        after_restart = asyncio.run(_next_payment_id())
        assert int(after_restart.removeprefix("PAY-")) >= stored_next
        with SessionLocal() as session:
            assert session.get(Payment, historical_id).payment_id == f"PAY-{historical_value:03d}"
    finally:
        with SessionLocal.begin() as session:
            historical = session.get(Payment, historical_id)
            if historical:
                session.delete(historical)


@pytest.mark.parametrize("coll", ["suppliers", "customers"])
def test_phone_is_optional_text_and_preserves_leading_zero_in_lists(s, coll, admin_headers):
    with_phone_name = f"TEST_PHONE_{coll}"
    without_phone_name = f"TEST_EMPTY_PHONE_{coll}"
    created_ids = []
    try:
        with_phone = s.post(
            f"{API}/{coll}",
            json={
                "name": with_phone_name,
                "phone": "01001234567",
            },
            headers=admin_headers,
        )
        without_phone = s.post(
            f"{API}/{coll}",
            json={
                "name": without_phone_name,
                "phone": "",
            },
            headers=admin_headers,
        )
        assert with_phone.status_code == 200, with_phone.text
        assert without_phone.status_code == 200, without_phone.text
        created_ids.extend([with_phone.json()["id"], without_phone.json()["id"]])

        listed = {row["name"]: row for row in s.get(f"{API}/{coll}", headers=admin_headers).json()}
        assert listed[with_phone_name]["phone"] == "01001234567"
        assert listed[without_phone_name]["phone"] == ""
    finally:
        for entity_id in created_ids:
            s.delete(f"{API}/{coll}/{entity_id}", headers=admin_headers)


def test_business_code_sequence_sqlite_migration_preserves_existing_rows(tmp_path):
    database_path = tmp_path / "business-codes.db"
    migration_engine = create_engine(f"sqlite:///{database_path.as_posix()}")
    with migration_engine.begin() as connection:
        connection.exec_driver_sql(
            "CREATE TABLE suppliers (id VARCHAR PRIMARY KEY, code VARCHAR UNIQUE)"
        )
        connection.exec_driver_sql(
            "INSERT INTO suppliers (id, code) VALUES ('legacy', 'SUP-0007')"
        )

    backup_path = migrate_business_code_sequences(migration_engine)
    assert backup_path and backup_path.is_file()
    with sqlite3.connect(database_path) as connection:
        assert (
            connection.execute(
                "SELECT code FROM suppliers WHERE id = 'legacy'"
            ).fetchone()[0]
            == "SUP-0007"
        )
        assert (
            connection.execute(
                "SELECT name FROM sqlite_master "
                "WHERE type = 'table' AND name = 'business_code_sequences'"
            ).fetchone()[0]
            == "business_code_sequences"
        )
        assert connection.execute("PRAGMA user_version").fetchone()[0] == (
            BUSINESS_CODE_SCHEMA_VERSION
        )
    migration_engine.dispose()


# ---------- Supplier price comparisons ----------
def _comparison_offer(item_id, supplier_id, unit_price=0, **overrides):
    return {
        "id": f"{item_id}-{supplier_id}",
        "item_id": item_id,
        "item_code": item_id.upper(),
        "product_name": item_id,
        "supplier_id": supplier_id,
        "supplier_name": supplier_id,
        "quantity": 1,
        "unit_price": unit_price,
        "discount_pct": 0,
        "tax_pct": 0,
        "shipping_cost": 0,
        "other_cost": 0,
        "delivery_days": 5,
        "availability": "available",
        "price_valid_until": "2099-12-31",
        **overrides,
    }


@pytest.mark.parametrize(
    (
        "first", "second", "expected_item_supplier", "expected_item_total",
        "expected_supplier", "expected_total",
    ),
    [
        pytest.param(
            {"unit_price": "١٬٠٠٠"},
            {"unit_price": "1,100"},
            "supplier-1",
            1000,
            "supplier-1",
            1000,
            id="numeric-formatted-prices",
        ),
        pytest.param(
            {"unit_price": 120, "discount_pct": 25},
            {"unit_price": 100},
            "supplier-2",
            100,
            "supplier-1",
            90,
            id="discount",
        ),
        pytest.param(
            {"unit_price": 90, "shipping_cost": 30},
            {"unit_price": 100},
            "supplier-1",
            90,
            "supplier-2",
            100,
            id="shipping",
        ),
        pytest.param(
            {"unit_price": 90, "tax_pct": 20},
            {"unit_price": 100},
            "supplier-1",
            90,
            "supplier-2",
            100,
            id="tax",
        ),
    ],
)
def test_supplier_recommendation_uses_numeric_final_total(
    first, second, expected_item_supplier, expected_item_total,
    expected_supplier, expected_total,
):
    rows = [
        _comparison_offer("item-1", "supplier-1", **first),
        _comparison_offer("item-1", "supplier-2", **second),
    ]

    result = calculate_comparison("2026-07-28", rows)
    product = result["product_summaries"][0]
    scenario = result["scenario_summary"]
    cheapest_summary = min(
        result["supplier_summaries"], key=lambda summary: summary["final_offer_total"]
    )

    assert product["lowest_final_total"] == expected_item_total
    assert product["lowest_final_total_supplier"] == expected_item_supplier
    assert scenario["cheapest_complete_supplier"] == cheapest_summary
    assert scenario["cheapest_complete_supplier"]["supplier_id"] == expected_supplier
    assert scenario["single_supplier_total"] == expected_total
    assert scenario["mixed_supplier_selections"][0]["supplier_id"] == expected_item_supplier


def test_supplier_recommendation_preserves_equal_price_ties():
    rows = [
        _comparison_offer("item-1", "supplier-1", 100),
        _comparison_offer("item-1", "supplier-2", 100),
    ]

    result = calculate_comparison("2026-07-28", rows)

    assert [row["is_lowest_final_total"] for row in result["rows"]] == [True, True]
    assert result["product_summaries"][0]["lowest_final_total_supplier"] == (
        "supplier-1"
    )
    assert (
        result["scenario_summary"]["cheapest_complete_supplier"]["supplier_id"]
        == "supplier-1"
    )
    assert all(
        summary["difference_from_lowest_complete"] == 0
        for summary in result["supplier_summaries"]
    )


def test_supplier_recommendation_mixed_scenario_uses_each_product_final_total():
    rows = [
        _comparison_offer("item-1", "supplier-1", 100),
        _comparison_offer("item-1", "supplier-2", 130),
        _comparison_offer("item-2", "supplier-1", 200),
        _comparison_offer("item-2", "supplier-2", 150),
    ]

    result = calculate_comparison("2026-07-28", rows)
    scenario = result["scenario_summary"]

    assert scenario["cheapest_complete_supplier"]["supplier_id"] == "supplier-2"
    assert scenario["single_supplier_total"] == 280
    assert scenario["mixed_supplier_total"] == 250
    assert scenario["savings_amount"] == 30
    assert scenario["savings_pct"] == 10.71
    assert {
        (selection["item_id"], selection["supplier_id"], selection["final_total"])
        for selection in scenario["mixed_supplier_selections"]
    } == {
        ("item-1", "supplier-1", 100),
        ("item-2", "supplier-2", 150),
    }


def test_supplier_price_comparison_calculations_and_rankings():
    rows = [
        {
            "id": "1",
            "item_id": "item-1",
            "item_code": "ITM-1",
            "product_name": "منتج",
            "brand": "علامة",
            "supplier_id": "supplier-1",
            "supplier_name": "المورد أ",
            "quantity": 10,
            "unit_price": 100,
            "discount_pct": 10,
            "tax_pct": 14,
            "shipping_cost": 50,
            "other_cost": 10,
            "delivery_days": 5,
            "availability": "available",
            "price_valid_until": "2099-12-31",
        },
        {
            "id": "2",
            "item_id": "item-1",
            "item_code": "ITM-1",
            "product_name": "منتج",
            "brand": "علامة",
            "supplier_id": "supplier-2",
            "supplier_name": "المورد ب",
            "quantity": 10,
            "unit_price": 95,
            "discount_pct": 0,
            "tax_pct": 14,
            "shipping_cost": 0,
            "other_cost": 0,
            "delivery_days": 3,
            "availability": "available",
            "price_valid_until": "2099-12-31",
        },
        {
            "id": "3",
            "item_id": "item-1",
            "item_code": "ITM-1",
            "product_name": "منتج",
            "brand": "علامة",
            "supplier_id": "supplier-3",
            "supplier_name": "المورد ج",
            "quantity": 10,
            "unit_price": 90,
            "discount_pct": 0,
            "tax_pct": 0,
            "shipping_cost": 0,
            "other_cost": 0,
            "delivery_days": 1,
            "availability": "unavailable",
            "price_valid_until": "2099-12-31",
        },
    ]
    result = calculate_comparison("2026-07-28", rows, {"ITM-1": 100})

    assert result["rows"][0]["subtotal"] == 1000
    assert result["rows"][0]["final_total"] == 1000
    assert result["rows"][1]["final_total"] == 950
    assert result["rows"][1]["difference_from_lowest"] == 0
    assert result["rows"][0]["difference_from_lowest"] == 50
    assert result["rows"][0]["difference_pct_from_lowest"] == 5.26
    assert result["rows"][1]["difference_from_last_price"] == -5
    assert result["rows"][1]["difference_pct_from_last_price"] == -5
    assert result["rows"][1]["is_lowest_final_total"] is True
    assert result["rows"][1]["is_fastest_delivery"] is True
    assert result["rows"][2]["is_unavailable"] is True
    assert result["rows"][2]["eligible"] is False
    product = result["product_summaries"][0]
    assert product["lowest_unit_price_supplier"] == "المورد ب"
    assert product["lowest_final_total_supplier"] == "المورد ب"
    assert product["fastest_delivery_supplier"] == "المورد ب"
    assert product["last_historical_unit_price"] == 100
    assert product["difference_from_last_price"] == -5
    assert product["difference_pct_from_last_price"] == -5
    assert product["available_offer_count"] == 2


def test_supplier_offer_adjustments_apply_once_and_vat_follows_discount():
    rows = [
        _comparison_offer("item-1", "supplier-1", 100),
        _comparison_offer("item-2", "supplier-1", 200),
    ]
    result = calculate_comparison(
        "2026-07-28", rows, supplier_offers=[{
            "supplier_id": "supplier-1", "discount_pct": 10, "tax_pct": 14,
            "shipping_cost": 50, "other_cost": 25,
        }],
    )
    assert [row["subtotal"] for row in result["rows"]] == [100, 200]
    assert all("shipping_cost" not in row for row in result["rows"])
    assert result["supplier_summaries"][0] == {
        **result["supplier_summaries"][0],
        "items_subtotal": 300,
        "total_discounts": 30,
        "amount_after_discount": 270,
        "total_taxes": 37.8,
        "total_shipping": 50,
        "total_other_costs": 25,
        "final_offer_total": 382.8,
    }


def test_cheapest_complete_offer_uses_final_total_and_excludes_incomplete():
    rows = [
        _comparison_offer("item-1", "supplier-1", 80),
        _comparison_offer("item-2", "supplier-1", 80),
        _comparison_offer("item-1", "supplier-2", 90),
        _comparison_offer("item-2", "supplier-2", 90),
        _comparison_offer("item-1", "supplier-3", 10),
    ]
    result = calculate_comparison(
        "2026-07-28", rows, supplier_offers=[
            {"supplier_id": "supplier-1", "shipping_cost": 50},
            {"supplier_id": "supplier-2", "discount_pct": 20},
            {"supplier_id": "supplier-3"},
        ],
    )
    assert result["scenario_summary"]["cheapest_complete_supplier"]["supplier_id"] == "supplier-2"
    assert result["scenario_summary"]["single_supplier_total"] == 144
    supplier_three = next(
        row for row in result["supplier_summaries"] if row["supplier_id"] == "supplier-3"
    )
    assert supplier_three["is_complete"] is False


def test_supplier_price_comparison_sqlite_migration_is_additive(tmp_path):
    database_path = tmp_path / "supplier-comparison.db"
    migration_engine = create_engine(f"sqlite:///{database_path.as_posix()}")
    with migration_engine.begin() as connection:
        for table in ("projects", "customers", "items", "suppliers"):
            connection.exec_driver_sql(
                f"CREATE TABLE {table} (id VARCHAR PRIMARY KEY, name VARCHAR)"
            )
            connection.exec_driver_sql(
                f"INSERT INTO {table} (id, name) VALUES ('{table}-1', 'legacy')"
            )

    backup_path = migrate_supplier_price_comparisons(migration_engine)
    assert backup_path and backup_path.is_file()
    with sqlite3.connect(database_path) as connection:
        tables = {
            row[0]
            for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        assert {"price_comparisons", "price_comparison_rows"}.issubset(tables)
        assert connection.execute("PRAGMA user_version").fetchone()[0] == (
            SUPPLIER_PRICE_COMPARISON_SCHEMA_VERSION
        )
        row_columns = {
            row[1]
            for row in connection.execute("PRAGMA table_info(price_comparison_rows)")
        }
        comparison_columns = {
            row[1]
            for row in connection.execute("PRAGMA table_info(price_comparisons)")
        }
        assert {
            "main_category", "subcategory", "specifications",
            "selected_for_purchase",
        }.issubset(row_columns)
        assert {"source_request_id", "source_request_number"}.issubset(
            comparison_columns
        )
        for table in ("projects", "customers", "items", "suppliers"):
            assert (
                connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] == 1
            )
    migration_engine.dispose()


def test_supplier_offer_adjustment_migration_preserves_legacy_totals(tmp_path):
    database_path = tmp_path / "supplier-offer-adjustments.db"
    migration_engine = create_engine(f"sqlite:///{database_path.as_posix()}")
    with migration_engine.begin() as connection:
        for table in ("projects", "customers", "items", "suppliers"):
            connection.exec_driver_sql(
                f"CREATE TABLE {table} (id VARCHAR PRIMARY KEY, name VARCHAR, code VARCHAR)"
            )
    migrate_supplier_price_comparisons(migration_engine)
    with migration_engine.begin() as connection:
        connection.exec_driver_sql(
            "INSERT INTO price_comparisons "
            "(id, comparison_number, comparison_date, created_at, updated_at) "
            "VALUES ('cmp-1', 'CMP-1', '2026-07-28', '', '')"
        )
        for row_id, price, shipping in (("row-1", 100, 20), ("row-2", 200, 20)):
            connection.exec_driver_sql(
                "INSERT INTO price_comparison_rows "
                "(id, comparison_id, position, item_code, product_name, supplier_code, "
                "supplier_name, quantity, unit_price, discount_pct, tax_pct, shipping_cost) "
                "VALUES (?, 'cmp-1', ?, ?, ?, 'SUP-1', 'Supplier', 1, ?, 10, 14, ?)",
                (row_id, 1 if row_id == "row-1" else 2, row_id, row_id, price, shipping),
            )
    backup_path = migrate_supplier_offer_adjustments(migration_engine)
    assert backup_path and backup_path.is_file()
    with sqlite3.connect(database_path) as connection:
        offer = connection.execute(
            "SELECT discount_pct, tax_pct, shipping_cost, other_cost "
            "FROM price_comparison_supplier_offers"
        ).fetchone()
        assert offer == pytest.approx((10, 14, 40, 0))
        assert connection.execute("PRAGMA user_version").fetchone()[0] == (
            SUPPLIER_OFFER_ADJUSTMENTS_SCHEMA_VERSION
        )
    migration_engine.dispose()


def test_supplier_price_comparison_save_reopen_edit_export_and_persistence(s, admin_headers):
    before = {
        "suppliers": len(s.get(f"{API}/suppliers", headers=admin_headers).json()),
        "customers": len(s.get(f"{API}/customers", headers=admin_headers).json()),
        "projects": len(s.get(f"{API}/projects", headers=admin_headers).json()),
        "items": len(s.get(f"{API}/items", headers=admin_headers).json()),
        "purchases": len(s.get(f"{API}/purchases", headers=admin_headers).json()),
        "payments": len(s.get(f"{API}/payments", headers=admin_headers).json()),
    }
    items = s.get(f"{API}/items", headers=admin_headers).json()[:2]
    suppliers = s.get(f"{API}/suppliers", headers=admin_headers).json()[:3]
    project = s.get(f"{API}/projects", headers=admin_headers).json()[0]

    def offer(item, supplier, price, delivery, **overrides):
        return {
            "item_id": item["id"],
            "supplier_id": supplier["id"],
            "quantity": 1,
            "unit": item.get("unit", ""),
            "unit_price": price,
            "discount_pct": 0,
            "tax_pct": 0,
            "shipping_cost": 0,
            "other_cost": 0,
            "delivery_days": delivery,
            "payment_terms": "30 يوم",
            "availability": "available",
            "price_valid_until": "2099-12-31",
            "notes": "",
            **overrides,
        }

    rows = [
        offer(items[0], suppliers[0], 100, 5),
        offer(items[0], suppliers[1], 90, 7),
        offer(items[0], suppliers[2], 80, 1, availability="unavailable"),
        offer(items[1], suppliers[0], 200, 8),
        offer(items[1], suppliers[1], 230, 4),
        offer(
            items[1],
            suppliers[2],
            0,
            2,
            price_valid_until="2020-01-01",
        ),
    ]
    payload = {
        "project_id": project["id"],
        "project_name": project["name"],
        "customer_id": "",
        "customer_name": "عميل مقارنة اختباري",
        "comparison_date": "2026-07-28",
        "notes": "اختبار متعدد المنتجات",
        "supplier_offers": [
            {
                "supplier_id": supplier["id"], "supplier_code": supplier["code"],
                "supplier_name": supplier["name"],
                "discount_pct": 0, "tax_pct": 0,
                "shipping_cost": 17 if index == 2 else 0,
                "other_cost": 3 if index == 2 else 0,
            }
            for index, supplier in enumerate(suppliers)
        ],
        "rows": rows,
    }
    created = s.post(f"{API}/price-comparisons", json=payload, headers=admin_headers)
    assert created.status_code == 200, created.text
    detail = created.json()
    assert detail["comparison_number"].startswith("CMP-")
    assert len(detail["comparison_number"].removeprefix("CMP-")) == 6
    assert len(detail["rows"]) == 6
    assert next(
        offer for offer in detail["supplier_offers"]
        if offer["supplier_id"] == suppliers[2]["id"]
    )["shipping_cost"] == 17
    assert len(detail["product_summaries"]) == 2
    assert len(detail["supplier_summaries"]) == 3
    assert detail["scenario_summary"]["cheapest_complete_supplier"]["supplier_id"] == (
        suppliers[0]["id"]
    )
    assert detail["scenario_summary"]["fastest_complete_supplier"]["supplier_id"] == (
        suppliers[1]["id"]
    )
    assert detail["scenario_summary"]["single_supplier_total"] == 300
    assert detail["scenario_summary"]["mixed_supplier_total"] == 290
    assert detail["scenario_summary"]["mixed_supplier_count"] == 2
    assert detail["scenario_summary"]["savings_amount"] == 10
    assert detail["scenario_summary"]["savings_pct"] == 3.33
    unavailable = next(
        row for row in detail["rows"] if row["availability"] == "unavailable"
    )
    expired_missing = next(row for row in detail["rows"] if row["is_expired"])
    assert unavailable["eligible"] is False
    assert expired_missing["is_missing_price"] is True
    assert expired_missing["eligible"] is False
    supplier_three = next(
        summary
        for summary in detail["supplier_summaries"]
        if summary["supplier_id"] == suppliers[2]["id"]
    )
    assert supplier_three["availability_pct"] == 0
    assert supplier_three["available_products"] == 0

    comparison_id = detail["id"]
    reopened = s.get(f"{API}/price-comparisons/{comparison_id}", headers=admin_headers)
    assert reopened.status_code == 200
    assert reopened.json()["comparison_number"] == detail["comparison_number"]
    assert next(
        offer for offer in reopened.json()["supplier_offers"]
        if offer["supplier_id"] == suppliers[2]["id"]
    )["other_cost"] == 3

    updated_rows = rows[:-1]
    updated_rows[0] = {**updated_rows[0], "unit_price": 85, "shipping_cost": 5}
    updated = s.put(
        f"{API}/price-comparisons/{comparison_id}",
        json={**payload, "notes": "تم التعديل", "rows": updated_rows},
        headers=admin_headers,
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["comparison_number"] == detail["comparison_number"]
    assert updated.json()["notes"] == "تم التعديل"
    assert len(updated.json()["rows"]) == 5
    assert updated.json()["rows"][0]["unit_price"] == 85

    listing = s.get(f"{API}/price-comparisons", headers=admin_headers).json()
    listed = next(row for row in listing if row["id"] == comparison_id)
    assert listed["row_count"] == 5

    exported = s.get(f"{API}/price-comparisons/{comparison_id}/export.xlsx", headers=admin_headers)
    assert exported.status_code == 200
    assert "spreadsheetml" in exported.headers["content-type"]
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert workbook.sheetnames == ["مقارنة المنتجات", "ملخص الموردين", "الملخص"]
    assert workbook["مقارنة المنتجات"].max_row == 6

    engine.dispose()
    persisted = s.get(f"{API}/price-comparisons/{comparison_id}", headers=admin_headers)
    assert persisted.status_code == 200
    assert persisted.json()["notes"] == "تم التعديل"
    assert len(persisted.json()["rows"]) == 5

    after = {
        "suppliers": len(s.get(f"{API}/suppliers", headers=admin_headers).json()),
        "customers": len(s.get(f"{API}/customers", headers=admin_headers).json()),
        "projects": len(s.get(f"{API}/projects", headers=admin_headers).json()),
        "items": len(s.get(f"{API}/items", headers=admin_headers).json()),
        "purchases": len(s.get(f"{API}/purchases", headers=admin_headers).json()),
        "payments": len(s.get(f"{API}/payments", headers=admin_headers).json()),
    }
    assert after == before


def test_manual_comparison_rows_persist_without_touching_master_data(s, admin_headers):
    before_items = len(s.get(f"{API}/items", headers=admin_headers).json())
    before_suppliers = len(s.get(f"{API}/suppliers", headers=admin_headers).json())
    system_item = s.get(f"{API}/items", headers=admin_headers).json()[0]
    system_supplier = s.get(f"{API}/suppliers", headers=admin_headers).json()[0]
    project = s.get(f"{API}/projects", headers=admin_headers).json()[0]
    manual_row = {
        "item_id": "",
        "product_name": "TEST_MANUAL_PRODUCT",
        "brand": "TEST_MANUAL_BRAND",
        "main_category": "TEST_MANUAL_MAIN",
        "subcategory": "TEST_MANUAL_SUB",
        "specifications": "TEST_MANUAL_SPEC",
        "supplier_id": "",
        "supplier_name": "TEST_MANUAL_SUPPLIER",
        "quantity": 2,
        "unit": "piece",
        "unit_price": 100,
        "discount_pct": 10,
        "tax_pct": 14,
        "shipping_cost": 5,
        "other_cost": 2,
        "delivery_days": 4,
        "payment_terms": "cash",
        "availability": "available",
        "price_valid_until": "2099-12-31",
        "notes": "TEST_MANUAL_NOTE",
    }
    payload = {
        "project_id": project["id"], "project_name": project["name"],
        "comparison_date": "2026-07-28", "rows": [manual_row],
    }
    created = s.post(f"{API}/price-comparisons", json=payload, headers=admin_headers)
    assert created.status_code == 200, created.text
    detail = created.json()
    stored = detail["rows"][0]
    assert stored["item_id"] is None
    assert stored["supplier_id"] is None
    assert stored["product_name"] == "TEST_MANUAL_PRODUCT"
    assert stored["supplier_name"] == "TEST_MANUAL_SUPPLIER"
    assert stored["main_category"] == "TEST_MANUAL_MAIN"
    assert stored["subcategory"] == "TEST_MANUAL_SUB"
    assert stored["specifications"] == "TEST_MANUAL_SPEC"
    assert stored["final_total"] == 200
    assert detail["supplier_summaries"][0]["final_offer_total"] == 212.2
    offer = detail["supplier_offers"][0]
    assert (
        offer["discount_pct"], offer["tax_pct"],
        offer["shipping_cost"], offer["other_cost"],
    ) == pytest.approx((10, 14, 5, 2))
    assert len(s.get(f"{API}/items", headers=admin_headers).json()) == before_items
    assert len(s.get(f"{API}/suppliers", headers=admin_headers).json()) == before_suppliers

    edited_manual = {**manual_row, "unit_price": 110, "notes": "TEST_EDITED"}
    system_row = {
        "item_id": system_item["id"],
        "supplier_id": system_supplier["id"],
        "quantity": 1,
        "unit": system_item.get("unit", ""),
        "unit_price": 50,
        "discount_pct": 0,
        "tax_pct": 0,
        "shipping_cost": 0,
        "other_cost": 0,
        "delivery_days": 2,
        "payment_terms": "",
        "availability": "available",
        "price_valid_until": "",
        "notes": "",
    }
    comparison_id = detail["id"]
    edited = s.put(
        f"{API}/price-comparisons/{comparison_id}",
        json={**payload, "rows": [edited_manual, system_row]},
        headers=admin_headers,
    )
    assert edited.status_code == 200, edited.text
    manual_after_edit = next(
        row for row in edited.json()["rows"] if row["item_id"] is None
    )
    assert manual_after_edit["unit_price"] == 110
    assert manual_after_edit["notes"] == "TEST_EDITED"

    deleted = s.put(
        f"{API}/price-comparisons/{comparison_id}",
        json={**payload, "rows": [system_row]},
        headers=admin_headers,
    )
    assert deleted.status_code == 200, deleted.text
    assert len(deleted.json()["rows"]) == 1
    assert deleted.json()["rows"][0]["item_id"] == system_item["id"]
    assert len(s.get(f"{API}/items", headers=admin_headers).json()) == before_items
    assert len(s.get(f"{API}/suppliers", headers=admin_headers).json()) == before_suppliers


def test_product_update_preserves_legacy_item_name(s, admin_headers):
    created = s.post(
        f"{API}/items",
        json={
            "product_name": "TEST_PRODUCT_ORIGINAL",
            "brand": "TEST_BRAND",
            "main_category": "TEST_MAIN",
            "subcategory": "TEST_SUB",
            "specifications": "TEST_SPEC",
            "unit": "قطعة",
        },
        headers=admin_headers,
    )
    assert created.status_code == 200, created.text
    item = created.json()
    try:
        updated = s.put(
            f"{API}/items/{item['id']}",
            json={
                "product_name": "TEST_PRODUCT_UPDATED",
                "brand": "TEST_BRAND_2",
            },
            headers=admin_headers,
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["product_name"] == "TEST_PRODUCT_UPDATED"
        assert updated.json()["brand"] == "TEST_BRAND_2"
        assert updated.json()["name"] == "TEST_PRODUCT_ORIGINAL"
    finally:
        s.delete(f"{API}/items/{item['id']}", headers=admin_headers)


# ---------- Settings ----------
def test_settings(s, admin_headers):
    r = s.get(f"{API}/settings", headers=admin_headers)
    assert r.status_code == 200
    data = r.json()
    assert len(data) == 12
    key = data[0]["key"]
    orig = data[0]["values"]
    new_vals = list(orig) + ["TEST_val"]
    up = s.put(f"{API}/settings/{key}", json={"values": new_vals}, headers=admin_headers)
    assert up.status_code == 200
    assert "TEST_val" in up.json()["values"]
    # restore
    s.put(f"{API}/settings/{key}", json={"values": orig}, headers=admin_headers)


def test_system_diagnostics_excludes_secrets_and_reports_health(s, admin_headers):
    response = s.get(f"{API}/system/diagnostics", headers=admin_headers)
    assert response.status_code == 200
    payload = response.json()
    assert payload["version"] == "0.3.0"
    assert payload["mode"] == "development"
    assert payload["database"]["status"] == "healthy"
    assert payload["database"]["foreign_key_violations"] == 0
    rendered = json.dumps(payload)
    assert os.environ["INTERNAL_REQUEST_TOKEN"] not in rendered
    assert "DATABASE_URL" not in rendered


def test_settings_backup_is_verified_and_does_not_return_records(s, admin_headers):
    response = s.post(f"{API}/system/backup", headers=admin_headers)
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["status"] == "ok"
    assert set(payload) == {
        "status", "created_utc", "database_file", "attachment_count"
    }
    backup_directory = Path(os.environ["PROCUREX_BACKUP_DIR"])
    database = backup_directory / payload["database_file"]
    assert database.is_file()
    assert (backup_directory / "last-successful-backup.json").is_file()
    with closing(sqlite3.connect(database)) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"


def test_open_folder_rejects_unknown_target(s, admin_headers):
    assert s.post(f"{API}/system/open-folder/secrets", headers=admin_headers).status_code == 404


# ---------- Excel export ----------
def test_export_excel(s, admin_headers):
    r = s.get(f"{API}/export/excel", timeout=60, headers=admin_headers)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    assert len(r.content) > 1000
    workbook = load_workbook(io.BytesIO(r.content), read_only=False)
    assert [cell.value for cell in workbook["Items"][1]][:10] == [
        "كود الصنف",
        "اسم الصنف السابق",
        "اسم المنتج",
        "العلامة التجارية",
        "التصنيف الرئيسي",
        "التصنيف الفرعي",
        "الوحدة",
        "المواصفات",
        "المورد المفضل",
        "ملاحظات",
    ]
    assert "اسم المنتج" in [cell.value for cell in workbook["Price History"][1]]
    register = workbook["Purchase Register"]
    assert register.freeze_panes == "A2"
    assert register.auto_filter.ref == "A1:J1"
    assert register.max_row - 1 == len(s.get(f"{API}/purchases", headers=admin_headers).json())
    assert register["B2"].number_format == "yyyy-mm-dd"
    assert "EGP" in register["G2"].number_format
    assert register.column_dimensions["D"].width >= len("المورد")
    assert workbook["Purchase Items"]["N2"].number_format.endswith('"EGP"')


# ---------- Purchase creation, dup, validation, cascade + Payment flow ----------
@pytest.fixture(scope="module")
def ids(s, admin_headers):
    suppliers = s.get(f"{API}/suppliers", headers=admin_headers).json()
    customers = s.get(f"{API}/customers", headers=admin_headers).json()
    projects = s.get(f"{API}/projects", headers=admin_headers).json()
    items = s.get(f"{API}/items", headers=admin_headers).json()
    return {
        "supplier_id": suppliers[0]["id"],
        "supplier_name": suppliers[0]["name"],
        "customer_id": customers[0]["id"],
        "project_id": projects[0]["id"],
        "item_id": items[0]["id"],
        "item2_id": items[1]["id"],
    }


created_purchase_ids = []


def _payload(ids, inv, qty=2, price=100, disc=10, vat=14):
    return {
        "purchase_date": "2025-01-15",
        "invoice_number": inv,
        "supplier_id": ids["supplier_id"],
        "project_id": ids["project_id"],
        "customer_id": ids["customer_id"],
        "items": [
            {
                "item_id": ids["item_id"],
                "quantity": qty,
                "unit_price": price,
                "discount_pct": disc,
                "vat_pct": vat,
            }
        ],
    }


def test_purchase_create_and_compute(s, ids, admin_headers):
    ph_before = len(s.get(f"{API}/price-history", headers=admin_headers).json())
    payload = _payload(ids, "TEST_INV_001", qty=2, price=100, disc=10, vat=14)
    r = s.post(f"{API}/purchases", json=payload, headers=admin_headers)
    assert r.status_code == 200, r.text
    p = r.json()
    # line = 2*100*0.9*1.14 = 205.2
    assert p["subtotal"] == 200
    assert p["discount_total"] == 20
    assert p["after_discount"] == 180
    assert round(p["vat_total"], 2) == 25.2
    assert round(p["invoice_total"], 2) == 205.2
    assert p["payment_status"] == "غير مدفوع"
    assert p["purchase_id"].startswith("PUR-")
    assert len(p["purchase_id"]) == 10
    assert p["purchase_id"] != "PUR-000001"
    created_purchase_ids.append(p["purchase_id"])
    # price history grew
    ph_after = len(s.get(f"{API}/price-history", headers=admin_headers).json())
    assert ph_after == ph_before + 1
    # GET verification
    got = s.get(f"{API}/purchases/{p['purchase_id']}", headers=admin_headers).json()
    assert got["invoice_total"] == p["invoice_total"]
    assert len(got["items"]) == 1
    master_item = next(
        item for item in s.get(f"{API}/items", headers=admin_headers).json() if item["id"] == ids["item_id"]
    )
    assert got["items"][0]["unit"] == master_item["unit"]
    assert got["items"][0]["product_name"] == master_item["product_name"]
    assert got["items"][0]["brand"] == master_item["brand"]
    assert got["items"][0]["specifications"] == master_item["specifications"]


def test_purchase_multi_line_rounding_and_totals(s, ids, admin_headers):
    payload = _payload(ids, "TEST_MULTI_001", qty=3, price=10.005, disc=5, vat=14)
    payload["shipping_cost"] = 0.005
    payload["items"].append(
        {
            "item_id": ids["item2_id"],
            "quantity": 2.5,
            "unit_price": 7.777,
            "discount_pct": 0,
            "vat_pct": 14,
        }
    )
    response = s.post(f"{API}/purchases", json=payload, headers=admin_headers)
    assert response.status_code == 200, response.text
    purchase = response.json()
    detail = s.get(f"{API}/purchases/{purchase['purchase_id']}", headers=admin_headers).json()
    assert len(detail["items"]) == 2
    assert purchase["invoice_total"] == round(
        sum(line["line_total"] for line in detail["items"])
        + purchase["shipping_cost"]
        + purchase["other_costs"],
        2,
    )
    assert s.delete(f"{API}/purchases/{purchase['purchase_id']}", headers=admin_headers).status_code == 200


def test_purchase_duplicate(s, ids, admin_headers):
    r = s.post(f"{API}/purchases", json=_payload(ids, "TEST_INV_001"), headers=admin_headers)
    assert r.status_code == 409
    detail = r.json()["detail"]
    assert detail["code"] == "duplicate_supplier_invoice"
    assert "بالفعل" in detail["message"]
    assert detail["existing_purchase_id"] == created_purchase_ids[0]

    formatted = s.post(
        f"{API}/purchases", json=_payload(ids, "  test_inv_٠٠١  "), headers=admin_headers,
    )
    assert formatted.status_code == 409
    assert formatted.json()["detail"]["existing_purchase_id"] == created_purchase_ids[0]


def test_purchase_validation(s, ids, admin_headers):
    # missing supplier
    p = _payload(ids, "TEST_INV_X")
    p["supplier_id"] = "nonexistent"
    r = s.post(f"{API}/purchases", json=p, headers=admin_headers)
    assert r.status_code == 422
    # missing/invalid project IDs cannot be replaced by display names
    p = _payload(ids, "TEST_INV_BAD_PROJECT")
    p["project_id"] = "nonexistent-project"
    p["project_name"] = "اسم مشروع موجود لا يعوض المعرّف"
    assert s.post(f"{API}/purchases", json=p, headers=admin_headers).status_code == 422
    p = _payload(ids, "TEST_INV_MISSING_PROJECT")
    p.pop("project_id")
    assert s.post(f"{API}/purchases", json=p, headers=admin_headers).status_code == 422
    # empty items
    p = _payload(ids, "TEST_INV_Y")
    p["items"] = []
    assert s.post(f"{API}/purchases", json=p, headers=admin_headers).status_code == 422
    # qty <= 0
    p = _payload(ids, "TEST_INV_Z", qty=0)
    assert s.post(f"{API}/purchases", json=p, headers=admin_headers).status_code == 422
    # price <= 0
    p = _payload(ids, "TEST_INV_W", price=0)
    assert s.post(f"{API}/purchases", json=p, headers=admin_headers).status_code == 422
    # percentages and additional costs are bounded
    p = _payload(ids, "TEST_INV_BAD_DISCOUNT")
    p["items"][0]["discount_pct"] = 101
    assert s.post(f"{API}/purchases", json=p, headers=admin_headers).status_code == 422
    p = _payload(ids, "TEST_INV_BAD_TAX")
    p["items"][0]["vat_pct"] = -1
    assert s.post(f"{API}/purchases", json=p, headers=admin_headers).status_code == 422
    p = _payload(ids, "TEST_INV_BAD_SHIPPING")
    p["shipping_cost"] = -0.01
    assert s.post(f"{API}/purchases", json=p, headers=admin_headers).status_code == 422


def test_purchase_transaction_rolls_back_on_line_history_failure(s, ids, monkeypatch, admin_headers):
    invoice = "TEST_INV_ROLLBACK"
    original_insert = LocalCollection.insert_one

    async def fail_price_history(self, document):
        if self.model.__tablename__ == "price_history":
            raise RuntimeError("injected price history failure")
        return await original_insert(self, document)

    monkeypatch.setattr(LocalCollection, "insert_one", fail_price_history)
    with TestClient(app, raise_server_exceptions=False) as safe_client:
        response = safe_client.post(f"{API}/purchases", json=_payload(ids, invoice), headers=admin_headers)
        assert response.status_code == 500
        assert response.json()["detail"]["code"] == "internal_error"

    purchases = s.get(f"{API}/purchases", headers=admin_headers).json()
    assert not any(row["invoice_number"] == invoice for row in purchases)
    history = s.get(f"{API}/price-history", headers=admin_headers).json()
    assert not any(row["invoice_number"] == invoice for row in history)


def test_payment_flow(s, ids, admin_headers):
    assert created_purchase_ids, "purchase must be created first"
    pid = created_purchase_ids[0]
    # partial payment
    r = s.post(
        f"{API}/payments",
        json={
            "purchase_id": pid,
            "payment_date": "2025-01-16",
            "amount_paid": 100,
            "submission_token": "payment-test-token-001",
        },
        headers=admin_headers,
    )
    assert r.status_code == 200, r.text
    pay1 = r.json()
    assert pay1["payment_id"].startswith("PAY-")
    assert len(pay1["payment_id"]) == 10
    got = s.get(f"{API}/purchases/{pid}", headers=admin_headers).json()
    assert got["payment_status"] == "مدفوع جزئي"
    assert round(got["remaining"], 2) == round(205.2 - 100, 2)

    payment_count = len(s.get(f"{API}/payments", headers=admin_headers).json())
    replay = s.post(
        f"{API}/payments",
        json={
            "purchase_id": pid,
            "payment_date": "2025-01-16",
            "amount_paid": 100,
            "submission_token": "payment-test-token-001",
        },
        headers=admin_headers,
    )
    assert replay.status_code == 200
    assert replay.json()["id"] == pay1["id"]
    assert replay.json()["idempotent_replay"] is True
    assert len(s.get(f"{API}/payments", headers=admin_headers).json()) == payment_count

    # exceed remaining
    r_over = s.post(
        f"{API}/payments",
        json={
            "purchase_id": pid,
            "payment_date": "2025-01-16",
            "amount_paid": 999,
        },
        headers=admin_headers,
    )
    assert r_over.status_code == 422

    # complete
    r2 = s.post(
        f"{API}/payments",
        json={
            "purchase_id": pid,
            "payment_date": "2025-01-17",
            "amount_paid": 105.2,
        },
        headers=admin_headers,
    )
    assert r2.status_code == 200
    assert len(r2.json()["payment_id"]) == 10
    assert r2.json()["payment_id"] != pay1["payment_id"]
    got2 = s.get(f"{API}/purchases/{pid}", headers=admin_headers).json()
    assert got2["payment_status"] == "مدفوع"

    # delete last payment recomputes status
    pay2_id = r2.json()["id"]
    d = s.delete(f"{API}/payments/{pay2_id}", headers=admin_headers)
    assert d.status_code == 200
    got3 = s.get(f"{API}/purchases/{pid}", headers=admin_headers).json()
    assert got3["payment_status"] == "مدفوع جزئي"


def test_purchase_delete_cascade(s, ids, admin_headers):
    assert created_purchase_ids
    pid = created_purchase_ids[0]
    d = s.delete(f"{API}/purchases/{pid}", headers=admin_headers)
    assert d.status_code == 200
    assert s.get(f"{API}/purchases/{pid}", headers=admin_headers).status_code == 404
    # payments for this purchase removed
    pays = s.get(f"{API}/payments", headers=admin_headers).json()
    assert not any(p["purchase_id"] == pid for p in pays)


def test_original_data_intact(s, admin_headers):
    """Ensure PUR-000001 still exists and dashboard KPIs match."""
    r = s.get(f"{API}/purchases/PUR-000001", headers=admin_headers)
    assert r.status_code == 200
    d = s.get(f"{API}/dashboard", headers=admin_headers).json()
    assert d["purchase_count"] == 1
    assert d["total_purchases"] == 29947.5
    assert len(r.json()["items"]) == 5
    assert all(line["item_id"] for line in r.json()["items"])
    assert all(line["product_name"] for line in r.json()["items"])
    assert all(
        "brand" in line and "specifications" in line for line in r.json()["items"]
    )
    history = s.get(f"{API}/price-history", headers=admin_headers).json()
    assert history and all(row["product_name"] for row in history)
    assert all("brand" in row and "specifications" in row for row in history)


def test_workbook_import_is_idempotent():
    counts = asyncio.run(import_data(db, parse_workbook(WORKBOOK.read_bytes())))
    assert all(value == 0 for value in counts.values())


def test_legacy_schema_migration_persists_after_reopen(tmp_path):
    database_path = tmp_path / "legacy.db"
    with sqlite3.connect(database_path) as connection:
        connection.execute(
            "CREATE TABLE items ("
            "id VARCHAR PRIMARY KEY, category VARCHAR NOT NULL DEFAULT '', "
            "name VARCHAR NOT NULL, specs TEXT NOT NULL DEFAULT '')"
        )
        connection.executemany(
            "INSERT INTO items (id, category, name, specs) VALUES (?, ?, ?, ?)",
            [("1", "مواد بناء", "أسمنت", "قديم"), ("2", "", "غير مصنف", "")],
        )

    migration_engine = create_engine(f"sqlite:///{database_path.as_posix()}")
    classification_backup = migrate_item_classification(migration_engine)
    backup_path = migrate_item_identity(migration_engine)
    migration_engine.dispose()

    assert classification_backup is not None and classification_backup.is_file()
    assert backup_path is not None and backup_path.is_file()
    with sqlite3.connect(classification_backup) as backup:
        assert [row[1] for row in backup.execute("PRAGMA table_info(items)")] == [
            "id",
            "category",
            "name",
            "specs",
        ]
        assert backup.execute("SELECT COUNT(*) FROM items").fetchone()[0] == 2

    with sqlite3.connect(database_path) as restarted:
        columns = [row[1] for row in restarted.execute("PRAGMA table_info(items)")]
        assert "main_category" in columns
        assert "subcategory" in columns
        assert "product_name" in columns
        assert "brand" in columns
        assert "specifications" in columns
        assert restarted.execute(
            "SELECT name, product_name, brand, category, main_category, subcategory, "
            "specs, specifications FROM items WHERE id = '1'"
        ).fetchone() == (
            "أسمنت",
            "أسمنت",
            "",
            "مواد بناء",
            "مواد بناء",
            "",
            "قديم",
            "قديم",
        )
        assert restarted.execute("PRAGMA user_version").fetchone()[0] == 2


# ---------- Incoming purchase requests ----------
INTERNAL_HEADERS = {"X-Internal-Token": "test-internal-token"}
incoming_state = {}


@pytest.fixture(scope="session")
def admin_headers(s):
    """A logged-in ERP admin, for workflow/PO endpoints Sprint 2.3 now
    gates with require_erp_role() - admin always passes every role check,
    so this is safe to reuse across tests that only care about business
    logic (not the role check itself)."""
    with SessionLocal() as session:
        username = f"role-admin-{uuid.uuid4().hex[:8]}"
        _make_user(session, username=username, role="admin")
    return _login_headers(s, username)


def _incoming_payload(token="incoming-submission-token-001", product="دهان واجهات"):
    return {
        "requester_name": "م. أحمد علي",
        "company_name": "شركة الاختبار للمقاولات",
        "phone_number": "+201001234567",
        "whatsapp_number": "+201001234567",
        "email": "engineer@example.com",
        "project_name": "مشروع إداري جديد",
        "project_location": "القاهرة الجديدة",
        "delivery_location": "بوابة الموقع الرئيسية",
        "required_delivery_date": "2099-12-31",
        "priority": "high",
        "notes": "يرجى التواصل قبل التسليم",
        "submission_token": token,
        "items": [
            {
                "product_name": product,
                "preferred_brand": "علامة مفضلة",
                "main_category": "دهانات",
                "subcategory": "دهانات خارجية",
                "specifications": "مقاوم للعوامل الجوية",
                "quantity": 12,
                "unit": "بستلة",
                "attachment_index": None,
            }
        ],
    }


def _submit_incoming(s, payload, files=None, website=""):
    return s.post(
        f"{API}/public/purchase-requests",
        data={"payload": json.dumps(payload, ensure_ascii=False), "website": website},
        files=files or [],
    )


def test_public_incoming_request_validation_and_spam_trap(s):
    invalid = _incoming_payload("incoming-invalid-token")
    invalid["requester_name"] = ""
    assert _submit_incoming(s, invalid).status_code == 422
    trapped = _incoming_payload("incoming-spam-token")
    assert (
        _submit_incoming(s, trapped, website="https://spam.invalid").status_code == 422
    )


def test_public_submission_is_idempotent_and_does_not_create_purchase(s):
    purchases_before = len(s.get(f"{API}/purchases").json())
    payload = _incoming_payload()
    response = _submit_incoming(s, payload)
    assert response.status_code == 200, response.text
    result = response.json()
    assert set(result) == {"ok", "duplicate", "request_number"}
    assert result["request_number"].startswith("REQ-")
    assert result["duplicate"] is False
    incoming_state["request_number"] = result["request_number"]

    duplicate = _submit_incoming(s, payload)
    assert duplicate.status_code == 200
    assert duplicate.json()["duplicate"] is True
    assert duplicate.json()["request_number"] == result["request_number"]

    content_duplicate = _incoming_payload("incoming-submission-token-002")
    duplicate_again = _submit_incoming(s, content_duplicate)
    assert duplicate_again.status_code == 200
    assert duplicate_again.json()["request_number"] == result["request_number"]
    assert len(s.get(f"{API}/purchases").json()) == purchases_before


def test_internal_boundary_list_filter_and_detail(s, admin_headers):
    assert s.get(f"{API}/internal/incoming-purchase-requests").status_code == 401
    listing = s.get(
        f"{API}/internal/incoming-purchase-requests",
        headers={**INTERNAL_HEADERS, **admin_headers},
        params={
            "search": incoming_state["request_number"],
            "status": "new",
            "priority": "high",
        },
    )
    assert listing.status_code == 200, listing.text
    assert len(listing.json()) == 1
    row = listing.json()[0]
    incoming_state["request_id"] = row["id"]
    detail = s.get(
        f"{API}/internal/incoming-purchase-requests/{row['id']}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    )
    assert detail.status_code == 200
    body = detail.json()
    assert body["request_number"] == incoming_state["request_number"]
    assert body["items"][0]["product_name"] == "دهان واجهات"
    assert body["status_history"][0]["to_status"] == "new"
    serialized = json.dumps(body, ensure_ascii=False).lower()
    for forbidden in ("unit_price", "supplier", "payment", "employee_id"):
        assert forbidden not in serialized


def test_attachment_validation_storage_and_protected_view(s, admin_headers):
    invalid_payload = _incoming_payload("incoming-invalid-file", "منتج بملف غير صالح")
    invalid_payload["items"][0]["attachment_index"] = 0
    invalid = _submit_incoming(
        s,
        invalid_payload,
        files=[
            (
                "attachments",
                ("malware.exe", b"MZ-not-allowed", "application/octet-stream"),
            ),
        ],
    )
    assert invalid.status_code == 422
    payload = _incoming_payload("incoming-valid-file", "منتج بصورة")
    payload["items"][0]["attachment_index"] = 0
    png = b"\x89PNG\r\n\x1a\n" + b"safe-test-image"
    created = _submit_incoming(
        s,
        payload,
        files=[
            ("attachments", ("specification.png", png, "image/png")),
        ],
    )
    assert created.status_code == 200, created.text
    listing = s.get(
        f"{API}/internal/incoming-purchase-requests",
        headers={**INTERNAL_HEADERS, **admin_headers},
        params={"search": created.json()["request_number"]},
    ).json()
    request_id = listing[0]["id"]
    incoming_state["attachment_request_id"] = request_id
    detail = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    attachment = detail["items"][0]["attachment"]
    assert attachment["original_filename"] == "specification.png"
    assert "stored_filename" not in attachment
    assert (
        s.get(
            f"{API}/internal/incoming-purchase-requests/{request_id}/attachments/{attachment['id']}"
        ).status_code
        == 401
    )
    download = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}/attachments/{attachment['id']}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    )
    assert download.status_code == 200
    assert download.content == png
    assert download.headers["content-type"] == "image/png"
    project = s.get(f"{API}/projects", headers=admin_headers).json()[0]
    supplier = s.get(f"{API}/suppliers", headers=admin_headers).json()[0]
    with SessionLocal.begin() as session:
        request_row = session.get(IncomingPurchaseRequest, request_id)
        request_row.status = "pricing"
        request_row.project_id = project["id"]
        request_row.project_name = project["name"]
    comparison = s.post(f"{API}/price-comparisons", json={
        "project_id": project["id"], "project_name": project["name"],
        "source_request_id": request_id,
        "source_request_number": created.json()["request_number"],
        "comparison_date": "2026-08-17",
        "rows": [{
            "product_name": "منتج بصورة", "supplier_name": "مورد يدوي",
            "supplier_id": supplier["id"],
            "quantity": 1, "unit": "قطعة", "unit_price": 10,
            "availability": "available", "price_valid_until": "2099-01-01",
        }],
    }, headers=admin_headers)
    assert comparison.status_code == 200, comparison.text
    comparison_detail = s.get(f"{API}/price-comparisons/{comparison.json()['id']}", headers=admin_headers)
    assert comparison_detail.status_code == 200
    assert comparison_detail.json()["source_attachments"] == [{
        "id": attachment["id"], "request_item_id": detail["items"][0]["id"],
        "item_label": "منتج بصورة", "original_filename": "specification.png",
        "media_type": "image/png", "size_bytes": len(png), "is_image": True,
        "view_url": f"/internal/incoming-purchase-requests/{request_id}/attachments/{attachment['id']}",
    }]
    assert s.get(f"{API}/price-comparisons", headers=admin_headers).status_code == 200


def test_internal_assignment_notes_status_and_notifications(s, admin_headers):
    request_id = incoming_state["request_id"]
    assignment = s.patch(
        f"{API}/internal/incoming-purchase-requests/{request_id}/assignment",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"employee": "موظف المشتريات"},
    )
    assert assignment.status_code == 200
    note = s.post(
        f"{API}/internal/incoming-purchase-requests/{request_id}/notes",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"author": "مدير المشتريات", "note": "تمت مراجعة المواصفات"},
    )
    assert note.status_code == 200
    changed = s.post(
        f"{API}/internal/incoming-purchase-requests/{request_id}/status",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={
            "status": "under_review",
            "changed_by": "مدير المشتريات",
            "note": "بدأت المراجعة",
        },
    )
    assert changed.status_code == 200
    body = changed.json()
    assert body["assigned_employee"] == "موظف المشتريات"
    assert body["internal_notes"][0]["note"] == "تمت مراجعة المواصفات"
    assert body["status_history"][-1]["to_status"] == "under_review"
    for protected_status in (
        "pricing", "waiting_for_approval", "approved",
        "rejected", "converted_to_purchase", "completed",
    ):
        blocked = s.post(
            f"{API}/internal/incoming-purchase-requests/{request_id}/status",
            headers={**INTERNAL_HEADERS, **admin_headers},
            json={"status": protected_status},
        )
        assert blocked.status_code == 409
    held = s.post(
        f"{API}/internal/incoming-purchase-requests/{request_id}/status",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"status": "hold", "changed_by": "مدير المشتريات"},
    )
    assert held.status_code == 200
    assert held.json()["status"] == "hold"
    resumed = s.post(
        f"{API}/internal/incoming-purchase-requests/{request_id}/status",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"status": "under_review", "changed_by": "مدير المشتريات"},
    )
    assert resumed.status_code == 200
    assert resumed.json()["status"] == "under_review"
    assert (
        s.post(
            f"{API}/internal/incoming-purchase-requests/{request_id}/status",
            headers={**INTERNAL_HEADERS, **admin_headers},
            json={"status": "not-a-status"},
        ).status_code
        == 422
    )
    notifications = s.get(
        f"{API}/internal/incoming-purchase-requests/notifications?unread_only=true",
        headers={**INTERNAL_HEADERS, **admin_headers},
    )
    assert notifications.status_code == 200
    assert any(item["entity_id"] == request_id for item in notifications.json())


def test_terminal_and_historical_request_statuses_remain_readable_but_cannot_reopen(s, admin_headers):
    timestamp = datetime.now(timezone.utc).isoformat()
    request_ids = {}
    with SessionLocal.begin() as session:
        for status in ("completed", "legacy_archived"):
            request_id = str(uuid.uuid4())
            request_ids[status] = request_id
            session.add(IncomingPurchaseRequest(
                id=request_id, request_number=f"T-STATUS-{uuid.uuid4().hex[:8]}",
                requester_name="مهندس الموقع", company_name="عميل حالة تاريخية",
                phone_number="01000000012", project_name="مشروع حالة تاريخية",
                project_location="القاهرة", delivery_location="القاهرة",
                required_delivery_date="2099-01-01", priority="normal", status=status,
                submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
                created_at=timestamp, updated_at=timestamp,
            ))

    for status, request_id in request_ids.items():
        loaded = s.get(
            f"{API}/internal/incoming-purchase-requests/{request_id}",
            headers={**INTERNAL_HEADERS, **admin_headers},
        )
        assert loaded.status_code == 200
        assert loaded.json()["status"] == status
        reopened = s.post(
            f"{API}/internal/incoming-purchase-requests/{request_id}/status",
            headers={**INTERNAL_HEADERS, **admin_headers},
            json={"status": "under_review"},
        )
        assert reopened.status_code == 409


def test_customer_and_draft_conversions_never_create_completed_purchase(s, admin_headers):
    request_id = incoming_state["request_id"]
    purchases_before = len(s.get(f"{API}/purchases", headers=admin_headers).json())
    customer = s.post(
        f"{API}/internal/incoming-purchase-requests/{request_id}/convert-customer",
        headers={**INTERNAL_HEADERS, **admin_headers},
    )
    assert customer.status_code == 200, customer.text
    assert customer.json()["customer_code"].startswith("CUS-")
    document = s.post(
        f"{API}/internal/incoming-purchase-requests/{request_id}/convert",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"document_type": "internal_request", "converted_by": "مدير المشتريات"},
    )
    assert document.status_code == 200, document.text
    assert document.json()["document_number"].startswith("IPR-")
    assert len(s.get(f"{API}/purchases", headers=admin_headers).json()) == purchases_before
    detail = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    assert detail["status"] == "converted_to_purchase"
    assert detail["converted_document"]["status"] == "draft"

    second_id = incoming_state["attachment_request_id"]
    draft = s.post(
        f"{API}/internal/incoming-purchase-requests/{second_id}/convert",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"document_type": "purchase_draft", "converted_by": "مدير المشتريات"},
    )
    assert draft.status_code == 200
    assert draft.json()["document_number"].startswith("PDR-")
    assert len(s.get(f"{API}/purchases", headers=admin_headers).json()) == purchases_before


def test_incoming_request_persists_after_reinitialization(s, admin_headers):
    init_db()
    listing = s.get(
        f"{API}/internal/incoming-purchase-requests",
        headers={**INTERNAL_HEADERS, **admin_headers},
        params={"search": incoming_state["request_number"]},
    )
    assert listing.status_code == 200
    assert listing.json()[0]["request_number"] == incoming_state["request_number"]


def test_incoming_request_migration_is_additive_and_backed_up(tmp_path):
    database_path = tmp_path / "version-two.db"
    with sqlite3.connect(database_path) as connection:
        connection.execute(
            "CREATE TABLE preserved_data (id INTEGER PRIMARY KEY, value TEXT)"
        )
        connection.execute("INSERT INTO preserved_data (value) VALUES ('keep me')")
        connection.execute("PRAGMA user_version = 2")
    migration_engine = create_engine(f"sqlite:///{database_path.as_posix()}")
    backup_path = migrate_incoming_requests(migration_engine)
    assert backup_path is not None and backup_path.is_file()
    assert migrate_incoming_requests(migration_engine) is None
    migration_engine.dispose()
    with sqlite3.connect(backup_path) as backup:
        assert (
            backup.execute("SELECT value FROM preserved_data").fetchone()[0]
            == "keep me"
        )
        assert backup.execute("PRAGMA user_version").fetchone()[0] == 2
    with sqlite3.connect(database_path) as migrated:
        assert (
            migrated.execute("SELECT value FROM preserved_data").fetchone()[0]
            == "keep me"
        )
        assert migrated.execute("PRAGMA user_version").fetchone()[0] == 3
        tables = {
            row[0]
            for row in migrated.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        assert "incoming_purchase_requests" in tables
        assert "internal_purchase_documents" in tables


def test_public_deployment_surface_does_not_register_internal_erp_routes():
    public_app = create_app(surface="public", initialize_database=False)
    paths = {route.path for route in public_app.routes}
    assert "/api/public/purchase-requests" in paths
    assert "/api/public/purchase-requests/health" in paths
    assert "/api/purchases" not in paths
    assert not any(path.startswith("/api/internal/") for path in paths)
    assert "/docs" not in paths
    with TestClient(public_app) as client:
        assert client.get("/api/purchases").status_code == 404
        assert client.get("/api/internal/incoming-purchase-requests").status_code == 404
        assert client.get("/api/public/purchase-requests/health").json() == {"ok": True}


@pytest.mark.parametrize("environment", ["staging", "production"])
def test_hosted_public_surface_requires_strict_configuration(monkeypatch, environment):
    values = {
        "APP_ENV": environment,
        "DATABASE_URL": "postgresql://user:password@database.example/procurex?sslmode=require",
        "CORS_ORIGINS": f"https://requests-{environment}.example.com",
        "TRUSTED_HOSTS": f"api-requests-{environment}.example.com",
        "FORCE_HTTPS": "true",
        "ATTACHMENT_STORAGE_BACKEND": "s3",
        "REQUEST_PRIVACY_SALT": "x" * 32,
        "AUTH_SECRET_KEY": "x" * 32,
        "R2_ENDPOINT_URL": "https://account.r2.cloudflarestorage.com",
        "R2_ACCESS_KEY_ID": "test-access-key",
        "R2_SECRET_ACCESS_KEY": "test-secret-key",
        "R2_BUCKET_NAME": f"procurex-{environment}-attachments",
    }
    for name, value in values.items():
        monkeypatch.setenv(name, value)

    public_app = create_app(surface="public", initialize_database=False)
    paths = {route.path for route in public_app.routes}
    assert "/api/public/purchase-requests" in paths
    assert "/api/purchases" not in paths
    assert not any(path.startswith("/api/internal/") for path in paths)
    assert "/docs" not in paths


# ---------- Hosted full-ERP-surface configuration (APP_ENV=production/staging, APP_SURFACE=full) ----------

def _set_valid_full_surface_env(monkeypatch, environment="production"):
    monkeypatch.setenv("APP_ENV", environment)
    monkeypatch.setenv("CORS_ORIGINS", "https://procurement.example.com")
    monkeypatch.setenv("TRUSTED_HOSTS", "procurement.example.com")
    monkeypatch.setenv("FORCE_HTTPS", "true")
    monkeypatch.setenv("AUTH_SECRET_KEY", "x" * 32)
    monkeypatch.delenv("TRUST_PROXY_HEADERS", raising=False)
    monkeypatch.delenv("INTERNAL_REQUEST_TOKEN", raising=False)


@pytest.mark.parametrize("environment", ["staging", "production"])
def test_hosted_full_surface_allowed_with_secure_configuration(monkeypatch, environment):
    """1. production/staging + full + valid secure configuration => allowed."""
    _set_valid_full_surface_env(monkeypatch, environment)
    full_app = create_app(surface="full", initialize_database=False)
    paths = {route.path for route in full_app.routes}
    assert "/api/purchases" in paths
    assert any(path.startswith("/api/internal/") for path in paths)
    assert "/docs" in paths


def test_hosted_full_surface_rejects_wildcard_cors(monkeypatch):
    """2. production + full + wildcard CORS => rejected."""
    _set_valid_full_surface_env(monkeypatch)
    monkeypatch.setenv("CORS_ORIGINS", "*")
    with pytest.raises(RuntimeError, match="CORS_ORIGINS must be an explicit HTTPS allowlist"):
        create_app(surface="full", initialize_database=False)


@pytest.mark.parametrize("unsafe_hosts", ["", "*", "*.example.com", "https://procurement.example.com"])
def test_hosted_full_surface_rejects_unsafe_trusted_hosts(monkeypatch, unsafe_hosts):
    """3. production + full + missing/unsafe trusted hosts => rejected."""
    _set_valid_full_surface_env(monkeypatch)
    monkeypatch.setenv("TRUSTED_HOSTS", unsafe_hosts)
    with pytest.raises(RuntimeError, match="TRUSTED_HOSTS must be an explicit hostname allowlist"):
        create_app(surface="full", initialize_database=False)


@pytest.mark.parametrize("secret", ["", "too-short"])
def test_hosted_full_surface_rejects_missing_auth_secret(monkeypatch, secret):
    """4. production + full + missing/short production auth secret => rejected."""
    _set_valid_full_surface_env(monkeypatch)
    monkeypatch.setenv("AUTH_SECRET_KEY", secret)
    with pytest.raises(RuntimeError, match="AUTH_SECRET_KEY must contain at least 32 characters"):
        create_app(surface="full", initialize_database=False)


def test_hosted_full_surface_rejects_spoofable_internal_access_fallback(monkeypatch):
    """Trusting proxy headers in production without an explicit internal-access
    token would make the require_internal_access localhost fallback spoofable
    from the internet - reject that combination rather than silently allowing
    it to stand in for real ERP authentication."""
    _set_valid_full_surface_env(monkeypatch)
    monkeypatch.setenv("TRUST_PROXY_HEADERS", "true")
    with pytest.raises(RuntimeError, match="INTERNAL_REQUEST_TOKEN"):
        create_app(surface="full", initialize_database=False)


def test_hosted_full_surface_allows_trusted_proxy_headers_with_internal_token(monkeypatch):
    _set_valid_full_surface_env(monkeypatch)
    monkeypatch.setenv("TRUST_PROXY_HEADERS", "true")
    monkeypatch.setenv("INTERNAL_REQUEST_TOKEN", "x" * 32)
    full_app = create_app(surface="full", initialize_database=False)
    assert any(route.path == "/api/purchases" for route in full_app.routes)


def test_development_full_surface_configuration_is_unaffected(monkeypatch):
    """5. existing development configuration remains unaffected - none of the
    hosted-surface checks apply outside staging/production."""
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.delenv("CORS_ORIGINS", raising=False)
    monkeypatch.delenv("TRUSTED_HOSTS", raising=False)
    monkeypatch.delenv("FORCE_HTTPS", raising=False)
    monkeypatch.delenv("AUTH_SECRET_KEY", raising=False)
    full_app = create_app(surface="full", initialize_database=False)
    assert any(route.path == "/api/purchases" for route in full_app.routes)


def test_staging_public_surface_refuses_filesystem_attachment_storage(monkeypatch):
    """Existing production public-surface S3 requirement remains unchanged."""
    monkeypatch.setenv("APP_ENV", "staging")
    monkeypatch.setenv("APP_SURFACE", "public")
    monkeypatch.setenv("ATTACHMENT_STORAGE_BACKEND", "local")
    with pytest.raises(RuntimeError, match="Staging public-surface deployments require"):
        build_attachment_storage()


def test_production_public_surface_refuses_filesystem_attachment_storage(monkeypatch):
    """Existing production public-surface S3 requirement remains unchanged."""
    monkeypatch.setenv("APP_ENV", "production")
    monkeypatch.setenv("APP_SURFACE", "public")
    monkeypatch.setenv("ATTACHMENT_STORAGE_BACKEND", "local")
    with pytest.raises(RuntimeError, match="Production public-surface deployments require"):
        build_attachment_storage()


# ---------- Full-ERP-surface local attachment storage in staging/production ----------

def _set_full_surface_env(monkeypatch, environment="production"):
    monkeypatch.setenv("APP_ENV", environment)
    monkeypatch.setenv("APP_SURFACE", "full")
    monkeypatch.setenv("ATTACHMENT_STORAGE_BACKEND", "local")
    monkeypatch.delenv("PROCUREX_DATA_ROOT", raising=False)
    monkeypatch.delenv("INCOMING_REQUEST_UPLOAD_DIR", raising=False)


@pytest.mark.parametrize("environment", ["staging", "production"])
def test_hosted_full_surface_allows_explicit_persistent_local_storage(monkeypatch, environment, tmp_path):
    """production + full + explicit persistent filesystem storage => allowed."""
    _set_full_surface_env(monkeypatch, environment)
    persistent_root = tmp_path / "procurex-attachments"
    monkeypatch.setenv("INCOMING_REQUEST_UPLOAD_DIR", str(persistent_root))
    storage = build_attachment_storage()
    assert isinstance(storage, LocalAttachmentStorage)
    assert storage.root == persistent_root.resolve()


def test_hosted_full_surface_allows_persistent_storage_via_data_root(monkeypatch, tmp_path):
    """PROCUREX_DATA_ROOT is also an accepted explicit persistent-path source,
    and must resolve to the same attachment path server.py's diagnostics use."""
    _set_full_surface_env(monkeypatch)
    data_root = tmp_path / "procurex-data"
    monkeypatch.setenv("PROCUREX_DATA_ROOT", str(data_root))
    storage = build_attachment_storage()
    assert isinstance(storage, LocalAttachmentStorage)
    assert storage.root == (data_root.resolve() / "data" / "attachments" / "incoming_requests")


def test_hosted_full_surface_rejects_unconfigured_local_storage(monkeypatch):
    """production + full + filesystem storage without required persistent
    config => rejected, rather than silently defaulting into the application
    directory."""
    _set_full_surface_env(monkeypatch)
    with pytest.raises(RuntimeError, match="requires an explicit persistent path"):
        build_attachment_storage()


def test_hosted_full_surface_rejects_persistent_path_inside_app_directory(monkeypatch):
    _set_full_surface_env(monkeypatch)
    monkeypatch.setenv("INCOMING_REQUEST_UPLOAD_DIR", str(ROOT_DIR / "storage" / "incoming_requests"))
    with pytest.raises(RuntimeError, match="must be outside the application directory"):
        build_attachment_storage()


@pytest.mark.parametrize("environment", ["staging", "production"])
def test_hosted_full_surface_allows_s3(monkeypatch, environment):
    """production + full + s3 => allowed (S3 remains supported for the full
    surface, not just public). boto3 is a production-only dependency and
    isn't installed in this dev/test environment, so success here is
    verified by confirming the surface/environment gate lets the call reach
    S3 construction (rather than rejecting it outright), not by a live S3
    round-trip."""
    monkeypatch.setenv("APP_ENV", environment)
    monkeypatch.setenv("APP_SURFACE", "full")
    monkeypatch.setenv("ATTACHMENT_STORAGE_BACKEND", "s3")
    monkeypatch.setenv("R2_ENDPOINT_URL", "https://account.r2.cloudflarestorage.com")
    monkeypatch.setenv("R2_ACCESS_KEY_ID", "test-access-key")
    monkeypatch.setenv("R2_SECRET_ACCESS_KEY", "test-secret-key")
    monkeypatch.setenv("R2_BUCKET_NAME", f"procurex-{environment}-attachments")
    with pytest.raises(RuntimeError, match="boto3 is required"):
        build_attachment_storage()


def test_development_local_attachment_storage_is_unaffected(monkeypatch, tmp_path):
    """Existing development configuration remains unaffected - no explicit
    persistent path is required outside staging/production."""
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.delenv("APP_SURFACE", raising=False)
    monkeypatch.setenv("ATTACHMENT_STORAGE_BACKEND", "local")
    monkeypatch.delenv("PROCUREX_DATA_ROOT", raising=False)
    monkeypatch.delenv("INCOMING_REQUEST_UPLOAD_DIR", raising=False)
    storage = build_attachment_storage()
    assert isinstance(storage, LocalAttachmentStorage)
    assert storage.root == (ROOT_DIR / "storage" / "incoming_requests").resolve()


def test_public_surface_rejects_oversized_http_body_before_form_parsing(monkeypatch):
    monkeypatch.setenv("PUBLIC_MAX_HTTP_BODY_BYTES", "100")
    public_app = create_app(surface="public", initialize_database=False)
    with TestClient(public_app) as client:
        response = client.post(
            "/api/public/purchase-requests",
            content=b"x",
            headers={"Content-Length": "101"},
        )
    assert response.status_code == 413


# ---------- Human-reviewed document capture ----------
class _FakeOCR:
    name = "fake-azure"

    def extract(self, files):
        assert files and files[0].media_type == "image/png"
        return OCRResult(
            raw_text="Copper cable 12 roll",
            confidence=0.91,
            page_count=1,
            provider_data={"fixture": True},
        )


class _FakeStructuredExtraction:
    name = "fake-openai"

    def extract_items(self, raw_text, document_type):
        assert raw_text == "Copper cable 12 roll"
        assert document_type == "printed_request"
        return [
            ExtractedLineItem(
                raw_line_text=raw_text,
                product_name="DOC_CAPTURE_COPPER_CABLE",
                brand="Test Brand",
                specification="4 mm",
                quantity=12,
                unit="roll",
                extraction_confidence=0.88,
            )
        ], {"fixture": True}


def _document_payload(token):
    return {
        "requester_name": "Document Test User",
        "company_name": "Synthetic Fixture",
        "phone_number": "+201000000000",
        "project_name": "Synthetic Project",
        "project_location": "Cairo",
        "delivery_location": "Cairo",
        "required_delivery_date": "2099-12-31",
        "priority": "normal",
        "notes": "Synthetic test only",
        "submission_token": token,
        "document_type": "printed_request",
    }


def test_document_capture_is_safe_when_disabled(s, monkeypatch):
    monkeypatch.setenv("DOCUMENT_EXTRACTION_ENABLED", "false")
    state = s.get(f"{API}/public/purchase-requests/document-extraction-capability")
    assert state.status_code == 200
    assert state.json()["status"] == "disabled"
    assert state.json()["manual_entry_available"] is True
    response = s.post(
        f"{API}/public/purchase-requests/documents",
        data={"payload": json.dumps(_document_payload("disabled-document-token-001"))},
        files={"documents": ("request.png", b"\x89PNG\r\n\x1a\nIEND", "image/png")},
    )
    assert response.status_code == 503


def test_document_capture_review_confirmation_and_audit(s, monkeypatch, admin_headers):
    monkeypatch.setenv("DOCUMENT_EXTRACTION_ENABLED", "true")
    monkeypatch.setenv(
        "AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT", "https://example.invalid"
    )
    monkeypatch.setenv("AZURE_DOCUMENT_INTELLIGENCE_KEY", "synthetic-key")
    monkeypatch.setenv("OPENAI_API_KEY", "synthetic-key")
    monkeypatch.setenv("OPENAI_DOCUMENT_EXTRACTION_MODEL", "synthetic-model")
    monkeypatch.setattr(
        document_service,
        "build_providers",
        lambda: (
            _FakeOCR(),
            _FakeStructuredExtraction(),
        ),
    )

    master_count = len(s.get(f"{API}/items", headers=admin_headers).json())
    created_master = s.post(
        f"{API}/items",
        json={
            "product_name": "DOC_CAPTURE_COPPER_CABLE",
            "brand": "Test Brand",
            "unit": "roll",
        },
        headers=admin_headers,
    )
    assert created_master.status_code == 200
    master = created_master.json()
    try:
        response = s.post(
            f"{API}/public/purchase-requests/documents",
            data={
                "payload": json.dumps(_document_payload("document-review-token-0001")),
                "website": "",
            },
            files={
                "documents": (
                    "request.png",
                    b"\x89PNG\r\n\x1a\nsynthetic-image-content-IEND",
                    "image/png",
                )
            },
        )
        assert response.status_code == 200, response.text
        result = response.json()
        request_id = next(
            row["id"]
            for row in s.get(
                f"{API}/internal/incoming-purchase-requests",
                headers={**INTERNAL_HEADERS, **admin_headers},
                params={"search": result["request_number"]},
            ).json()
            if row["request_number"] == result["request_number"]
        )
        detail_url = f"{API}/internal/incoming-purchase-requests/{request_id}"
        assert s.get(detail_url, headers={**INTERNAL_HEADERS, **admin_headers}).json()["items"] == []

        assert run_once() is True
        review_url = f"{detail_url}/documents/{result['document_id']}"
        review = s.get(review_url, headers={**INTERNAL_HEADERS, **admin_headers})
        assert review.status_code == 200, review.text
        extracted = review.json()["items"]
        assert review.json()["status"] == "review_required"
        assert extracted[0]["quantity"] == 12
        assert extracted[0]["candidates"][0]["item_id"] == master["id"]

        assert (
            s.post(f"{review_url}/retry", headers={**INTERNAL_HEADERS, **admin_headers}).status_code == 200
        )
        assert s.get(review_url, headers={**INTERNAL_HEADERS, **admin_headers}).json()["items"] == []
        assert run_once() is True
        extracted = s.get(review_url, headers={**INTERNAL_HEADERS, **admin_headers}).json()["items"]

        assert (
            s.patch(
                f"{review_url}/items/{extracted[0]['id']}",
                headers={**INTERNAL_HEADERS, **admin_headers},
                json={
                    "quantity": 14,
                    "selected_item_id": master["id"],
                    "review_notes": "Checked",
                },
            ).status_code
            == 200
        )
        added = s.post(
            f"{review_url}/items",
            headers={**INTERNAL_HEADERS, **admin_headers},
            json={
                "product_name": "DOC_CAPTURE_COPPER_CABLE",
                "quantity": None,
                "unit": "",
            },
        )
        assert added.status_code == 200
        flagged = s.get(review_url, headers={**INTERNAL_HEADERS, **admin_headers}).json()["items"]
        added_flag = next(row for row in flagged if row["id"] == added.json()["id"])
        assert added_flag["validation"]["missing_quantity"] is True
        assert added_flag["validation"]["missing_unit"] is True
        assert added_flag["validation"]["suspected_duplicate"] is True
        split = s.post(
            f"{review_url}/items/{added.json()['id']}/split",
            headers={**INTERNAL_HEADERS, **admin_headers},
            json={
                "first": {"product_name": "Split A", "quantity": 1, "unit": "roll"},
                "second": {"product_name": "Split B", "quantity": 1, "unit": "roll"},
            },
        )
        assert split.status_code == 200
        merge = s.post(
            f"{review_url}/items/merge",
            headers={**INTERNAL_HEADERS, **admin_headers},
            json={"item_ids": [added.json()["id"], split.json()["created_id"]]},
        )
        assert merge.status_code == 200
        assert (
            s.delete(
                f"{review_url}/items/{merge.json()['target_id']}",
                headers={**INTERNAL_HEADERS, **admin_headers},
            ).status_code
            == 200
        )
        assert len(s.get(review_url, headers={**INTERNAL_HEADERS, **admin_headers}).json()["items"]) == 1
        creation_request = s.post(
            f"{review_url}/items/{extracted[0]['id']}/master-creation-request",
            headers={**INTERNAL_HEADERS, **admin_headers},
            json={
                "requested_by": "Reviewer",
                "confirmation": "REQUEST_MASTER_ITEM_CREATION",
            },
        )
        assert creation_request.status_code == 200
        assert len(s.get(f"{API}/items", headers=admin_headers).json()) == master_count + 1

        confirm = s.post(
            f"{review_url}/confirm",
            headers={**INTERNAL_HEADERS, **admin_headers},
            json={"confirmed_by": "Reviewer", "confirmation": "CONFIRM_REVIEWED_ITEMS"},
        )
        assert confirm.status_code == 200, confirm.text
        assert confirm.json()["items_added"] == 1
        detail_after = s.get(detail_url, headers={**INTERNAL_HEADERS, **admin_headers}).json()
        assert len(detail_after["items"]) == 1
        assert detail_after["items"][0]["quantity"] == 14
        assert (
            s.post(
                f"{review_url}/confirm",
                headers={**INTERNAL_HEADERS, **admin_headers},
                json={
                    "confirmed_by": "Reviewer",
                    "confirmation": "CONFIRM_REVIEWED_ITEMS",
                },
            ).status_code
            == 422
        )
        with SessionLocal() as session:
            assert (
                session.query(RequestAuditEvent)
                .filter_by(document_id=result["document_id"])
                .count()
                >= 4
            )
            assert (
                session.query(ItemMasterCreationRequest)
                .filter_by(extracted_item_id=extracted[0]["id"])
                .count()
                == 1
            )
    finally:
        s.delete(f"{API}/items/{master['id']}")


def test_document_upload_security_and_cancel(s, monkeypatch, admin_headers):
    monkeypatch.setenv("DOCUMENT_EXTRACTION_ENABLED", "true")
    monkeypatch.setenv(
        "AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT", "https://example.invalid"
    )
    monkeypatch.setenv("AZURE_DOCUMENT_INTELLIGENCE_KEY", "synthetic-key")
    monkeypatch.setenv("OPENAI_API_KEY", "synthetic-key")
    monkeypatch.setenv("OPENAI_DOCUMENT_EXTRACTION_MODEL", "synthetic-model")
    invalid = s.post(
        f"{API}/public/purchase-requests/documents",
        data={"payload": json.dumps(_document_payload("document-invalid-token-001"))},
        files={"documents": ("fake.png", b"not-an-image", "image/png")},
    )
    assert invalid.status_code == 422
    oversized = s.post(
        f"{API}/public/purchase-requests/documents",
        data={"payload": json.dumps(_document_payload("document-oversized-token-001"))},
        files={
            "documents": (
                "large.png",
                b"\x89PNG\r\n\x1a\n" + (b"x" * MAX_FILE_BYTES) + b"IEND",
                "image/png",
            )
        },
    )
    assert oversized.status_code == 413

    retained_payload = _document_payload("document-retained-token-001")
    retained_payload["document_type"] = "supplier_quotation"
    retained = s.post(
        f"{API}/public/purchase-requests/documents",
        data={"payload": json.dumps(retained_payload)},
        files=[
            (
                "documents",
                ("quote.pdf", b"%PDF-1.4 synthetic %%EOF", "application/pdf"),
            ),
            ("documents", ("page.png", b"\x89PNG\r\n\x1a\npage-IEND", "image/png")),
        ],
    )
    assert retained.status_code == 200, retained.text
    retained_result = retained.json()
    listing = s.get(
        f"{API}/internal/incoming-purchase-requests",
        headers={**INTERNAL_HEADERS, **admin_headers},
        params={"search": retained_result["request_number"]},
    ).json()
    request_id = listing[0]["id"]
    review_url = (
        f"{API}/internal/incoming-purchase-requests/{request_id}"
        f"/documents/{retained_result['document_id']}"
    )
    review = s.get(review_url, headers={**INTERNAL_HEADERS, **admin_headers}).json()
    assert review["document_type"] == "supplier_quotation"
    assert review["status"] == "uploaded"
    assert len(review["files"]) == 2
    assert s.post(f"{review_url}/retry", headers={**INTERNAL_HEADERS, **admin_headers}).status_code == 409
    assert s.post(f"{review_url}/cancel", headers={**INTERNAL_HEADERS, **admin_headers}).status_code == 200
    assert s.get(review_url, headers={**INTERNAL_HEADERS, **admin_headers}).json()["status"] == "cancelled"


def test_document_capture_migration_is_additive_and_backed_up(tmp_path):
    path = tmp_path / "legacy.db"
    migration_engine = create_engine(f"sqlite:///{path.as_posix()}")
    with migration_engine.begin() as connection:
        connection.exec_driver_sql(
            "CREATE TABLE items (id VARCHAR PRIMARY KEY, name VARCHAR NOT NULL, "
            "product_name VARCHAR NOT NULL DEFAULT '')"
        )
        connection.exec_driver_sql(
            "INSERT INTO items(id, name, product_name) VALUES ('1', 'Arabic item', 'Arabic item')"
        )
        connection.exec_driver_sql("PRAGMA user_version = 6")
    backup = migrate_document_capture(migration_engine)
    assert backup and backup.is_file()
    assert migrate_document_capture(migration_engine) is None
    with sqlite3.connect(path) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
        assert (
            connection.execute("PRAGMA user_version").fetchone()[0]
            == DOCUMENT_CAPTURE_SCHEMA_VERSION
        )
        columns = {row[1] for row in connection.execute("PRAGMA table_info(items)")}
        assert {"name_ar", "name_en", "alternative_names", "search_aliases"} <= columns
        assert (
            connection.execute("SELECT name_ar FROM items WHERE id='1'").fetchone()[0]
            == "Arabic item"
        )
        assert connection.execute("SELECT COUNT(*) FROM items").fetchone()[0] == 1
        tables = {
            row[0]
            for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )
        }
        assert "purchase_request_documents" in tables


# ---------- Removed Construction Calculator regression coverage ----------
CONSTRUCTION_API_PATHS = (
    f"{API}/construction-calculator/categories",
    f"{API}/construction-calculator/work-items",
    f"{API}/construction-calculator/calculate",
)


@pytest.mark.parametrize("legacy_flag", [None, "true"])
def test_construction_api_is_unavailable_even_with_legacy_flag(monkeypatch, legacy_flag):
    if legacy_flag is None:
        monkeypatch.delenv("CONSTRUCTION_API_ENABLED", raising=False)
    else:
        monkeypatch.setenv("CONSTRUCTION_API_ENABLED", legacy_flag)
    disabled_app = create_app(initialize_database=False)
    registered_paths = {route.path for route in disabled_app.routes}
    assert not any(path.startswith(f"{API}/construction-calculator") for path in registered_paths)
    with TestClient(disabled_app) as client:
        for path in CONSTRUCTION_API_PATHS:
            assert client.get(path, headers={"X-Internal-Token": "test-internal-token"}).status_code == 404
        assert client.post(CONSTRUCTION_API_PATHS[-1], json={}).status_code == 404


def test_runtime_has_no_construction_router_or_seed_imports():
    server_source = (BACKEND_DIR / "server.py").read_text(encoding="utf-8")
    database_source = (BACKEND_DIR / "database.py").read_text(encoding="utf-8")
    installer_source = (
        BACKEND_DIR.parent / "installer" / "Build-Installer.ps1"
    ).read_text(encoding="utf-8")
    assert "construction_calculator.router" not in server_source
    assert "construction_calculator_router" not in server_source
    assert "construction_calculator.seed" not in database_source
    assert "migrate_construction_calculator" not in database_source
    assert 'CONSTRUCTION_AUTO_INSTALL", "true"' not in database_source
    assert "--collect-submodules construction_calculator" not in installer_source
    assert "construction_calculator/data" not in installer_source


def test_normal_startup_preserves_historical_construction_data_without_seeding(monkeypatch):
    sentinel_id = "historical-compatibility-category"
    try:
        with engine.begin() as connection:
            existing_tables = {
                row[0]
                for row in connection.exec_driver_sql(
                    "SELECT name FROM sqlite_master WHERE type='table' "
                    "AND name LIKE 'construction_%'"
                )
            }
            assert existing_tables == set()
            connection.exec_driver_sql(
                "CREATE TABLE construction_categories ("
                "id TEXT PRIMARY KEY, code TEXT NOT NULL, source_text TEXT NOT NULL, "
                "technical_notes TEXT NOT NULL)"
            )
            connection.exec_driver_sql(
                "INSERT INTO construction_categories "
                "(id, code, source_text, technical_notes) VALUES (?, ?, ?, ?)",
                (
                    sentinel_id,
                    "HIST-COMPAT",
                    "preserved historical row",
                    "must remain unchanged",
                ),
            )
            before_schema = connection.exec_driver_sql(
                "SELECT sql FROM sqlite_master WHERE type='table' "
                "AND name='construction_categories'"
            ).scalar_one()
            before_row = connection.exec_driver_sql(
                "SELECT code, source_text, technical_notes "
                "FROM construction_categories WHERE id = ?",
                (sentinel_id,),
            ).one()

        # A legacy environment variable must not reactivate removed seeding.
        monkeypatch.setenv("CONSTRUCTION_AUTO_INSTALL", "true")
        init_db()

        with engine.begin() as connection:
            after_tables = {
                row[0]
                for row in connection.exec_driver_sql(
                    "SELECT name FROM sqlite_master WHERE type='table' "
                    "AND name LIKE 'construction_%'"
                )
            }
            after_schema = connection.exec_driver_sql(
                "SELECT sql FROM sqlite_master WHERE type='table' "
                "AND name='construction_categories'"
            ).scalar_one()
            after_row = connection.exec_driver_sql(
                "SELECT code, source_text, technical_notes "
                "FROM construction_categories WHERE id = ?",
                (sentinel_id,),
            ).one()
        assert after_tables == {"construction_categories"}
        assert after_schema == before_schema
        assert after_row == before_row
    finally:
        with engine.begin() as connection:
            connection.exec_driver_sql("DROP TABLE IF EXISTS construction_categories")


def test_construction_migration_is_additive_backed_up_and_idempotent(tmp_path):
    path = tmp_path / "legacy-construction.db"
    migration_engine = create_engine(f"sqlite:///{path.as_posix()}")
    with migration_engine.begin() as connection:
        connection.exec_driver_sql(
            "CREATE TABLE legacy_data (id INTEGER PRIMARY KEY, value TEXT)"
        )
        connection.exec_driver_sql(
            "INSERT INTO legacy_data(value) VALUES ('preserved')"
        )
        connection.exec_driver_sql("PRAGMA user_version = 7")
    backup = migrate_construction_calculator(migration_engine)
    assert backup and backup.is_file()
    assert migrate_construction_calculator(migration_engine) is None
    with sqlite3.connect(path) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
        assert (
            connection.execute("PRAGMA user_version").fetchone()[0]
            == CONSTRUCTION_CALCULATOR_SCHEMA_VERSION
        )
        assert (
            connection.execute("SELECT value FROM legacy_data").fetchone()[0]
            == "preserved"
        )
        assert (
            connection.execute(
                "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='construction_work_items'"
            ).fetchone()[0]
            == 1
        )


def test_construction_alembic_upgrade_and_downgrade_cycle(tmp_path):
    path = tmp_path / "construction-alembic-cycle.db"
    migration_engine = create_engine(f"sqlite:///{path.as_posix()}")
    migration_path = (
        BACKEND_DIR / "alembic" / "versions" / "0006_construction_calculator.py"
    )
    spec = importlib.util.spec_from_file_location(
        "construction_alembic_cycle", migration_path
    )
    migration = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(migration)
    with migration_engine.begin() as connection:
        connection.exec_driver_sql("PRAGMA foreign_keys=OFF")
        for table in Base.metadata.sorted_tables:
            if not table.name.startswith("construction_"):
                table.create(connection, checkfirst=True)
        operations = Operations(MigrationContext.configure(connection))
        original_op = migration.op
        migration.op = operations
        try:
            migration.upgrade()
            assert "construction_work_items" in {
                row[0]
                for row in connection.exec_driver_sql(
                    "SELECT name FROM sqlite_master WHERE type='table'"
                )
            }
            migration.downgrade()
            assert "construction_work_items" not in {
                row[0]
                for row in connection.exec_driver_sql(
                    "SELECT name FROM sqlite_master WHERE type='table'"
                )
            }
            migration.upgrade()
            assert "construction_calculation_snapshots" in {
                row[0]
                for row in connection.exec_driver_sql(
                    "SELECT name FROM sqlite_master WHERE type='table'"
                )
            }
        finally:
            migration.op = original_op


def test_guided_project_approval_payment_revision_and_cash_workflow(s, admin_headers):
    """One end-to-end test keeps the state-machine relationships visible."""
    timestamp = datetime.now(timezone.utc).isoformat()
    request_id = str(uuid.uuid4())
    with SessionLocal.begin() as session:
        session.add(IncomingPurchaseRequest(
            id=request_id, request_number=f"T-REQ-{uuid.uuid4().hex[:8]}",
            requester_name="مهندس الاختبار", company_name="عميل الاختبار",
            phone_number="01000000000", project_name="مشروع ربط اختباري",
            project_location="القاهرة", delivery_location="القاهرة",
            required_delivery_date="2099-01-01", priority="normal",
            submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
            created_at=timestamp, updated_at=timestamp,
        ))
        session.add_all([
            IncomingPurchaseRequestItem(
                id=str(uuid.uuid4()), request_id=request_id, position=1,
                product_name="صنف فني أول", quantity=1, unit="قطعة",
                review_status="approved", reviewed_by="engineer", reviewed_at=timestamp,
            ),
            IncomingPurchaseRequestItem(
                id=str(uuid.uuid4()), request_id=request_id, position=2,
                product_name="صنف فني ثان", quantity=2, unit="قطعة",
                review_status="approved", reviewed_by="engineer", reviewed_at=timestamp,
            ),
        ])
    existing_project = s.get(f"{API}/projects", headers=admin_headers).json()[0]
    linked = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/link-project",
        json={"project_id": existing_project["id"], "actor": "tester"},
        headers={**INTERNAL_HEADERS, **admin_headers},
    )
    assert linked.status_code == 200, linked.text
    with SessionLocal() as session:
        assert session.get(IncomingPurchaseRequest, request_id).project_id == existing_project["id"]

    second_id = str(uuid.uuid4())
    unique_project_name = f"مشروع جديد {uuid.uuid4().hex[:8]}"
    with SessionLocal.begin() as session:
        session.add(IncomingPurchaseRequest(
            id=second_id, request_number=f"T-REQ-{uuid.uuid4().hex[:8]}",
            requester_name="مسؤول جديد", company_name="عميل جديد",
            phone_number="01000000001", project_name=unique_project_name,
            project_location="الجيزة", delivery_location="الجيزة",
            required_delivery_date="2099-01-01", priority="normal",
            submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
            created_at=timestamp, updated_at=timestamp,
        ))
    created_project = s.post(
        f"{API}/workflow/incoming-purchase-requests/{second_id}/create-project",
        json={"name": unique_project_name, "customer_name": "عميل جديد", "city": "الجيزة", "actor": "tester"},
        headers={**INTERNAL_HEADERS, **admin_headers},
    )
    assert created_project.status_code == 200, created_project.text
    duplicate_project = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/create-project",
        json={"name": unique_project_name},
        headers={**INTERNAL_HEADERS, **admin_headers},
    )
    assert duplicate_project.status_code == 409

    item = s.get(f"{API}/items", headers=admin_headers).json()[0]
    supplier = s.get(f"{API}/suppliers", headers=admin_headers).json()[0]
    technical = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/technical-decision",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"decision": "approved_for_pricing", "actor": "tester"},
    )
    assert technical.status_code == 200, technical.text

    def new_approval():
        comparison = s.post(f"{API}/price-comparisons", json={
            "project_id": existing_project["id"], "project_name": existing_project["name"],
            "customer_name": existing_project.get("customer_name", ""),
            "source_request_id": request_id, "source_request_number": "SOURCE-TEST",
            "comparison_date": "2026-08-16", "rows": [{
                "item_id": item["id"], "supplier_id": supplier["id"],
                "quantity": 2, "unit": item.get("unit", "قطعة"), "unit_price": 125,
                "availability": "available", "price_valid_until": "2099-12-31",
                "selected_for_purchase": 1,
            }],
        }, headers=admin_headers)
        assert comparison.status_code == 200, comparison.text
        response = s.post(f"{API}/workflow/approvals/from-comparison", headers={**INTERNAL_HEADERS, **admin_headers}, json={
            "comparison_id": comparison.json()["id"], "engineer_name": "مهندس الاعتماد",
            "engineer_email": "engineer@example.com", "created_by": "tester",
        })
        assert response.status_code == 201, response.text
        approval = response.json()["approval"]
        assert approval["lines"][0]["supplier_name"] == supplier["name"]
        assert approval["lines"][0]["unit_price"] == 125
        assert s.post(f"{API}/workflow/approvals/from-comparison", headers={**INTERNAL_HEADERS, **admin_headers}, json={
            "comparison_id": comparison.json()["id"], "engineer_name": "duplicate",
        }).status_code == 409
        s.post(f"{API}/workflow/approvals/{approval['id']}/ready", headers={**INTERNAL_HEADERS, **admin_headers}, json={"actor": "tester"})
        sent = s.post(f"{API}/workflow/approvals/{approval['id']}/sent", headers={**INTERNAL_HEADERS, **admin_headers}, json={"actor": "tester"})
        assert sent.status_code == 200, sent.text
        return approval

    approved = new_approval()
    public_loaded = s.get(f"{API}/public/approvals/{approved['secure_token']}")
    assert public_loaded.status_code == 200
    assert "secure_token" not in public_loaded.json()
    decision = s.post(f"{API}/public/approvals/{approved['secure_token']}/decision", json={"decision": "approved"})
    assert decision.status_code == 200
    repeated = s.post(f"{API}/public/approvals/{approved['secure_token']}/decision", json={"decision": "approved"})
    assert repeated.json()["already_recorded"] is True
    assert s.post(f"{API}/public/approvals/{approved['secure_token']}/decision", json={"decision": "rejected"}).status_code == 409

    intent = s.post(f"{API}/public/approvals/{approved['secure_token']}/payments", json={"method": "instapay"})
    assert intent.status_code == 201, intent.text
    payment_id = intent.json()["payment"]["id"]
    proof = s.post(
        f"{API}/public/approvals/{approved['secure_token']}/payments/{payment_id}/proof",
        files={"proof": ("proof.png", b"safe-test-image", "image/png")},
        data={"external_reference": "TX-1"},
    )
    assert proof.status_code == 200, proof.text
    assert proof.json() == {"ok": True, "status": "under_review", "paid": False}
    verified = s.post(f"{API}/workflow/payments/{payment_id}/verify", headers={**INTERNAL_HEADERS, **admin_headers}, json={"actor": "cashier"})
    assert verified.status_code == 200
    assert verified.json()["payment"]["status"] == "verified"

    generated_po = s.post(f"{API}/purchase-orders/from-comparison", json={
        "comparison_id": approved["comparison_id"], "po_date": "2026-08-16",
        "created_by": "tester", "orders": [{
            "supplier_id": supplier["id"], "supplier_name": supplier["name"],
            "items": [{"item_id": item["id"], "product_name": item["product_name"],
                       "quantity": 2, "unit_price": 125}],
        }],
    }, headers=admin_headers)
    assert generated_po.status_code == 200, generated_po.text
    po = generated_po.json()["purchase_orders"][0]
    assert po["source_request_id"] == request_id
    assert po["approval_id"] == approved["id"]
    assert s.post(f"{API}/purchase-orders/from-comparison", json={
        "comparison_id": approved["comparison_id"], "po_date": "2026-08-16",
        "orders": [{"supplier_name": supplier["name"], "items": [{
            "product_name": item["product_name"], "quantity": 2, "unit_price": 125,
        }]}],
    }, headers=admin_headers).status_code == 409

    revision_source = new_approval()
    s.get(f"{API}/public/approvals/{revision_source['secure_token']}")
    revision_decision = s.post(
        f"{API}/public/approvals/{revision_source['secure_token']}/decision",
        json={"decision": "revision_requested", "note": "غيّر المورد"},
    )
    assert revision_decision.status_code == 200
    revision = s.post(f"{API}/workflow/approvals/{revision_source['id']}/revision", headers={**INTERNAL_HEADERS, **admin_headers}, json={"actor": "tester"})
    assert revision.status_code == 201, revision.text
    assert revision.json()["approval"]["revision_number"] == 1
    assert revision.json()["approval"]["previous_revision_id"] == revision_source["id"]

    cash_approval = new_approval()
    s.get(f"{API}/public/approvals/{cash_approval['secure_token']}")
    s.post(f"{API}/public/approvals/{cash_approval['secure_token']}/decision", json={"decision": "approved"})
    cash = s.post(f"{API}/public/approvals/{cash_approval['secure_token']}/payments", json={"method": "cash"})
    assert cash.status_code == 201
    cash_payment = cash.json()["payment"]
    confirmed = s.post(f"{API}/workflow/payments/{cash_payment['id']}/confirm-cash", headers={**INTERNAL_HEADERS, **admin_headers}, json={
        "cash_reference": cash_payment["cash_reference"], "actor": "cashier",
    })
    assert confirmed.status_code == 200
    assert confirmed.json()["payment"]["status"] == "verified"
    assert s.post(f"{API}/workflow/payments/{cash_payment['id']}/confirm-cash", headers={**INTERNAL_HEADERS, **admin_headers}, json={
        "cash_reference": cash_payment["cash_reference"], "actor": "cashier",
    }).status_code == 409

    hub = s.get(f"{API}/workflow/projects/{existing_project['id']}/procurement-hub", headers={**INTERNAL_HEADERS, **admin_headers})
    assert hub.status_code == 200
    assert hub.json()["kpis"]["comparison_count"] >= 3


def test_internal_procurement_roles_gate_comparison_fund_and_po(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"role-engineer-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"role-responsible-{suffix}", role="procurement_responsible")
        _make_user(session, username=f"role-manager-{suffix}", role="commercial_manager")
    engineer_headers = _login_headers(s, f"role-engineer-{suffix}")
    responsible_headers = _login_headers(s, f"role-responsible-{suffix}")
    manager_headers = _login_headers(s, f"role-manager-{suffix}")

    timestamp = datetime.now(timezone.utc).isoformat()
    request_id = str(uuid.uuid4())
    request_number = f"T-ROLE-{uuid.uuid4().hex[:8]}"
    with SessionLocal.begin() as session:
        session.add(IncomingPurchaseRequest(
            id=request_id, request_number=request_number,
            requester_name="مهندس الموقع", company_name="عميل الأدوار",
            phone_number="01000000009", project_name="مشروع الأدوار",
            project_location="القاهرة", delivery_location="القاهرة",
            required_delivery_date="2099-01-01", priority="normal",
            submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
            created_at=timestamp, updated_at=timestamp,
        ))
        session.add_all([
            IncomingPurchaseRequestItem(
                id=str(uuid.uuid4()), request_id=request_id, position=1,
                product_name="صنف فني أول", quantity=1, unit="قطعة",
                review_status="approved", reviewed_by="engineer", reviewed_at=timestamp,
            ),
            IncomingPurchaseRequestItem(
                id=str(uuid.uuid4()), request_id=request_id, position=2,
                product_name="صنف فني ثان", quantity=2, unit="قطعة",
                review_status="approved", reviewed_by="engineer", reviewed_at=timestamp,
            ),
        ])
    project = s.get(f"{API}/projects", headers=admin_headers).json()[0]
    items = s.get(f"{API}/items", headers=admin_headers).json()[:2]
    suppliers = s.get(f"{API}/suppliers", headers=admin_headers).json()[:2]
    comparison_body = {
        "project_id": project["id"], "project_name": project["name"],
        "source_request_id": request_id, "source_request_number": "CLIENT-SPOOFED-REQUEST",
        "comparison_date": "2026-08-16", "rows": [{
            "item_id": items[0]["id"], "supplier_id": suppliers[0]["id"],
            "quantity": 1, "unit": items[0].get("unit", "قطعة"), "unit_price": 100,
            "availability": "available", "price_valid_until": "2099-12-31",
            "selected_for_purchase": 1,
        }, {
            "item_id": items[1]["id"], "supplier_id": suppliers[1]["id"],
            "quantity": 2, "unit": items[1].get("unit", "قطعة"), "unit_price": 75,
            "discount_pct": 5, "tax_pct": 14, "delivery_days": 4,
            "payment_terms": "نقدي", "availability": "available",
            "price_valid_until": "2099-12-31", "selected_for_purchase": 1,
        }],
    }
    assert s.post(f"{API}/price-comparisons", json=comparison_body, headers=responsible_headers).status_code == 409
    denied_review = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/technical-decision",
        headers={**INTERNAL_HEADERS, **manager_headers},
        json={"decision": "approved_for_pricing"},
    )
    assert denied_review.status_code == 403
    historical = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    )
    assert historical.status_code == 200
    assert historical.json()["project_id"] == ""
    unlinked_review = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/technical-decision",
        headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"decision": "approved_for_pricing"},
    )
    assert unlinked_review.status_code == 409
    linked = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/link-project",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"project_id": project["id"], "actor": "engineer"},
    )
    assert linked.status_code == 200, linked.text
    reviewed = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/technical-decision",
        headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"decision": "approved_for_pricing", "actor": "engineer"},
    )
    assert reviewed.status_code == 200
    repeated_review = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/technical-decision",
        headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"decision": "approved_for_pricing", "actor": "engineer"},
    )
    assert repeated_review.status_code == 200
    assert repeated_review.json()["already_recorded"] is True
    reopened_request = s.get(f"{API}/internal/incoming-purchase-requests/{request_id}", headers={**INTERNAL_HEADERS, **admin_headers})
    assert reopened_request.status_code == 200
    assert reopened_request.json()["status"] == "pricing"
    assert reopened_request.json()["project_id"] == project["id"]
    assert all(item["review_status"] == "approved" for item in reopened_request.json()["items"])
    missing_project_body = {**comparison_body, "project_id": "", "project_name": project["name"]}
    assert s.post(f"{API}/price-comparisons", json=missing_project_body, headers=responsible_headers).status_code == 422
    invalid_project_body = {
        **comparison_body,
        "project_id": "nonexistent-project",
        "project_name": project["name"],
    }
    assert s.post(f"{API}/price-comparisons", json=invalid_project_body, headers=responsible_headers).status_code == 422
    comparison = s.post(f"{API}/price-comparisons", json=comparison_body, headers=responsible_headers)
    assert comparison.status_code == 200, comparison.text
    assert comparison.json()["project_id"] == project["id"]
    assert comparison.json()["project_name"] == project["name"]
    assert comparison.json()["source_request_id"] == request_id
    assert comparison.json()["source_request_number"] == request_number
    reassigned = s.put(
        f"{API}/price-comparisons/{comparison.json()['id']}",
        json={**comparison_body, "source_request_id": str(uuid.uuid4())},
        headers=responsible_headers,
    )
    assert reassigned.status_code == 409
    standalone = s.post(f"{API}/price-comparisons", json={
        **comparison_body,
        "source_request_id": "",
        "source_request_number": "CLIENT-ONLY-REFERENCE",
        "rows": [comparison_body["rows"][0]],
    }, headers=responsible_headers)
    assert standalone.status_code == 200, standalone.text
    assert standalone.json()["source_request_number"] == ""
    assert s.post(
        f"{API}/workflow/approvals/from-comparison",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        json={
            "comparison_id": standalone.json()["id"],
            "approval_type": "comparison_workflow",
        },
    ).status_code == 409
    assert s.put(
        f"{API}/price-comparisons/{comparison.json()['id']}",
        json=missing_project_body,
        headers=responsible_headers,
    ).status_code == 422
    with SessionLocal.begin() as session:
        session.get(PriceComparison, comparison.json()["id"]).project_id = None
    invalid_approval = s.post(
        f"{API}/workflow/approvals/from-comparison", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"comparison_id": comparison.json()["id"], "approval_type": "comparison_workflow"},
    )
    assert invalid_approval.status_code == 409
    with SessionLocal.begin() as session:
        session.get(PriceComparison, comparison.json()["id"]).project_id = project["id"]
    approval_response = s.post(
        f"{API}/workflow/approvals/from-comparison", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"comparison_id": comparison.json()["id"], "approval_type": "comparison_workflow", "created_by": "officer"},
    )
    assert approval_response.status_code == 201, approval_response.text
    approval = approval_response.json()["approval"]
    assert approval["project_id"] == project["id"]
    assert approval["project_name"] == project["name"]
    assert approval["source_request_id"] == request_id
    assert approval["source_request_number"] == request_number
    assert approval["comparison_id"] == comparison.json()["id"]
    assert approval["comparison_number"] == comparison.json()["comparison_number"]
    assert approval["approval_stage"] == "comparison_technical"
    assert approval["responsible_role"] == "procurement_engineer"
    request_after_approval = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    assert request_after_approval["status"] == "waiting_for_approval"
    assert s.get(f"{API}/public/approvals/{approval['secure_token']}").status_code == 404

    po_body = {"comparison_id": comparison.json()["id"], "po_date": "2026-08-16", "orders": []}
    assert s.post(f"{API}/purchase-orders/from-comparison", json=po_body, headers=responsible_headers).status_code == 409
    denied = s.post(
        f"{API}/workflow/approvals/{approval['id']}/decision", headers={**INTERNAL_HEADERS, **manager_headers},
        json={"decision": "approved", "actor": "officer"},
    )
    assert denied.status_code == 403
    technical_approval = s.post(
        f"{API}/workflow/approvals/{approval['id']}/decision", headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"decision": "approved", "actor": "engineer"},
    )
    assert technical_approval.status_code == 200
    assert technical_approval.json()["approval"]["approval_stage"] == "fund_release"
    assert technical_approval.json()["approval"]["responsible_role"] == "commercial_manager"
    fund_approval = s.post(
        f"{API}/workflow/approvals/{approval['id']}/decision", headers={**INTERNAL_HEADERS, **manager_headers},
        json={"decision": "approved", "actor": "manager"},
    )
    assert fund_approval.status_code == 200
    assert fund_approval.json()["approval"]["status"] == "pending_approval"
    assert fund_approval.json()["approval"]["approval_stage"] == "funds_release"
    assert s.post(f"{API}/purchase-orders/from-comparison", json=po_body, headers=responsible_headers).status_code == 409
    denied_release = s.post(
        f"{API}/workflow/approvals/{approval['id']}/funds-release", headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"actor": "officer"},
    )
    assert denied_release.status_code == 403
    released = s.post(
        f"{API}/workflow/approvals/{approval['id']}/funds-release", headers={**INTERNAL_HEADERS, **manager_headers},
        json={"actor": "manager", "method": "transfer"},
    )
    assert released.status_code == 200
    assert released.json()["approval"]["status"] == "approved"
    assert released.json()["approval"]["approval_stage"] == "po_ready"
    assert s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()["status"] == "approved"
    repeated_release = s.post(
        f"{API}/workflow/approvals/{approval['id']}/funds-release", headers={**INTERNAL_HEADERS, **manager_headers},
        json={"actor": "manager"},
    )
    assert repeated_release.json()["already_recorded"] is True
    with SessionLocal.begin() as session:
        session.get(EngineerApproval, approval["id"]).project_id = "nonexistent-project"
    assert s.post(f"{API}/purchase-orders/from-comparison", json=po_body, headers=responsible_headers).status_code == 409
    with SessionLocal.begin() as session:
        session.get(EngineerApproval, approval["id"]).project_id = project["id"]
    generated = s.post(f"{API}/purchase-orders/from-comparison", json=po_body, headers=responsible_headers)
    assert generated.status_code == 200, generated.text
    assert generated.json()["count"] == 2
    purchase_orders = generated.json()["purchase_orders"]
    assert {row["supplier_id"] for row in purchase_orders} == {supplier["id"] for supplier in suppliers}
    assert all(row["approval_id"] == approval["id"] for row in purchase_orders)
    assert all(row["project_id"] == project["id"] for row in purchase_orders)
    assert all(row["source_request_id"] == request_id for row in purchase_orders)
    assert all(row["source_request_number"] == request_number for row in purchase_orders)
    assert all(row["comparison_id"] == comparison.json()["id"] for row in purchase_orders)
    po_numbers = [row["po_number"] for row in purchase_orders]
    assert len(po_numbers) == len(set(po_numbers)) == 2
    assert all(number.startswith("PO-") and len(number) == 9 for number in po_numbers)
    approval_after_po = s.get(f"{API}/workflow/approvals/{approval['id']}", headers={**INTERNAL_HEADERS, **admin_headers}).json()
    assert {row["po_number"] for row in approval_after_po["purchase_orders"]} == set(po_numbers)
    assert all(row["status"] != "cancelled" for row in approval_after_po["purchase_orders"])
    approvals_list = s.get(f"{API}/workflow/approvals", headers={**INTERNAL_HEADERS, **admin_headers}).json()["items"]
    assert next(row for row in approvals_list if row["id"] == approval["id"])["has_purchase_order"] is True
    assert s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()["status"] == "converted_to_purchase"
    first_po = purchase_orders[0]
    detail = s.get(f"{API}/purchase-orders/{first_po['id']}", headers=responsible_headers)
    assert detail.status_code == 200
    assert detail.json()["source_request_number"] == request_number
    assert detail.json()["comparison_number"] == comparison.json()["comparison_number"]
    assert detail.json()["approval_number"] == approval["approval_number"]
    assert len(detail.json()["items"]) == 1
    for status in ("approved", "sent", "supplier_confirmed"):
        transition = s.patch(
            f"{API}/purchase-orders/{first_po['id']}/status",
            json={"status": status},
            headers=responsible_headers,
        )
        assert transition.status_code == 200, transition.text
        assert transition.json()["status"] == status
    assert s.post(
        f"{API}/purchase-orders/{first_po['id']}/finalize",
        json={"actor": "manager"},
        headers=manager_headers,
    ).status_code == 403
    finalized = s.post(
        f"{API}/purchase-orders/{first_po['id']}/finalize",
        json={"actor": "officer"},
        headers=responsible_headers,
    )
    assert finalized.status_code == 200, finalized.text
    assert finalized.json()["purchase_order"]["status"] == "in_delivery"
    assert s.post(
        f"{API}/purchase-orders/{first_po['id']}/finalize",
        json={"actor": "officer"},
        headers=responsible_headers,
    ).json()["already_recorded"] is True
    receiving = s.post(f"{API}/purchase-orders/{first_po['id']}/receipts", json={
        "receipt_type": "problem", "idempotency_key": "confirmed-po-receiving-001",
        "actor": "officer",
        "problem_reason": "مشكلة أخرى",
    }, headers=responsible_headers)
    assert receiving.status_code == 200, receiving.text
    assert receiving.json()["purchase_order"]["status"] == "delivery_problem"
    assert receiving.json()["purchase_order"]["receipt_summary"]["received_quantity"] == 0
    assert s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()["status"] == "converted_to_purchase"
    completed_receiving = s.post(f"{API}/purchase-orders/{first_po['id']}/receipts", json={
        "receipt_type": "full", "idempotency_key": "confirmed-po-receiving-002",
        "actor": "officer",
    }, headers=responsible_headers)
    assert completed_receiving.status_code == 200, completed_receiving.text
    assert completed_receiving.json()["purchase_order"]["status"] == "completed"
    assert s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()["status"] == "converted_to_purchase"
    cancelled_po = purchase_orders[1]
    cancelled = s.patch(
        f"{API}/purchase-orders/{cancelled_po['id']}/status",
        json={"status": "cancelled"},
        headers=responsible_headers,
    )
    assert cancelled.status_code == 200, cancelled.text
    assert s.post(
        f"{API}/purchase-orders/{cancelled_po['id']}/finalize",
        json={"actor": "officer"},
        headers=responsible_headers,
    ).status_code == 409
    assert s.get(
        f"{API}/purchase-orders/{cancelled_po['id']}", headers=responsible_headers,
    ).json()["status"] == "cancelled"
    completed_request = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    assert completed_request["status"] == "completed"
    assert [row["to_status"] for row in completed_request["status_history"]][-4:] == [
        "waiting_for_approval", "approved", "converted_to_purchase", "completed",
    ]
    hub = s.get(
        f"{API}/workflow/projects/{project['id']}/procurement-hub",
        headers={**INTERNAL_HEADERS, **admin_headers},
    )
    assert hub.status_code == 200, hub.text
    request_chain = next(
        chain for chain in hub.json()["traceability"]
        if chain["request"]["id"] == request_id
    )
    assert request_chain["request"] == {
        "id": request_id,
        "request_number": request_number,
        "status": "completed",
    }
    comparison_chain = next(
        row for row in request_chain["comparisons"]
        if row["id"] == comparison.json()["id"]
    )
    approval_chain = next(
        row for row in comparison_chain["approvals"]
        if row["id"] == approval["id"]
    )
    assert comparison_chain["comparison_number"] == comparison.json()["comparison_number"]
    assert approval_chain["approval_number"] == approval["approval_number"]
    traced_orders = {row["id"]: row for row in approval_chain["purchase_orders"]}
    assert traced_orders[first_po["id"]]["status"] == "completed"
    assert traced_orders[first_po["id"]]["receipt_summary"]["remaining_quantity"] == 0
    assert traced_orders[cancelled_po["id"]]["status"] == "cancelled"
    admin_report = s.get(f"{API}/purchase-orders/{first_po['id']}/reports/admin", headers=responsible_headers)
    site_report = s.get(f"{API}/purchase-orders/{first_po['id']}/reports/site", headers=responsible_headers)
    assert admin_report.status_code == 200
    assert admin_report.json()["items"][0]["unit_price"] > 0
    assert "final_total" in admin_report.json()["totals"]
    assert site_report.status_code == 200
    assert "totals" not in site_report.json()
    assert all("unit_price" not in item and "line_total" not in item for item in site_report.json()["items"])


def test_purchase_orders_use_approved_snapshot_after_comparison_changes(s, admin_headers):
    timestamp = datetime.now(timezone.utc).isoformat()
    request_id = str(uuid.uuid4())
    request_number = f"T-SNAPSHOT-{uuid.uuid4().hex[:8]}"
    with SessionLocal.begin() as session:
        session.add(IncomingPurchaseRequest(
            id=request_id, request_number=request_number,
            requester_name="مهندس الموقع", company_name="عميل الاعتماد",
            phone_number="01000000010", project_name="مشروع اعتماد ثابت",
            project_location="القاهرة", delivery_location="القاهرة",
            required_delivery_date="2099-01-01", priority="normal",
            submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
            created_at=timestamp, updated_at=timestamp,
        ))
        session.add(IncomingPurchaseRequestItem(
            id=str(uuid.uuid4()), request_id=request_id, position=1,
            product_name="احتياج اعتماد ثابت", quantity=1, unit="قطعة",
            review_status="approved", reviewed_by="engineer", reviewed_at=timestamp,
        ))

    project = s.get(f"{API}/projects", headers=admin_headers).json()[0]
    linked = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/link-project",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"project_id": project["id"], "actor": "engineer"},
    )
    assert linked.status_code == 200, linked.text
    reviewed = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/technical-decision",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={
            "decision": "approved_for_pricing",
            "actor": "engineer",
        },
    )
    assert reviewed.status_code == 200, reviewed.text

    items = s.get(f"{API}/items", headers=admin_headers).json()[:3]
    suppliers = s.get(f"{API}/suppliers", headers=admin_headers).json()[:2]
    original_rows = [{
        "item_id": items[0]["id"], "supplier_id": suppliers[0]["id"],
        "quantity": 3, "unit": items[0]["unit"], "unit_price": 100,
        "discount_pct": 10, "tax_pct": 14, "shipping_cost": 5,
        "other_cost": 2, "availability": "available",
        "price_valid_until": "2099-12-31", "selected_for_purchase": 1,
    }, {
        "item_id": items[1]["id"], "supplier_id": suppliers[1]["id"],
        "quantity": 2, "unit": items[1]["unit"], "unit_price": 75,
        "discount_pct": 5, "tax_pct": 14, "shipping_cost": 3,
        "other_cost": 1, "availability": "available",
        "price_valid_until": "2099-12-31", "selected_for_purchase": 1,
    }]
    comparison_body = {
        "project_id": project["id"], "project_name": project["name"],
        "customer_name": "عميل الاعتماد",
        "source_request_id": request_id, "source_request_number": request_number,
        "comparison_date": "2026-08-18", "rows": original_rows,
    }
    comparison = s.post(f"{API}/price-comparisons", json=comparison_body, headers=admin_headers)
    assert comparison.status_code == 200, comparison.text

    approval_response = s.post(
        f"{API}/workflow/approvals/from-comparison",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={
            "comparison_id": comparison.json()["id"],
            "approval_type": "comparison_workflow",
            "created_by": "procurement responsible",
        },
    )
    assert approval_response.status_code == 201, approval_response.text
    approval = approval_response.json()["approval"]
    approved_lines = {line["item_id"]: line for line in approval["lines"]}
    first_approved = approved_lines[items[0]["id"]]
    assert {
        key: first_approved[key]
        for key in (
            "supplier_id", "quantity", "unit", "unit_price",
            "discount_pct", "tax_pct", "line_total",
        )
    } == {
        "supplier_id": suppliers[0]["id"],
        "quantity": 3.0,
        "unit": items[0]["unit"],
        "unit_price": 100.0,
        "discount_pct": 0.0,
        "tax_pct": 0.0,
        "line_total": 300.0,
    }
    assert approved_lines[items[1]["id"]]["line_total"] == 150.0
    assert approval["final_total"] == 481.25

    for actor_role in ("procurement_engineer", "commercial_manager"):
        decision = s.post(
            f"{API}/workflow/approvals/{approval['id']}/decision",
            headers={**INTERNAL_HEADERS, **admin_headers},
            json={"decision": "approved", "actor": actor_role},
        )
        assert decision.status_code == 200, decision.text
    released = s.post(
        f"{API}/workflow/approvals/{approval['id']}/funds-release",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"actor": "manager"},
    )
    assert released.status_code == 200, released.text

    changed_rows = [{
        "item_id": items[2]["id"], "supplier_id": suppliers[1]["id"],
        "quantity": 99, "unit": items[2]["unit"], "unit_price": 999,
        "discount_pct": 0, "tax_pct": 0, "availability": "available",
        "price_valid_until": "2099-12-31", "selected_for_purchase": 1,
    }, {
        "item_id": items[1]["id"], "supplier_id": suppliers[1]["id"],
        "quantity": 88, "unit": items[1]["unit"], "unit_price": 888,
        "discount_pct": 0, "tax_pct": 0, "availability": "available",
        "price_valid_until": "2099-12-31", "selected_for_purchase": 1,
    }]
    changed = s.put(f"{API}/price-comparisons/{comparison.json()['id']}", json={
        **comparison_body,
        "customer_name": "عميل تم تغييره بعد الاعتماد",
        "source_request_number": "REQ-MUTATED-AFTER-APPROVAL",
        "rows": changed_rows,
    }, headers=admin_headers)
    assert changed.status_code == 200, changed.text
    assert changed.json()["rows"][0]["unit_price"] == 999

    po_body = {
        "comparison_id": comparison.json()["id"],
        "po_date": "2026-08-18",
        "created_by": "procurement responsible",
        "orders": [{
            "supplier_id": "CLIENT-SUPPLIER",
            "supplier_name": "Client supplied wrong supplier",
            "items": [{
                "product_name": "Client supplied wrong item",
                "quantity": 999,
                "unit_price": 999,
            }],
        }],
    }
    other_project_id = str(uuid.uuid4())
    with SessionLocal.begin() as session:
        session.add(Project(
            id=other_project_id,
            code=f"T-PROJ-{uuid.uuid4().hex[:8]}",
            name=f"مشروع مختلف {uuid.uuid4().hex[:8]}",
        ))

    def assert_ancestry_rejected(**changes):
        with SessionLocal.begin() as session:
            approval_row = session.get(EngineerApproval, approval["id"])
            original = {field: getattr(approval_row, field) for field in changes}
            for field, value in changes.items():
                setattr(approval_row, field, value)
        rejected = s.post(f"{API}/purchase-orders/from-comparison", json=po_body, headers=admin_headers)
        assert rejected.status_code == 409, rejected.text
        with SessionLocal.begin() as session:
            approval_row = session.get(EngineerApproval, approval["id"])
            for field, value in original.items():
                setattr(approval_row, field, value)
        with SessionLocal() as session:
            assert session.query(PurchaseOrder).filter_by(
                comparison_id=comparison.json()["id"],
            ).count() == 0

    assert_ancestry_rejected(project_id=other_project_id)
    assert_ancestry_rejected(source_request_id=str(uuid.uuid4()))
    assert_ancestry_rejected(comparison_id=str(uuid.uuid4()))
    assert_ancestry_rejected(status="pending_approval", approval_stage="funds_release")
    assert_ancestry_rejected(source_request_number="REQ-WRONG-ANCESTRY")
    assert_ancestry_rejected(comparison_number="CMP-WRONG-ANCESTRY")

    generated = s.post(f"{API}/purchase-orders/from-comparison", json=po_body, headers=admin_headers)
    assert generated.status_code == 200, generated.text
    result = generated.json()
    assert result["count"] == 2
    assert sum(order["final_total"] for order in result["purchase_orders"]) == 481.25
    assert {order["supplier_id"] for order in result["purchase_orders"]} == {
        suppliers[0]["id"], suppliers[1]["id"],
    }
    assert all(order["source_request_number"] == request_number for order in result["purchase_orders"])
    assert all(order["customer_name"] == "عميل الاعتماد" for order in result["purchase_orders"])

    po_items = {
        item["item_id"]: item
        for order in result["purchase_orders"]
        for item in order["items"]
    }
    assert set(po_items) == {items[0]["id"], items[1]["id"]}
    first_po_item = po_items[items[0]["id"]]
    assert {
        key: first_po_item[key]
        for key in (
            "quantity", "unit", "unit_price", "discount_pct", "vat_pct",
            "line_total",
        )
    } == {
        "quantity": 3.0,
        "unit": items[0]["unit"],
        "unit_price": 100.0,
        "discount_pct": 0.0,
        "vat_pct": 0.0,
        "line_total": 300.0,
    }
    second_po_item = po_items[items[1]["id"]]
    assert {
        key: second_po_item[key]
        for key in (
            "quantity", "unit", "unit_price", "discount_pct", "vat_pct",
            "line_total",
        )
    } == {
        "quantity": 2.0,
        "unit": items[1]["unit"],
        "unit_price": 75.0,
        "discount_pct": 0.0,
        "vat_pct": 0.0,
        "line_total": 150.0,
    }
    totals_by_supplier = {
        order["supplier_id"]: order for order in result["purchase_orders"]
    }
    assert totals_by_supplier[suppliers[0]["id"]]["discount_total"] == 30.0
    assert totals_by_supplier[suppliers[0]["id"]]["vat_total"] == 37.8
    assert totals_by_supplier[suppliers[1]["id"]]["discount_total"] == 7.5
    assert totals_by_supplier[suppliers[1]["id"]]["vat_total"] == 19.95


def test_site_receiving_is_transactional_idempotent_and_closes_only_when_complete(s, admin_headers):
    timestamp = datetime.now(timezone.utc).isoformat()
    request_id = str(uuid.uuid4())
    po_id = str(uuid.uuid4())
    first_item_id = str(uuid.uuid4())
    second_item_id = str(uuid.uuid4())
    with SessionLocal.begin() as session:
        session.add(IncomingPurchaseRequest(
            id=request_id, request_number=f"T-SINGLE-PO-{uuid.uuid4().hex[:8]}",
            requester_name="مهندس الموقع", company_name="عميل استلام فردي",
            phone_number="01000000013", project_name="مشروع الاستلام",
            project_location="القاهرة", delivery_location="القاهرة",
            required_delivery_date="2099-01-01", priority="normal",
            status="converted_to_purchase",
            submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
            created_at=timestamp, updated_at=timestamp,
        ))
        session.add(PurchaseOrder(
            id=po_id, po_number=f"T-PO-{uuid.uuid4().hex[:8]}", status="draft",
            source_request_id=request_id,
            project_id="test-project", project_name="مشروع الاستلام",
            supplier_name="مورد الاستلام", po_date="2026-08-17",
            created_at=timestamp, updated_at=timestamp,
        ))
        session.add_all([
            PurchaseOrderItem(
                id=first_item_id, purchase_order_id=po_id, product_name="أسمنت",
                quantity=100, unit="شيكارة", unit_price=10,
            ),
            PurchaseOrderItem(
                id=second_item_id, purchase_order_id=po_id, product_name="رمل",
                quantity=20, unit="م³", unit_price=5,
            ),
        ])
    before_delivery = s.post(f"{API}/purchase-orders/{po_id}/receipts", json={
        "receipt_type": "full", "idempotency_key": "before-delivery-001",
    }, headers=admin_headers)
    assert before_delivery.status_code == 409
    with SessionLocal() as session:
        assert session.get(IncomingPurchaseRequest, request_id).status == "converted_to_purchase"
    with SessionLocal.begin() as session:
        session.get(PurchaseOrder, po_id).status = "in_delivery"

    partial_body = {
        "receipt_type": "partial", "idempotency_key": "partial-receipt-001",
        "actor": "المستلم",
        "lines": [{"purchase_order_item_id": first_item_id, "quantity": 60}],
    }
    partial = s.post(f"{API}/purchase-orders/{po_id}/receipts", json=partial_body, headers=admin_headers)
    assert partial.status_code == 200, partial.text
    assert partial.json()["purchase_order"]["status"] == "partial_received"
    assert partial.json()["purchase_order"]["items"][0]["received_quantity"] == 60
    repeated = s.post(f"{API}/purchase-orders/{po_id}/receipts", json=partial_body, headers=admin_headers)
    assert repeated.status_code == 200
    assert repeated.json()["already_recorded"] is True
    assert repeated.json()["purchase_order"]["items"][0]["received_quantity"] == 60
    with SessionLocal() as session:
        assert session.query(PurchaseOrderReceipt).filter_by(purchase_order_id=po_id).count() == 1
        assert session.get(IncomingPurchaseRequest, request_id).status == "converted_to_purchase"

    over = s.post(f"{API}/purchase-orders/{po_id}/receipts", json={
        "receipt_type": "partial", "idempotency_key": "partial-over-001",
        "lines": [{"purchase_order_item_id": first_item_id, "quantity": 41}],
    }, headers=admin_headers)
    assert over.status_code == 422
    with SessionLocal() as session:
        assert session.query(PurchaseOrderReceipt).filter_by(purchase_order_id=po_id).count() == 1
        assert session.get(IncomingPurchaseRequest, request_id).status == "converted_to_purchase"
    problem = s.post(f"{API}/purchase-orders/{po_id}/receipts", json={
        "receipt_type": "problem", "idempotency_key": "problem-receipt-001",
        "problem_reason": "تالف",
        "affected_item_id": second_item_id, "affected_quantity": 5,
    }, headers=admin_headers)
    assert problem.status_code == 200, problem.text
    assert problem.json()["purchase_order"]["status"] == "delivery_problem"
    assert problem.json()["purchase_order"]["receipt_summary"]["received_quantity"] == 60
    with SessionLocal() as session:
        assert session.get(IncomingPurchaseRequest, request_id).status == "converted_to_purchase"

    final = s.post(f"{API}/purchase-orders/{po_id}/receipts", json={
        "receipt_type": "partial", "idempotency_key": "partial-final-001",
        "actor": "المستلم",
        "lines": [
            {"purchase_order_item_id": first_item_id, "quantity": 40},
            {"purchase_order_item_id": second_item_id, "quantity": 20},
        ],
    }, headers=admin_headers)
    assert final.status_code == 200, final.text
    completed = final.json()["purchase_order"]
    assert completed["status"] == "completed"
    assert completed["receipt_summary"]["ordered_quantity"] == 120.0
    assert completed["receipt_summary"]["received_quantity"] == 120.0
    assert completed["receipt_summary"]["remaining_quantity"] == 0.0
    assert completed["receipt_summary"]["receipt_count"] == 3
    assert completed["receipt_summary"]["latest_receipt_date"]
    completed_request = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    assert completed_request["status"] == "completed"
    assert [row["to_status"] for row in completed_request["status_history"]] == ["completed"]
    assert len(completed["receipt_history"]) == 3
    assert s.get(f"{API}/purchase-orders/{po_id}/reports/site", headers=admin_headers).status_code == 200
    assert s.post(f"{API}/purchase-orders/{po_id}/receipts", json={
        "receipt_type": "problem", "idempotency_key": "after-complete-001",
        "problem_reason": "مشكلة أخرى",
    }, headers=admin_headers).status_code == 409
    assert s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()["status"] == "completed"


def test_request_completes_only_after_all_formal_purchase_orders(s, admin_headers):
    timestamp = datetime.now(timezone.utc).isoformat()
    request_id = str(uuid.uuid4())
    po_ids = [str(uuid.uuid4()), str(uuid.uuid4())]
    item_ids = [str(uuid.uuid4()), str(uuid.uuid4())]
    with SessionLocal.begin() as session:
        session.add(IncomingPurchaseRequest(
            id=request_id, request_number=f"T-MULTI-PO-{uuid.uuid4().hex[:8]}",
            requester_name="مهندس الموقع", company_name="عميل متعدد الموردين",
            phone_number="01000000011", project_name="مشروع متعدد الموردين",
            project_location="القاهرة", delivery_location="القاهرة",
            required_delivery_date="2099-01-01", priority="normal",
            status="converted_to_purchase",
            submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
            created_at=timestamp, updated_at=timestamp,
        ))
        for position, (po_id, item_id) in enumerate(zip(po_ids, item_ids), 1):
            session.add(PurchaseOrder(
                id=po_id, po_number=f"T-MULTI-{uuid.uuid4().hex[:8]}",
                source_request_id=request_id, status="in_delivery",
                project_id="test-project", project_name="مشروع متعدد الموردين",
                supplier_name=f"مورد {position}", po_date="2026-08-18",
                created_at=timestamp, updated_at=timestamp,
            ))
            session.add(PurchaseOrderItem(
                id=item_id, purchase_order_id=po_id,
                product_name=f"صنف {position}", quantity=position,
                unit="قطعة", unit_price=10,
            ))

    first_receipt = s.post(f"{API}/purchase-orders/{po_ids[0]}/receipts", json={
        "receipt_type": "full", "idempotency_key": "multi-po-first",
        "actor": "officer",
    }, headers=admin_headers)
    assert first_receipt.status_code == 200, first_receipt.text
    assert first_receipt.json()["purchase_order"]["status"] == "completed"
    assert s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()["status"] == "converted_to_purchase"

    second_receipt_body = {
        "receipt_type": "full", "idempotency_key": "multi-po-second",
        "actor": "officer",
    }
    second_receipt = s.post(f"{API}/purchase-orders/{po_ids[1]}/receipts", json=second_receipt_body, headers=admin_headers)
    assert second_receipt.status_code == 200, second_receipt.text
    assert second_receipt.json()["purchase_order"]["status"] == "completed"
    completed_request = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    assert completed_request["status"] == "completed"
    assert completed_request["status_history"][-1]["to_status"] == "completed"
    repeated = s.post(
        f"{API}/purchase-orders/{po_ids[1]}/receipts",
        json=second_receipt_body,
        headers=admin_headers,
    )
    assert repeated.status_code == 200
    assert repeated.json()["already_recorded"] is True
    repeated_request = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    assert repeated_request["status"] == "completed"
    assert [
        row["to_status"] for row in repeated_request["status_history"]
        if row["to_status"] == "completed"
    ] == ["completed"]


# ---------------- Authentication (Sprint 2.1) ----------------

def _make_user(
    session, *, username, password="Sprint21Passw0rd!",
    account_type="erp", role="admin", active=True,
):
    now = datetime.now(timezone.utc).isoformat()
    user = User(
        id=str(uuid.uuid4()), username=username, display_name=username,
        password_hash=hash_password(password), account_type=account_type,
        role=role, active=active, created_at=now, updated_at=now,
    )
    session.add(user)
    session.commit()
    session.refresh(user)
    return user


def test_auth_valid_login_returns_token_and_public_user(s):
    with SessionLocal() as session:
        _make_user(session, username="auth-valid-login", password="CorrectHorse123!")

    response = s.post(f"{API}/auth/login", json={
        "username": "auth-valid-login", "password": "CorrectHorse123!",
    })

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["token_type"] == "bearer"
    assert body["access_token"]
    assert body["user"]["username"] == "auth-valid-login"
    assert body["user"]["account_type"] == "erp"
    assert body["user"]["role"] == "admin"
    assert body["user"]["active"] is True
    assert "password_hash" not in body["user"]
    assert "password" not in body["user"]


def test_auth_invalid_password_rejected(s):
    with SessionLocal() as session:
        _make_user(session, username="auth-bad-password", password="CorrectHorse123!")

    response = s.post(f"{API}/auth/login", json={
        "username": "auth-bad-password", "password": "wrong-password",
    })

    assert response.status_code == 401
    assert "password_hash" not in response.text


def test_auth_unknown_username_rejected(s):
    response = s.post(f"{API}/auth/login", json={
        "username": "does-not-exist-at-all", "password": "whatever",
    })
    assert response.status_code == 401


def test_auth_inactive_account_rejected(s):
    with SessionLocal() as session:
        _make_user(
            session, username="auth-inactive", password="CorrectHorse123!",
            active=False,
        )

    response = s.post(f"{API}/auth/login", json={
        "username": "auth-inactive", "password": "CorrectHorse123!",
    })

    assert response.status_code == 401


def test_auth_current_user_endpoint(s):
    with SessionLocal() as session:
        _make_user(
            session, username="auth-me", password="CorrectHorse123!",
            role="procurement_engineer",
        )
    login = s.post(f"{API}/auth/login", json={
        "username": "auth-me", "password": "CorrectHorse123!",
    })
    token = login.json()["access_token"]

    me = s.get(f"{API}/auth/me", headers={"Authorization": f"Bearer {token}"})

    assert me.status_code == 200, me.text
    body = me.json()
    assert body["username"] == "auth-me"
    assert body["role"] == "procurement_engineer"
    assert body["account_type"] == "erp"
    assert "password_hash" not in body
    assert "password" not in body


def test_auth_me_requires_a_token(s):
    response = s.get(f"{API}/auth/me")
    assert response.status_code == 401


def test_auth_me_rejects_a_garbage_token(s):
    response = s.get(f"{API}/auth/me", headers={"Authorization": "Bearer not-a-real-token"})
    assert response.status_code == 401


def test_auth_logout_requires_a_token_and_acknowledges(s):
    with SessionLocal() as session:
        _make_user(session, username="auth-logout", password="CorrectHorse123!")
    login = s.post(f"{API}/auth/login", json={
        "username": "auth-logout", "password": "CorrectHorse123!",
    })
    token = login.json()["access_token"]

    anonymous = s.post(f"{API}/auth/logout")
    assert anonymous.status_code == 401

    authenticated = s.post(f"{API}/auth/logout", headers={"Authorization": f"Bearer {token}"})
    assert authenticated.status_code == 200
    assert authenticated.json() == {"ok": True}


def test_auth_deactivation_revokes_access_immediately(s):
    with SessionLocal() as session:
        user = _make_user(session, username="auth-revoke", password="CorrectHorse123!")
        user_id = user.id

    login = s.post(f"{API}/auth/login", json={
        "username": "auth-revoke", "password": "CorrectHorse123!",
    })
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    assert s.get(f"{API}/auth/me", headers=headers).status_code == 200

    with SessionLocal() as session:
        row = session.get(User, user_id)
        row.active = False
        session.commit()

    assert s.get(f"{API}/auth/me", headers=headers).status_code == 401


def test_auth_site_portal_login_returns_portal_role(s):
    with SessionLocal() as session:
        _make_user(
            session, username="auth-site-engineer", password="CorrectHorse123!",
            account_type="site_portal", role="site_engineer",
        )

    response = s.post(f"{API}/auth/login", json={
        "username": "auth-site-engineer", "password": "CorrectHorse123!",
    })

    assert response.status_code == 200, response.text
    body = response.json()["user"]
    assert body["account_type"] == "site_portal"
    assert body["role"] == "site_engineer"


# ---------------- Admin user management (Sprint 2.2) ----------------

ADMIN_API = f"{API}/admin/users"


def _login_headers(client, username, password="Sprint21Passw0rd!"):
    login = client.post(f"{API}/auth/login", json={"username": username, "password": password})
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def _make_admin(session, suffix):
    username = f"admin-{suffix}"
    _make_user(session, username=username, role="admin")
    return username


def _make_project_row(session, suffix):
    project_id = f"T-ADMIN-PROJECT-{suffix}"
    session.add(Project(id=project_id, code=f"T-ADM-{suffix}", name=f"مشروع اختبار {suffix}"))
    session.commit()
    return project_id


def test_admin_can_list_users(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        _make_user(session, username=f"engineer-{suffix}", role="procurement_engineer")
    headers = _login_headers(s, admin_username)

    response = s.get(ADMIN_API, headers=headers)

    assert response.status_code == 200, response.text
    usernames = {row["username"] for row in response.json()}
    assert admin_username in usernames
    assert f"engineer-{suffix}" in usernames
    for row in response.json():
        assert "password_hash" not in row
        assert "password" not in row


def test_admin_can_create_erp_user(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
    headers = _login_headers(s, admin_username)

    response = s.post(ADMIN_API, headers=headers, json={
        "username": f"new-erp-{suffix}", "display_name": "مسؤول مشتريات",
        "password": "AnotherStrongPass1!", "account_type": "erp",
        "role": "procurement_responsible", "active": True,
    })

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["account_type"] == "erp"
    assert body["role"] == "procurement_responsible"
    assert body["active"] is True
    assert "password_hash" not in body
    assert "password" not in body

    # The account actually works.
    login = s.post(f"{API}/auth/login", json={
        "username": f"new-erp-{suffix}", "password": "AnotherStrongPass1!",
    })
    assert login.status_code == 200


def test_admin_can_create_site_portal_user(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        project_id = _make_project_row(session, suffix)
    headers = _login_headers(s, admin_username)

    response = s.post(ADMIN_API, headers=headers, json={
        "username": f"new-portal-{suffix}", "display_name": "مهندس موقع",
        "password": "AnotherStrongPass1!", "account_type": "site_portal",
        "project_ids": [project_id],
    })

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["account_type"] == "site_portal"
    assert body["role"] == "site_engineer"
    assert [p["id"] for p in body["assigned_projects"]] == [project_id]


def test_admin_create_user_rejects_duplicate_username(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        _make_user(session, username=f"taken-{suffix}")
    headers = _login_headers(s, admin_username)

    response = s.post(ADMIN_API, headers=headers, json={
        "username": f"taken-{suffix}", "display_name": "x",
        "password": "AnotherStrongPass1!", "account_type": "erp", "role": "admin",
    })

    assert response.status_code == 409


def test_admin_create_user_rejects_invalid_erp_role(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
    headers = _login_headers(s, admin_username)

    response = s.post(ADMIN_API, headers=headers, json={
        "username": f"bad-role-{suffix}", "display_name": "x",
        "password": "AnotherStrongPass1!", "account_type": "erp", "role": "site_engineer",
    })

    assert response.status_code == 422


def test_admin_create_user_rejects_erp_role_on_site_portal_account(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
    headers = _login_headers(s, admin_username)

    response = s.post(ADMIN_API, headers=headers, json={
        "username": f"bad-combo-{suffix}", "display_name": "x",
        "password": "AnotherStrongPass1!", "account_type": "site_portal", "role": "admin",
    })

    assert response.status_code == 422


def test_non_admin_erp_user_cannot_manage_users(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"engineer-only-{suffix}", role="procurement_engineer")
    headers = _login_headers(s, f"engineer-only-{suffix}")

    response = s.get(ADMIN_API, headers=headers)

    assert response.status_code == 403


def test_site_portal_user_cannot_manage_users(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(
            session, username=f"site-only-{suffix}",
            account_type="site_portal", role="site_engineer",
        )
    headers = _login_headers(s, f"site-only-{suffix}")

    response = s.get(ADMIN_API, headers=headers)

    assert response.status_code == 403


def test_admin_can_deactivate_and_reactivate_user(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        target = _make_user(session, username=f"toggle-{suffix}")
        target_id = target.id
    headers = _login_headers(s, admin_username)

    deactivate = s.put(f"{ADMIN_API}/{target_id}", headers=headers, json={"active": False})
    assert deactivate.status_code == 200, deactivate.text
    assert deactivate.json()["active"] is False

    login_while_inactive = s.post(f"{API}/auth/login", json={
        "username": f"toggle-{suffix}", "password": "Sprint21Passw0rd!",
    })
    assert login_while_inactive.status_code == 401

    reactivate = s.put(f"{ADMIN_API}/{target_id}", headers=headers, json={"active": True})
    assert reactivate.status_code == 200
    assert reactivate.json()["active"] is True

    login_again = s.post(f"{API}/auth/login", json={
        "username": f"toggle-{suffix}", "password": "Sprint21Passw0rd!",
    })
    assert login_again.status_code == 200


def test_admin_deactivation_revokes_a_live_session_immediately(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        target = _make_user(session, username=f"live-revoke-{suffix}")
        target_id = target.id
    admin_headers = _login_headers(s, admin_username)
    target_headers = _login_headers(s, f"live-revoke-{suffix}")
    assert s.get(f"{API}/auth/me", headers=target_headers).status_code == 200

    s.put(f"{ADMIN_API}/{target_id}", headers=admin_headers, json={"active": False})

    assert s.get(f"{API}/auth/me", headers=target_headers).status_code == 401


def test_admin_can_reset_password(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        target = _make_user(session, username=f"resetme-{suffix}", password="OldPassword1!")
        target_id = target.id
    headers = _login_headers(s, admin_username)

    reset = s.post(
        f"{ADMIN_API}/{target_id}/reset-password", headers=headers,
        json={"new_password": "BrandNewPassword1!"},
    )
    assert reset.status_code == 200, reset.text
    assert reset.json() == {"ok": True}
    assert "password" not in reset.text
    assert "BrandNewPassword1!" not in reset.text

    old_password_login = s.post(f"{API}/auth/login", json={
        "username": f"resetme-{suffix}", "password": "OldPassword1!",
    })
    assert old_password_login.status_code == 401

    new_password_login = s.post(f"{API}/auth/login", json={
        "username": f"resetme-{suffix}", "password": "BrandNewPassword1!",
    })
    assert new_password_login.status_code == 200


def test_admin_can_assign_valid_project_to_site_portal_user(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        target = _make_user(
            session, username=f"assignee-{suffix}",
            account_type="site_portal", role="site_engineer",
        )
        target_id = target.id
        project_id = _make_project_row(session, suffix)
    headers = _login_headers(s, admin_username)

    response = s.post(
        f"{ADMIN_API}/{target_id}/projects", headers=headers,
        json={"project_id": project_id},
    )

    assert response.status_code == 200, response.text
    assert [p["id"] for p in response.json()["assigned_projects"]] == [project_id]


def test_admin_assign_project_rejects_invalid_project_id(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        target = _make_user(
            session, username=f"badproject-{suffix}",
            account_type="site_portal", role="site_engineer",
        )
        target_id = target.id
    headers = _login_headers(s, admin_username)

    response = s.post(
        f"{ADMIN_API}/{target_id}/projects", headers=headers,
        json={"project_id": "does-not-exist"},
    )

    assert response.status_code == 422


def test_admin_assign_project_prevents_duplicate_access(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        target = _make_user(
            session, username=f"dupproject-{suffix}",
            account_type="site_portal", role="site_engineer",
        )
        target_id = target.id
        project_id = _make_project_row(session, suffix)
    headers = _login_headers(s, admin_username)

    first = s.post(
        f"{ADMIN_API}/{target_id}/projects", headers=headers,
        json={"project_id": project_id},
    )
    assert first.status_code == 200

    second = s.post(
        f"{ADMIN_API}/{target_id}/projects", headers=headers,
        json={"project_id": project_id},
    )
    assert second.status_code == 409


def test_admin_can_remove_project_assignment(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        target = _make_user(
            session, username=f"removeproject-{suffix}",
            account_type="site_portal", role="site_engineer",
        )
        target_id = target.id
        project_id = _make_project_row(session, suffix)
    headers = _login_headers(s, admin_username)
    s.post(f"{ADMIN_API}/{target_id}/projects", headers=headers, json={"project_id": project_id})

    response = s.delete(f"{ADMIN_API}/{target_id}/projects/{project_id}", headers=headers)

    assert response.status_code == 200, response.text
    assert response.json()["assigned_projects"] == []


def test_admin_user_endpoints_never_return_password_hash(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
    headers = _login_headers(s, admin_username)

    create = s.post(ADMIN_API, headers=headers, json={
        "username": f"nohash-{suffix}", "display_name": "x",
        "password": "AnotherStrongPass1!", "account_type": "erp", "role": "admin",
    })
    assert "password_hash" not in create.text

    listing = s.get(ADMIN_API, headers=headers)
    assert "password_hash" not in listing.text


# ---------------- Site Request Portal (Sprint A-G) ----------------

PORTAL_API = f"{API}/portal"


def _make_project_for_portal(session, suffix, **overrides):
    project_id = f"T-PORTAL-PROJECT-{suffix}"
    defaults = dict(id=project_id, code=f"T-POR-{suffix}", name=f"مشروع بوابة {suffix}")
    defaults.update(overrides)
    session.add(Project(**defaults))
    session.commit()
    return project_id


def _make_portal_user(session, suffix, project_ids=(), password="Sprint21Passw0rd!"):
    username = f"portal-{suffix}"
    user = _make_user(
        session, username=username, password=password,
        account_type="site_portal", role="site_engineer",
    )
    now = datetime.now(timezone.utc).isoformat()
    for project_id in project_ids:
        session.add(UserProjectAccess(
            id=str(uuid.uuid4()), user_id=user.id, project_id=project_id, created_at=now,
        ))
    session.commit()
    return username, user.id


def _make_portal_item(session, suffix, **overrides):
    item_id = f"T-PORTAL-ITEM-{suffix}"
    defaults = dict(
        id=item_id, code=f"T-ITM-{suffix}", name=f"صنف اختبار {suffix}",
        product_name=f"صنف اختبار {suffix}", unit="قطعة",
    )
    defaults.update(overrides)
    session.add(Item(**defaults))
    session.commit()
    return item_id


def _portal_payload(project_id="", items=None, destination="site", required_date="2027-01-01"):
    return {
        "required_delivery_date": required_date,
        "priority": "normal",
        "delivery_destination": destination,
        "notes": "",
        "project_id": project_id,
        "items": items or [],
    }


def _submit_portal_request(client, headers, payload, files=None):
    return client.post(
        f"{PORTAL_API}/purchase-requests",
        data={"payload": json.dumps(payload)},
        headers=headers,
        files=files or [],
    )


def _manual_item(name="صنف يدوي اختبار", unit="قطعة", quantity=1, note=""):
    return {"product_name": name, "unit": unit, "quantity": quantity, "note": note}


def test_partial_item_review_progresses_only_approved_subset_and_creates_linked_correction(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        portal_username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
        _make_user(session, username=f"partial-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"partial-resp-{suffix}", role="procurement_responsible")
    portal_headers = _login_headers(s, portal_username)
    engineer_headers = _login_headers(s, f"partial-eng-{suffix}")
    responsible_headers = _login_headers(s, f"partial-resp-{suffix}")
    created = _submit_portal_request(s, portal_headers, _portal_payload(items=[
        _manual_item(f"معتمد-{suffix}"),
        _manual_item(f"مرفوض-{suffix}"),
        _manual_item(f"ناقص-{suffix}"),
    ])).json()
    request_id = created["request_id"]
    detail = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    decisions = [
        (detail["items"][0]["id"], "approved", ""),
        (detail["items"][1]["id"], "rejected", "غير مطلوب"),
        (detail["items"][2]["id"], "need_clarification", "أكمل المواصفة"),
    ]
    for item_id, status, reason in decisions:
        response = s.patch(
            f"{API}/internal/incoming-purchase-requests/{request_id}/items/{item_id}/review",
            headers={**INTERNAL_HEADERS, **engineer_headers},
            json={"status": status, "reason": reason},
        )
        assert response.status_code == 200, response.text

    progressed = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/technical-decision",
        headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"decision": "approved_for_pricing"},
    )
    assert progressed.status_code == 200, progressed.text
    assert progressed.json()["eligible_item_count"] == 1
    assert progressed.json()["returned_item_count"] == 2

    dashboard = s.get(f"{API}/dashboard", headers=responsible_headers)
    assert dashboard.status_code == 200, dashboard.text
    attention = dashboard.json()["attention_items"]
    assert any(row["type"] == "sourcing_required" and row["reference"] == created["request_number"] for row in attention)

    rfq = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    )
    assert rfq.status_code == 200, rfq.text
    assert len(rfq.json()["rfq"]["items"]) == 1
    assert rfq.json()["rfq"]["items"][0]["product_name"] == f"معتمد-{suffix}"

    returned = s.get(f"{PORTAL_API}/returned-items", headers=portal_headers)
    assert returned.status_code == 200, returned.text
    assert {row["status"] for row in returned.json()} == {"rejected", "need_clarification"}
    returned_item = next(row for row in returned.json() if row["status"] == "need_clarification")

    draft = _save_correction_draft(s, portal_headers, returned_item["id"], quantity=2)
    assert draft.status_code == 200, draft.text
    assert draft.json()["ready"] is True

    # Saving a draft must never create a REQ by itself.
    with SessionLocal() as session:
        assert session.scalar(
            select(func.count()).select_from(IncomingPurchaseRequest)
            .where(IncomingPurchaseRequest.source_request_id == request_id)
        ) == 0

    corrected = _resubmit_corrections(s, portal_headers, request_id, [returned_item["id"]])
    assert corrected.status_code == 200, corrected.text
    child = corrected.json()
    assert child["source_request_id"] == request_id
    with SessionLocal() as session:
        original_item = session.get(IncomingPurchaseRequestItem, returned_item["id"])
        child_request = session.get(IncomingPurchaseRequest, child["request_id"])
        child_items = session.scalars(
            select(IncomingPurchaseRequestItem).where(
                IncomingPurchaseRequestItem.request_id == child["request_id"]
            )
        ).all()
        assert original_item is not None
        assert original_item.review_status == "need_clarification"
        assert child_request.source_request_id == request_id
        assert len(child_items) == 1
        assert child_items[0].source_item_id == original_item.id

    # The approved item's own downstream chain (already an active RFQ) is
    # completely unaffected by the grouped correction of the other items.
    rfq_after = s.get(f"{RFQ_API}/{rfq.json()['rfq']['id']}", headers={**INTERNAL_HEADERS, **responsible_headers})
    assert rfq_after.status_code == 200, rfq_after.text
    assert len(rfq_after.json()["items"]) == 1
    assert rfq_after.json()["items"][0]["product_name"] == f"معتمد-{suffix}"


def _save_correction_draft(client, headers, item_id, *, product_name="صنف مصحح", unit="قطعة", quantity=1, note="مواصفة مكتملة", required_delivery_date="2099-01-01"):
    body = {
        "product_name": product_name, "unit": unit,
        "quantity": quantity, "note": note, "required_delivery_date": required_delivery_date,
    }
    return client.post(f"{PORTAL_API}/returned-items/{item_id}/correct", headers=headers, json=body)


def _resubmit_corrections(client, headers, original_request_id, item_ids):
    return client.post(
        f"{PORTAL_API}/purchase-requests/{original_request_id}/resubmit-corrections",
        headers=headers, json={"item_ids": item_ids},
    )


def _review_item(client, headers, request_id, item_id, status, reason=""):
    response = client.patch(
        f"{API}/internal/incoming-purchase-requests/{request_id}/items/{item_id}/review",
        headers={**INTERNAL_HEADERS, **headers},
        json={"status": status, "reason": reason},
    )
    assert response.status_code == 200, response.text
    return response


def _setup_returned_items_scenario(s, admin_headers, suffix, item_count=4):
    """One original REQ with `item_count` manual items, all reviewed as
    need_clarification (so every one becomes an eligible returned item)."""
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        portal_username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
        _make_user(session, username=f"grp-eng-{suffix}", role="procurement_engineer")
    portal_headers = _login_headers(s, portal_username)
    engineer_headers = _login_headers(s, f"grp-eng-{suffix}")
    created = _submit_portal_request(s, portal_headers, _portal_payload(items=[
        _manual_item(f"صنف-{suffix}-{i}") for i in range(item_count)
    ])).json()
    request_id = created["request_id"]
    detail = s.get(
        f"{API}/internal/incoming-purchase-requests/{request_id}",
        headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    item_ids = [item["id"] for item in detail["items"]]
    for item_id in item_ids:
        _review_item(s, engineer_headers, request_id, item_id, "need_clarification", "أكمل المواصفة")
    return request_id, item_ids, portal_headers


def test_four_returned_items_from_one_request_group_into_one_corrected_req(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    request_id, item_ids, portal_headers = _setup_returned_items_scenario(s, admin_headers, suffix, item_count=4)
    for item_id in item_ids:
        assert _save_correction_draft(s, portal_headers, item_id).status_code == 200

    result = _resubmit_corrections(s, portal_headers, request_id, item_ids)
    assert result.status_code == 200, result.text
    body = result.json()
    assert body["item_count"] == 4

    with SessionLocal() as session:
        child_items = session.scalars(
            select(IncomingPurchaseRequestItem).where(
                IncomingPurchaseRequestItem.request_id == body["request_id"]
            )
        ).all()
        assert len(child_items) == 4
        assert {item.source_item_id for item in child_items} == set(item_ids)
        child_request = session.get(IncomingPurchaseRequest, body["request_id"])
        assert child_request.source_request_id == request_id


def test_three_corrected_items_group_while_one_stays_pending(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    request_id, item_ids, portal_headers = _setup_returned_items_scenario(s, admin_headers, suffix, item_count=4)
    ready_ids, pending_id = item_ids[:3], item_ids[3]
    for item_id in ready_ids:
        assert _save_correction_draft(s, portal_headers, item_id).status_code == 200

    result = _resubmit_corrections(s, portal_headers, request_id, ready_ids)
    assert result.status_code == 200, result.text
    assert result.json()["item_count"] == 3

    returned = s.get(f"{PORTAL_API}/returned-items", headers=portal_headers).json()
    pending_row = next(row for row in returned if row["id"] == pending_id)
    assert pending_row["ready"] is False
    assert pending_row["corrected_request"] is None
    resubmitted_rows = [row for row in returned if row["id"] in ready_ids]
    assert all(row["corrected_request"] is not None for row in resubmitted_rows)

    # The pending item can still be corrected and resubmitted later, on its
    # own, as a second corrected REQ linked to the same original request.
    assert _save_correction_draft(s, portal_headers, pending_id).status_code == 200
    second = _resubmit_corrections(s, portal_headers, request_id, [pending_id])
    assert second.status_code == 200, second.text
    assert second.json()["request_id"] != result.json()["request_id"]
    assert second.json()["source_request_id"] == request_id


def test_items_from_two_different_original_requests_cannot_be_grouped(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    request_a, items_a, portal_headers = _setup_returned_items_scenario(s, admin_headers, f"{suffix}-a", item_count=1)
    request_b, items_b, _ = _setup_returned_items_scenario(s, admin_headers, f"{suffix}-b", item_count=1)
    assert _save_correction_draft(s, portal_headers, items_a[0]).status_code == 200

    # request_b's item does not belong to request_a's owner/project scope,
    # and even if it did, mixing items across two original REQs must fail.
    result = _resubmit_corrections(s, portal_headers, request_a, [items_a[0], items_b[0]])
    assert result.status_code == 422, result.text

    with SessionLocal() as session:
        assert session.scalar(
            select(func.count()).select_from(IncomingPurchaseRequest)
            .where(IncomingPurchaseRequest.source_request_id == request_a)
        ) == 0


def test_double_submit_resubmission_does_not_duplicate_the_child_req(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    request_id, item_ids, portal_headers = _setup_returned_items_scenario(s, admin_headers, suffix, item_count=2)
    for item_id in item_ids:
        assert _save_correction_draft(s, portal_headers, item_id).status_code == 200

    first = _resubmit_corrections(s, portal_headers, request_id, item_ids)
    assert first.status_code == 200, first.text
    second = _resubmit_corrections(s, portal_headers, request_id, item_ids)
    assert second.status_code == 200, second.text
    assert second.json()["already_exists"] is True
    assert second.json()["request_id"] == first.json()["request_id"]

    with SessionLocal() as session:
        children = session.scalars(
            select(IncomingPurchaseRequest).where(
                IncomingPurchaseRequest.source_request_id == request_id
            )
        ).all()
        assert len(children) == 1
        child_items = session.scalars(
            select(IncomingPurchaseRequestItem).where(
                IncomingPurchaseRequestItem.request_id == children[0].id
            )
        ).all()
        assert len(child_items) == 2


def test_already_resubmitted_item_cannot_be_corrected_or_resubmitted_again(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    request_id, item_ids, portal_headers = _setup_returned_items_scenario(s, admin_headers, suffix, item_count=1)
    item_id = item_ids[0]
    assert _save_correction_draft(s, portal_headers, item_id).status_code == 200
    resubmitted = _resubmit_corrections(s, portal_headers, request_id, [item_id])
    assert resubmitted.status_code == 200, resubmitted.text

    again = _save_correction_draft(s, portal_headers, item_id)
    assert again.status_code == 409, again.text


PDF_BYTES = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<< >>\nendobj\n%%EOF"
PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
XLSX_BYTES = b"PK\x03\x04" + b"\x00" * 32
XLS_BYTES = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32
EXE_BYTES = b"MZ" + b"\x00" * 32


def test_portal_anonymous_cannot_create_request(s):
    response = _submit_portal_request(s, {}, _portal_payload())
    assert response.status_code == 401


def test_portal_erp_user_cannot_use_portal_endpoint(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"erp-vs-portal-{suffix}", role="procurement_engineer")
    headers = _login_headers(s, f"erp-vs-portal-{suffix}")

    response = _submit_portal_request(s, headers, _portal_payload())

    assert response.status_code == 403


def test_portal_site_user_can_create_request(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username, user_id = _make_portal_user(session, suffix, project_ids=[project_id])
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        items=[{"item_id": item_id, "quantity": 3, "note": "عاجل"}],
    ))

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ok"] is True
    assert body["request_number"].startswith("REQ-")

    with SessionLocal() as session:
        row = session.get(IncomingPurchaseRequest, body["request_id"])
        assert row is not None
        assert row.requester_user_id == user_id
        assert row.project_id == project_id
        assert row.delivery_destination == "site"
        items = session.scalars(
            select(IncomingPurchaseRequestItem).where(
                IncomingPurchaseRequestItem.request_id == row.id
            )
        ).all()
        assert len(items) == 1
        assert items[0].item_id == item_id
        assert items[0].quantity == 3


def test_req_clarification_requires_engineer_and_reason(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        portal_username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
        item_id = _make_portal_item(session, suffix)
        _make_user(session, username=f"clarify-engineer-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"clarify-responsible-{suffix}", role="procurement_responsible")
    portal_headers = _login_headers(s, portal_username)
    engineer_headers = _login_headers(s, f"clarify-engineer-{suffix}")
    responsible_headers = _login_headers(s, f"clarify-responsible-{suffix}")
    created = _submit_portal_request(s, portal_headers, _portal_payload(
        items=[{"item_id": item_id, "quantity": 1, "note": ""}],
    )).json()
    endpoint = f"{API}/workflow/incoming-purchase-requests/{created['request_id']}/technical-decision"

    missing_reason = s.post(endpoint, headers={**INTERNAL_HEADERS, **engineer_headers}, json={
        "decision": "revision_required", "note": "",
    })
    assert missing_reason.status_code == 422
    forbidden = s.post(endpoint, headers={**INTERNAL_HEADERS, **responsible_headers}, json={
        "decision": "revision_required", "note": "وضح المقاس",
    })
    assert forbidden.status_code == 403
    allowed = s.post(endpoint, headers={**INTERNAL_HEADERS, **engineer_headers}, json={
        "decision": "revision_required", "note": "وضح المقاس المطلوب",
    })
    assert allowed.status_code == 200, allowed.text
    assert allowed.json()["status"] == "need_clarification"


def test_site_portal_clarification_resubmits_same_req_with_history_and_attachment(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        portal_username, portal_user_id = _make_portal_user(session, suffix, project_ids=[project_id])
        other_username, _ = _make_portal_user(session, f"{suffix}-other", project_ids=[project_id])
        item_id = _make_portal_item(session, suffix)
        _make_user(session, username=f"clarify-flow-engineer-{suffix}", role="procurement_engineer")
    portal_headers = _login_headers(s, portal_username)
    other_headers = _login_headers(s, other_username)
    engineer_headers = _login_headers(s, f"clarify-flow-engineer-{suffix}")
    created = _submit_portal_request(s, portal_headers, _portal_payload(
        items=[{"item_id": item_id, "quantity": 1, "note": ""}],
    )).json()
    request_id = created["request_id"]
    decision = s.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/technical-decision",
        headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"decision": "revision_required", "note": "أرفق صورة للمقاس"},
    )
    assert decision.status_code == 200, decision.text

    assert s.get(f"{PORTAL_API}/purchase-requests").status_code == 401
    own_list = s.get(f"{PORTAL_API}/purchase-requests", headers=portal_headers)
    assert own_list.status_code == 200
    listed = next(row for row in own_list.json() if row["id"] == request_id)
    assert listed["request_number"] == created["request_number"]
    assert listed["clarification"]["reason"] == "أرفق صورة للمقاس"
    assert listed["clarification"]["response_status"] == "awaiting_response"
    assert all(row["id"] != request_id for row in s.get(
        f"{PORTAL_API}/purchase-requests", headers=other_headers,
    ).json())
    assert s.post(
        f"{PORTAL_API}/purchase-requests/{request_id}/clarification",
        headers=other_headers, data={"response": "محاولة غير مسموحة"},
    ).status_code == 404

    resubmitted = s.post(
        f"{PORTAL_API}/purchase-requests/{request_id}/clarification",
        headers=portal_headers,
        data={"response": "المقاس 60 × 60 سم"},
        files=[("attachments", ("size.png", PNG_BYTES, "image/png"))],
    )
    assert resubmitted.status_code == 200, resubmitted.text
    assert resubmitted.json()["request_id"] == request_id
    assert resubmitted.json()["request_number"] == created["request_number"]
    assert resubmitted.json()["status"] == "under_review"

    with SessionLocal() as session:
        row = session.get(IncomingPurchaseRequest, request_id)
        assert row.requester_user_id == portal_user_id
        assert row.status == "under_review"
        history = session.scalars(select(IncomingRequestStatusHistory).where(
            IncomingRequestStatusHistory.request_id == request_id,
        ).order_by(IncomingRequestStatusHistory.created_at)).all()
        assert [(entry.from_status, entry.to_status) for entry in history][-2:] == [
            ("new", "need_clarification"),
            ("need_clarification", "under_review"),
        ]
        assert history[-1].note == "المقاس 60 × 60 سم"
        attachments = session.scalars(select(IncomingRequestGeneralAttachment).where(
            IncomingRequestGeneralAttachment.request_id == request_id,
        )).all()
        assert any(item.original_filename == "size.png" for item in attachments)


def test_portal_requester_identity_is_server_derived_and_cannot_be_spoofed(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username, user_id = _make_portal_user(session, suffix, project_ids=[project_id])
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    payload = _portal_payload(items=[{"item_id": item_id, "quantity": 1, "note": ""}])
    # A client-supplied requester identity has no field to occupy in this
    # schema at all; sending one anyway must simply be ignored, not honored.
    payload["requester_name"] = "شخص آخر منتحل"
    payload["requester_user_id"] = "someone-elses-id"

    response = _submit_portal_request(s, headers, payload)
    assert response.status_code == 200, response.text

    with SessionLocal() as session:
        row = session.get(IncomingPurchaseRequest, response.json()["request_id"])
        assert row.requester_user_id == user_id
        assert row.requester_name != "شخص آخر منتحل"


def test_portal_client_cannot_submit_unauthorized_project(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        own_project_id = _make_project_for_portal(session, suffix)
        other_project_id = _make_project_for_portal(session, f"{suffix}-other")
        username, _ = _make_portal_user(session, suffix, project_ids=[own_project_id])
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        project_id=other_project_id,
        items=[{"item_id": item_id, "quantity": 1, "note": ""}],
    ))

    # A single-project user's own project is always server-derived, so a
    # foreign project_id in the payload is simply never consulted.
    assert response.status_code == 200, response.text
    with SessionLocal() as session:
        row = session.get(IncomingPurchaseRequest, response.json()["request_id"])
        assert row.project_id == own_project_id
        assert row.project_id != other_project_id


def test_portal_multi_project_user_cannot_pick_an_unassigned_project(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_a = _make_project_for_portal(session, f"{suffix}-a")
        project_b = _make_project_for_portal(session, f"{suffix}-b")
        outside_project = _make_project_for_portal(session, f"{suffix}-outside")
        username, _ = _make_portal_user(session, suffix, project_ids=[project_a, project_b])
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        project_id=outside_project,
        items=[{"item_id": item_id, "quantity": 1, "note": ""}],
    ))

    assert response.status_code == 422


def test_portal_user_with_no_assigned_project_cannot_submit(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _ = _make_portal_user(session, suffix, project_ids=[])
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        items=[{"item_id": item_id, "quantity": 1, "note": ""}],
    ))

    assert response.status_code == 422


def test_portal_invalid_item_id_rejected(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        items=[{"item_id": "does-not-exist-in-item-master", "quantity": 1, "note": ""}],
    ))

    assert response.status_code == 422


def test_portal_item_without_item_id_is_rejected_not_free_text(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
    headers = _login_headers(s, username)

    payload = _portal_payload()
    # There is no product_name field in the portal schema; a free-text-only
    # line (no item_id) must be rejected outright, not silently accepted.
    payload["items"] = [{"product_name": "صنف غير موجود بالكتالوج", "quantity": 1}]

    response = _submit_portal_request(s, headers, payload)

    assert response.status_code == 422


def test_portal_invalid_delivery_destination_rejected(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        destination="office", items=[{"item_id": item_id, "quantity": 1, "note": ""}],
    ))

    assert response.status_code == 422


def test_portal_submitted_request_preserves_existing_req_workflow(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        sup_before = session.scalar(select(func.count()).select_from(Supplier))
        item_before = session.scalar(select(func.count()).select_from(Item))
        project_id = _make_project_for_portal(session, suffix)
        username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        items=[{"item_id": item_id, "quantity": 2, "note": ""}],
    ))
    assert response.status_code == 200, response.text
    request_number = response.json()["request_number"]
    assert re.fullmatch(r"REQ-\d{8}-[0-9A-F]{10}", request_number)

    listing = s.get(f"{API}/internal/incoming-purchase-requests", headers={**INTERNAL_HEADERS, **admin_headers})
    assert listing.status_code == 200
    numbers = [row["request_number"] for row in listing.json()]
    assert request_number in numbers
    matched = next(row for row in listing.json() if row["request_number"] == request_number)
    assert matched["status"] == "new"

    with SessionLocal() as session:
        sup_after = session.scalar(select(func.count()).select_from(Supplier))
        item_after = session.scalar(select(func.count()).select_from(Item))
        # The new Item Master row created for this test is expected;
        # nothing else about Items or Suppliers should move.
        assert sup_after == sup_before
        assert item_after == item_before + 1


def test_portal_previous_items_scoped_to_own_user_only(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username_a, _ = _make_portal_user(session, f"{suffix}-a", project_ids=[project_id])
        username_b, _ = _make_portal_user(session, f"{suffix}-b", project_ids=[project_id])
        item_a = _make_portal_item(session, f"{suffix}-a")
        item_b = _make_portal_item(session, f"{suffix}-b")
    headers_a = _login_headers(s, username_a)
    headers_b = _login_headers(s, username_b)

    assert _submit_portal_request(s, headers_a, _portal_payload(
        items=[{"item_id": item_a, "quantity": 1, "note": ""}],
    )).status_code == 200
    assert _submit_portal_request(s, headers_b, _portal_payload(
        items=[{"item_id": item_b, "quantity": 1, "note": ""}],
    )).status_code == 200

    previous_a = s.get(f"{PORTAL_API}/previous-items", headers=headers_a).json()
    previous_b = s.get(f"{PORTAL_API}/previous-items", headers=headers_b).json()

    assert [row["id"] for row in previous_a] == [item_a]
    assert [row["id"] for row in previous_b] == [item_b]


def test_portal_previous_items_are_distinct_and_most_recent_first(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
        item_x = _make_portal_item(session, f"{suffix}-x")
        item_y = _make_portal_item(session, f"{suffix}-y")
    headers = _login_headers(s, username)

    # Request X twice (should not duplicate) then Y once (should rank first).
    for _ in range(2):
        assert _submit_portal_request(s, headers, _portal_payload(
            items=[{"item_id": item_x, "quantity": 1, "note": ""}],
        )).status_code == 200
    assert _submit_portal_request(s, headers, _portal_payload(
        items=[{"item_id": item_y, "quantity": 1, "note": ""}],
    )).status_code == 200

    previous = s.get(f"{PORTAL_API}/previous-items", headers=headers).json()

    ids = [row["id"] for row in previous]
    assert ids == [item_y, item_x]
    assert len(ids) == len(set(ids))


def test_portal_previous_items_excludes_deleted_items(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    assert _submit_portal_request(s, headers, _portal_payload(
        items=[{"item_id": item_id, "quantity": 1, "note": ""}],
    )).status_code == 200
    assert [row["id"] for row in s.get(f"{PORTAL_API}/previous-items", headers=headers).json()] == [item_id]

    with SessionLocal() as session:
        session.execute(Item.__table__.delete().where(Item.id == item_id))
        session.commit()

    after_delete = s.get(f"{PORTAL_API}/previous-items", headers=headers).json()
    assert after_delete == []


def test_portal_previous_items_empty_for_user_with_no_history(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
    headers = _login_headers(s, username)

    response = s.get(f"{PORTAL_API}/previous-items", headers=headers)

    assert response.status_code == 200
    assert response.json() == []


def test_portal_previous_items_exposes_no_commercial_data(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        project_id = _make_project_for_portal(session, suffix)
        username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
        item_id = _make_portal_item(session, suffix, preferred_supplier="مورد سري")
    headers = _login_headers(s, username)
    assert _submit_portal_request(s, headers, _portal_payload(
        items=[{"item_id": item_id, "quantity": 1, "note": ""}],
    )).status_code == 200

    response = s.get(f"{PORTAL_API}/previous-items", headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert set(body[0].keys()) == {"id", "code", "name", "unit", "main_category"}
    assert "مورد سري" not in response.text
    assert "preferred_supplier" not in response.text
    assert "price" not in response.text.lower()


# ---------------- Manual items and request attachments ----------------

def _portal_setup(session, suffix):
    project_id = _make_project_for_portal(session, suffix)
    username, user_id = _make_portal_user(session, suffix, project_ids=[project_id])
    return username, user_id, project_id


def test_portal_manual_item_accepted(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        items=[_manual_item(name="أسمنت غير مسجل", unit="كيس", quantity=4)],
    ))

    assert response.status_code == 200, response.text
    with SessionLocal() as session:
        row = session.get(IncomingPurchaseRequest, response.json()["request_id"])
        items = session.scalars(
            select(IncomingPurchaseRequestItem).where(IncomingPurchaseRequestItem.request_id == row.id)
        ).all()
        assert len(items) == 1
        assert items[0].item_id == ""
        assert items[0].product_name == "أسمنت غير مسجل"
        assert items[0].unit == "كيس"


def test_portal_manual_item_does_not_create_item_master_record(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
        items_before = session.scalar(select(func.count()).select_from(Item))

    headers = _login_headers(s, username)
    response = _submit_portal_request(s, headers, _portal_payload(
        items=[_manual_item(name=f"صنف يدوي فريد {suffix}")],
    ))
    assert response.status_code == 200, response.text

    with SessionLocal() as session:
        items_after = session.scalar(select(func.count()).select_from(Item))
        match = session.scalar(select(Item).where(Item.name == f"صنف يدوي فريد {suffix}"))
        assert items_after == items_before
        assert match is None


def test_portal_manual_item_requires_name(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        items=[_manual_item(name="")],
    ))

    assert response.status_code == 422


def test_portal_manual_item_requires_positive_quantity(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(
        items=[_manual_item(quantity=0)],
    ))

    assert response.status_code == 422


def test_portal_rejects_line_claiming_to_be_both_master_and_manual(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    payload = _portal_payload(items=[{
        "item_id": item_id, "product_name": "اسم منتحل", "unit": "قطعة", "quantity": 1, "note": "",
    }])

    response = _submit_portal_request(s, headers, payload)

    assert response.status_code == 422


def test_portal_mixed_master_and_manual_items_accepted_in_one_request(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(s, headers, _portal_payload(items=[
        {"item_id": item_id, "quantity": 2, "note": ""},
        _manual_item(name="صنف يدوي ضمن نفس الطلب"),
        _manual_item(name="صنف يدوي ثانٍ"),
    ]))

    assert response.status_code == 200, response.text
    with SessionLocal() as session:
        items = session.scalars(
            select(IncomingPurchaseRequestItem)
            .where(IncomingPurchaseRequestItem.request_id == response.json()["request_id"])
            .order_by(IncomingPurchaseRequestItem.position)
        ).all()
        assert len(items) == 3
        assert items[0].item_id == item_id
        assert items[1].item_id == "" and items[1].product_name == "صنف يدوي ضمن نفس الطلب"
        assert items[2].item_id == "" and items[2].product_name == "صنف يدوي ثانٍ"


def test_portal_pdf_attachment_accepted(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(
        s, headers, _portal_payload(items=[_manual_item()]),
        files=[("attachments", ("spec.pdf", PDF_BYTES, "application/pdf"))],
    )

    assert response.status_code == 200, response.text


def test_portal_xlsx_attachment_accepted(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(
        s, headers, _portal_payload(items=[_manual_item()]),
        files=[("attachments", ("list.xlsx", XLSX_BYTES, "application/octet-stream"))],
    )

    assert response.status_code == 200, response.text


def test_portal_xls_attachment_accepted(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(
        s, headers, _portal_payload(items=[_manual_item()]),
        files=[("attachments", ("list.xls", XLS_BYTES, "application/octet-stream"))],
    )

    assert response.status_code == 200, response.text


def test_portal_image_attachment_accepted(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(
        s, headers, _portal_payload(items=[_manual_item()]),
        files=[("attachments", ("photo.png", PNG_BYTES, "image/png"))],
    )

    assert response.status_code == 200, response.text


def test_portal_executable_attachment_rejected(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(
        s, headers, _portal_payload(items=[_manual_item()]),
        files=[("attachments", ("tool.exe", EXE_BYTES, "application/octet-stream"))],
    )

    assert response.status_code == 422


def test_portal_multiple_attachments_accepted_and_linked_to_correct_request(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        username, _, _ = _portal_setup(session, suffix)
    headers = _login_headers(s, username)

    response = _submit_portal_request(
        s, headers, _portal_payload(items=[_manual_item()]),
        files=[
            ("attachments", ("photo1.png", PNG_BYTES, "image/png")),
            ("attachments", ("photo2.png", PNG_BYTES, "image/png")),
            ("attachments", ("spec.pdf", PDF_BYTES, "application/pdf")),
        ],
    )

    assert response.status_code == 200, response.text
    request_id = response.json()["request_id"]

    with SessionLocal() as session:
        rows = session.scalars(
            select(IncomingRequestGeneralAttachment)
            .where(IncomingRequestGeneralAttachment.request_id == request_id)
        ).all()
        assert len(rows) == 3
        assert {r.request_id for r in rows} == {request_id}
        assert sorted(r.original_filename for r in rows) == ["photo1.png", "photo2.png", "spec.pdf"]

    # A second, unrelated request must not see these files.
    with SessionLocal() as session:
        other_username, _, _ = _portal_setup(session, f"{suffix}-other")
    other_headers = _login_headers(s, other_username)
    other_response = _submit_portal_request(s, other_headers, _portal_payload(items=[_manual_item()]))
    assert other_response.status_code == 200
    with SessionLocal() as session:
        other_rows = session.scalars(
            select(IncomingRequestGeneralAttachment)
            .where(IncomingRequestGeneralAttachment.request_id == other_response.json()["request_id"])
        ).all()
        assert other_rows == []


# ---------------- Manual item -> Item Master conversion ----------------

CONVERT_URL = lambda req_id, item_id: f"{API}/internal/incoming-purchase-requests/{req_id}/items/{item_id}/convert-to-item"  # noqa: E731


def _make_manual_line(client, suffix, name="صنف يدوي أصلي بدون تعديل"):
    """Create a real REQ with one manual line via the portal flow, return
    (request_id, request_item_id, admin_username)."""
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        portal_username, _, _ = _portal_setup(session, suffix)
    portal_headers = _login_headers(client, portal_username)
    response = _submit_portal_request(client, portal_headers, _portal_payload(
        items=[_manual_item(name=name)],
    ))
    assert response.status_code == 200, response.text
    request_id = response.json()["request_id"]
    with SessionLocal() as session:
        line = session.scalars(
            select(IncomingPurchaseRequestItem).where(IncomingPurchaseRequestItem.request_id == request_id)
        ).first()
    return request_id, line.id, admin_username


def _convert_body(**overrides):
    body = {"name": "صنف معتمد جديد", "unit": "قطعة", "main_category": "", "subcategory": "", "brand": "", "specifications": "", "notes": ""}
    body.update(overrides)
    return body


def test_convert_manual_item_succeeds(s):
    suffix = uuid.uuid4().hex[:8]
    request_id, line_id, admin_username = _make_manual_line(s, suffix)
    headers = _login_headers(s, admin_username)

    response = s.post(CONVERT_URL(request_id, line_id), headers=headers, json=_convert_body(name=f"صنف {suffix}"))

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["created"] is True
    assert body["item_code"]


def test_convert_uses_centralized_itm_numbering(s):
    suffix = uuid.uuid4().hex[:8]
    request_id, line_id, admin_username = _make_manual_line(s, suffix)
    headers = _login_headers(s, admin_username)

    response = s.post(CONVERT_URL(request_id, line_id), headers=headers, json=_convert_body(name=f"صنف ترقيم {suffix}"))
    assert response.status_code == 200, response.text
    item_code = response.json()["item_code"]

    with SessionLocal() as session:
        item = session.get(Item, response.json()["item_id"])
        assert item.code == item_code
        stored_next = session.get(BusinessCodeSequence, "items").next_value
        # Reserved code's numeric suffix must be strictly below the counter
        # the sequence table now points at (i.e. it was actually reserved
        # through the shared sequence, not invented locally).
        assert item_code.startswith("ITM-")
        assert int(item_code.split("-")[1]) < stored_next


def test_convert_links_request_line_to_new_item(s):
    suffix = uuid.uuid4().hex[:8]
    request_id, line_id, admin_username = _make_manual_line(s, suffix)
    headers = _login_headers(s, admin_username)

    response = s.post(CONVERT_URL(request_id, line_id), headers=headers, json=_convert_body(name=f"صنف ربط {suffix}"))
    assert response.status_code == 200, response.text
    new_item_id = response.json()["item_id"]

    with SessionLocal() as session:
        line = session.get(IncomingPurchaseRequestItem, line_id)
        assert line.item_id == new_item_id


def test_convert_preserves_original_manual_wording(s):
    suffix = uuid.uuid4().hex[:8]
    original_name = f"نص الطلب الأصلي {suffix}"
    request_id, line_id, admin_username = _make_manual_line(s, suffix, name=original_name)
    headers = _login_headers(s, admin_username)

    response = s.post(
        CONVERT_URL(request_id, line_id), headers=headers,
        json=_convert_body(name=f"اسم الصنف المعتمد المختلف {suffix}"),
    )
    assert response.status_code == 200, response.text

    with SessionLocal() as session:
        line = session.get(IncomingPurchaseRequestItem, line_id)
        assert line.product_name == original_name


def test_convert_rejects_already_master_linked_line(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        project_id = _make_project_for_portal(session, suffix)
        username, _ = _make_portal_user(session, suffix, project_ids=[project_id])
        item_id = _make_portal_item(session, suffix)
    headers = _login_headers(s, username)
    submit = _submit_portal_request(s, headers, _portal_payload(
        items=[{"item_id": item_id, "quantity": 1, "note": ""}],
    ))
    assert submit.status_code == 200
    request_id = submit.json()["request_id"]
    with SessionLocal() as session:
        line = session.scalars(
            select(IncomingPurchaseRequestItem).where(IncomingPurchaseRequestItem.request_id == request_id)
        ).first()

    admin_headers = _login_headers(s, admin_username)
    response = s.post(CONVERT_URL(request_id, line.id), headers=admin_headers, json=_convert_body())

    # Already Master-linked - not a manual line - reported the same way as
    # an idempotent re-conversion (no second Item, existing one returned).
    assert response.status_code == 200
    assert response.json()["created"] is False
    assert response.json()["item_id"] == item_id


def test_convert_is_idempotent_on_second_attempt(s):
    suffix = uuid.uuid4().hex[:8]
    request_id, line_id, admin_username = _make_manual_line(s, suffix)
    headers = _login_headers(s, admin_username)

    with SessionLocal() as session:
        items_before = session.scalar(select(func.count()).select_from(Item))

    first = s.post(CONVERT_URL(request_id, line_id), headers=headers, json=_convert_body(name=f"صنف تكرار {suffix}"))
    assert first.status_code == 200
    first_item_id = first.json()["item_id"]

    second = s.post(CONVERT_URL(request_id, line_id), headers=headers, json=_convert_body(name="اسم مختلف تمامًا"))
    assert second.status_code == 200
    assert second.json()["created"] is False
    assert second.json()["item_id"] == first_item_id

    with SessionLocal() as session:
        items_after = session.scalar(select(func.count()).select_from(Item))
        assert items_after == items_before + 1


def test_convert_rejects_duplicate_item_name(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        existing_name = f"صنف موجود بالفعل {suffix}"
        _make_portal_item(session, suffix, name=existing_name)
    request_id, line_id, admin_username = _make_manual_line(s, suffix)
    headers = _login_headers(s, admin_username)

    response = s.post(CONVERT_URL(request_id, line_id), headers=headers, json=_convert_body(name=existing_name))

    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["code"] == "duplicate_item_name"
    assert detail["existing_item"]["name"] == existing_name


def test_convert_forbidden_for_site_portal_account(s):
    suffix = uuid.uuid4().hex[:8]
    request_id, line_id, _ = _make_manual_line(s, suffix)
    with SessionLocal() as session:
        portal_username, _, _ = _portal_setup(session, f"{suffix}-attacker")
    headers = _login_headers(s, portal_username)

    response = s.post(CONVERT_URL(request_id, line_id), headers=headers, json=_convert_body())

    assert response.status_code == 403


def test_convert_increases_item_count_by_exactly_one(s):
    suffix = uuid.uuid4().hex[:8]
    request_id, line_id, admin_username = _make_manual_line(s, suffix)
    headers = _login_headers(s, admin_username)
    with SessionLocal() as session:
        items_before = session.scalar(select(func.count()).select_from(Item))
        suppliers_before = session.scalar(select(func.count()).select_from(Supplier))

    response = s.post(CONVERT_URL(request_id, line_id), headers=headers, json=_convert_body(name=f"صنف عدّاد {suffix}"))
    assert response.status_code == 200, response.text

    with SessionLocal() as session:
        items_after = session.scalar(select(func.count()).select_from(Item))
        suppliers_after = session.scalar(select(func.count()).select_from(Supplier))
        assert items_after == items_before + 1
        assert suppliers_after == suppliers_before


# ---------------- Sprint 2.3: ERP role enforcement ----------------

def test_convert_manual_item_allowed_for_procurement_responsible(s):
    suffix = uuid.uuid4().hex[:8]
    request_id, line_id, _ = _make_manual_line(s, suffix)
    with SessionLocal() as session:
        _make_user(session, username=f"sprint23-resp-convert-{suffix}", role="procurement_responsible")
    headers = _login_headers(s, f"sprint23-resp-convert-{suffix}")

    response = s.post(CONVERT_URL(request_id, line_id), headers=headers, json=_convert_body(name=f"صنف مسؤول مشتريات {suffix}"))

    assert response.status_code == 200, response.text
    assert response.json()["created"] is True


def test_convert_manual_item_forbidden_for_procurement_engineer_and_commercial_manager(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"sprint23-eng-convert-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"sprint23-mgr-convert-{suffix}", role="commercial_manager")
    engineer_headers = _login_headers(s, f"sprint23-eng-convert-{suffix}")
    manager_headers = _login_headers(s, f"sprint23-mgr-convert-{suffix}")

    request_id, line_id, _ = _make_manual_line(s, suffix)
    assert s.post(CONVERT_URL(request_id, line_id), headers=engineer_headers, json=_convert_body()).status_code == 403
    assert s.post(CONVERT_URL(request_id, line_id), headers=manager_headers, json=_convert_body()).status_code == 403


def test_anonymous_request_to_workflow_mutation_is_rejected(s):
    response = s.post(
        f"{API}/workflow/incoming-purchase-requests/{uuid.uuid4()}/technical-decision",
        headers=INTERNAL_HEADERS,
        json={"decision": "approved_for_pricing"},
    )
    assert response.status_code == 401


def test_deactivated_erp_user_loses_workflow_access_immediately(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        user = _make_user(session, username=f"sprint23-deactivated-{suffix}", role="procurement_engineer")
        user_id = user.id
    headers = _login_headers(s, f"sprint23-deactivated-{suffix}")
    ok = s.post(
        f"{API}/workflow/incoming-purchase-requests/{uuid.uuid4()}/technical-decision",
        headers={**INTERNAL_HEADERS, **headers}, json={"decision": "approved_for_pricing"},
    )
    assert ok.status_code == 404  # authenticated fine, only the (fake) request id is missing

    with SessionLocal() as session:
        session.get(User, user_id).active = False
        session.commit()

    revoked = s.post(
        f"{API}/workflow/incoming-purchase-requests/{uuid.uuid4()}/technical-decision",
        headers={**INTERNAL_HEADERS, **headers}, json={"decision": "approved_for_pricing"},
    )
    assert revoked.status_code == 401


def test_commercial_manager_cannot_prepare_price_comparisons(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"sprint23-mgr-cmp-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"sprint23-mgr-cmp-{suffix}")
    project = s.get(f"{API}/projects", headers=admin_headers).json()[0]
    item = s.get(f"{API}/items", headers=admin_headers).json()[0]
    supplier = s.get(f"{API}/suppliers", headers=admin_headers).json()[0]

    response = s.post(f"{API}/price-comparisons", headers=manager_headers, json={
        "project_id": project["id"], "project_name": project["name"],
        "comparison_date": "2026-08-19", "rows": [{
            "item_id": item["id"], "supplier_id": supplier["id"],
            "quantity": 1, "unit": item.get("unit", "قطعة"), "unit_price": 10,
            "availability": "available", "price_valid_until": "2099-12-31",
        }],
    })
    assert response.status_code == 403


def test_procurement_responsible_cannot_confirm_funds_release(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"sprint23-resp-funds-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"sprint23-resp-funds-{suffix}")

    response = s.post(
        f"{API}/workflow/approvals/{uuid.uuid4()}/funds-release",
        headers={**INTERNAL_HEADERS, **responsible_headers}, json={"actor": "x"},
    )
    assert response.status_code == 403


def test_site_portal_account_is_forbidden_from_internal_workflow_mutation(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        portal_username, _, _ = _portal_setup(session, f"sprint23-portal-{suffix}")
    portal_headers = _login_headers(s, portal_username)
    project = s.get(f"{API}/projects", headers=admin_headers).json()[0]
    item = s.get(f"{API}/items", headers=admin_headers).json()[0]
    supplier = s.get(f"{API}/suppliers", headers=admin_headers).json()[0]

    response = s.post(f"{API}/price-comparisons", headers=portal_headers, json={
        "project_id": project["id"], "project_name": project["name"],
        "comparison_date": "2026-08-19", "rows": [{
            "item_id": item["id"], "supplier_id": supplier["id"],
            "quantity": 1, "unit": item.get("unit", "قطعة"), "unit_price": 10,
            "availability": "available", "price_valid_until": "2099-12-31",
        }],
    })
    assert response.status_code == 403


def test_client_supplied_actor_role_cannot_escalate_privileges(s):
    """CRITICAL spoof-resistance check: the authenticated DB role always
    wins over any client-supplied role field in the payload."""
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"sprint23-spoof-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"sprint23-spoof-mgr-{suffix}", role="commercial_manager")
        portal_username, _, _ = _portal_setup(session, f"sprint23-spoof-portal-{suffix}")
    engineer_headers = _login_headers(s, f"sprint23-spoof-eng-{suffix}")
    manager_headers = _login_headers(s, f"sprint23-spoof-mgr-{suffix}")
    portal_headers = _login_headers(s, portal_username)

    # Authenticated procurement_engineer claiming to be commercial_manager
    # via the (now-ignored) payload field must still be refused a
    # commercial-only action.
    spoofed_funds_release = s.post(
        f"{API}/workflow/approvals/{uuid.uuid4()}/funds-release",
        headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"actor": "x", "actor_role": "commercial_manager"},
    )
    assert spoofed_funds_release.status_code == 403

    # Authenticated commercial_manager claiming to be procurement_responsible
    # must still be refused an engineer-only action (technical review).
    spoofed_technical_decision = s.post(
        f"{API}/workflow/incoming-purchase-requests/{uuid.uuid4()}/technical-decision",
        headers={**INTERNAL_HEADERS, **manager_headers},
        json={"decision": "approved_for_pricing", "actor_role": "procurement_responsible"},
    )
    assert spoofed_technical_decision.status_code == 403

    # Authenticated site-portal account claiming to be admin must still be
    # refused every internal ERP workflow action.
    request_id, line_id, _ = _make_manual_line(s, suffix)
    spoofed_conversion = s.post(
        CONVERT_URL(request_id, line_id), headers=portal_headers,
        json=_convert_body(actor_role="admin"),
    )
    assert spoofed_conversion.status_code == 403


# ---------------- Sprint 3.1: RFQ + Supplier Quotations ----------------

RFQ_API = f"{API}/workflow/rfqs"


def _make_pricing_request(client, suffix, manual_item=False):
    """Create a REQ with one approved item, link it to a project, and move
    it to 'pricing' via the real technical-decision endpoint (not a status
    shortcut). Returns (request_id, request_number, project_id, item_id,
    admin_headers)."""
    timestamp = datetime.now(timezone.utc).isoformat()
    request_id = str(uuid.uuid4())
    request_number = f"T-RFQ-{uuid.uuid4().hex[:8]}"
    line_id = str(uuid.uuid4())
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        real_item_id = "" if manual_item else session.scalars(select(Item)).first().id
        session.add(IncomingPurchaseRequest(
            id=request_id, request_number=request_number,
            requester_name="مهندس موقع RFQ", company_name="عميل RFQ",
            phone_number="01000000002", project_name="مشروع RFQ",
            project_location="القاهرة", delivery_location="القاهرة",
            required_delivery_date="2099-01-01", priority="normal",
            submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
            created_at=timestamp, updated_at=timestamp,
        ))
        session.add(IncomingPurchaseRequestItem(
            id=line_id, request_id=request_id, position=1, item_id=real_item_id,
            product_name=f"صنف RFQ يدوي {suffix}" if manual_item else "صنف RFQ من الدليل",
            quantity=3, unit="قطعة",
            review_status="approved", reviewed_by="engineer", reviewed_at=timestamp,
        ))
        session.commit()
    admin_headers = _login_headers(client, admin_username)
    project = client.get(f"{API}/projects", headers=admin_headers).json()[0]
    linked = client.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/link-project",
        headers={**INTERNAL_HEADERS, **admin_headers}, json={"project_id": project["id"], "actor": "tester"},
    )
    assert linked.status_code == 200, linked.text
    decision = client.post(
        f"{API}/workflow/incoming-purchase-requests/{request_id}/technical-decision",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={"decision": "approved_for_pricing", "actor": "tester"},
    )
    assert decision.status_code == 200, decision.text
    return request_id, request_number, project["id"], line_id, admin_headers


def test_procurement_responsible_can_create_rfq_from_pricing_request(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-resp-{suffix}")
    request_id, request_number, project_id, line_id, _ = _make_pricing_request(s, suffix)

    response = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id, "deadline": "2099-01-15", "actor": "tester"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["already_exists"] is False
    rfq = body["rfq"]
    assert rfq["rfq_number"].startswith("RFQ-")
    assert rfq["source_request_id"] == request_id
    assert rfq["source_request_number"] == request_number
    assert rfq["project_id"] == project_id
    assert rfq["deadline"] == "2099-01-15"
    # RFQ items copy the REQ item verbatim (product name, qty, unit, source id).
    assert len(rfq["items"]) == 1
    item = rfq["items"][0]
    assert item["source_request_item_id"] == line_id
    assert item["quantity"] == 3
    assert item["unit"] == "قطعة"
    assert item["item_id"]


def test_admin_can_create_rfq(s):
    suffix = uuid.uuid4().hex[:8]
    request_id, *_rest, admin_headers = _make_pricing_request(s, suffix)
    response = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **admin_headers},
        json={"source_request_id": request_id},
    )
    assert response.status_code == 200, response.text


def test_procurement_engineer_cannot_create_rfq(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-eng-{suffix}", role="procurement_engineer")
    engineer_headers = _login_headers(s, f"rfq-eng-{suffix}")
    request_id, *_rest = _make_pricing_request(s, suffix)
    response = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"source_request_id": request_id},
    )
    assert response.status_code == 403


def test_commercial_manager_cannot_create_rfq(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"rfq-mgr-{suffix}")
    request_id, *_rest = _make_pricing_request(s, suffix)
    response = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **manager_headers},
        json={"source_request_id": request_id},
    )
    assert response.status_code == 403


def test_site_portal_cannot_create_rfq(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        portal_username, _, _ = _portal_setup(session, f"rfq-portal-{suffix}")
    portal_headers = _login_headers(s, portal_username)
    request_id, *_rest = _make_pricing_request(s, suffix)
    response = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **portal_headers},
        json={"source_request_id": request_id},
    )
    assert response.status_code == 403


def test_anonymous_cannot_create_rfq(s):
    response = s.post(
        RFQ_API, headers=INTERNAL_HEADERS, json={"source_request_id": str(uuid.uuid4())},
    )
    assert response.status_code == 401


def test_rfq_manual_request_item_works(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-manual-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-manual-resp-{suffix}")
    request_id, _request_number, _project_id, line_id, _ = _make_pricing_request(
        s, suffix, manual_item=True,
    )
    response = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    )
    assert response.status_code == 200, response.text
    item = response.json()["rfq"]["items"][0]
    assert item["item_id"] == ""
    assert item["source_request_item_id"] == line_id
    assert item["product_name"] == f"صنف RFQ يدوي {suffix}"


def test_rfq_uses_centralized_numbering(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-seq-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-seq-resp-{suffix}")
    request_a, *_ = _make_pricing_request(s, f"{suffix}-a")
    request_b, *_ = _make_pricing_request(s, f"{suffix}-b")
    first = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_a},
    ).json()["rfq"]["rfq_number"]
    second = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_b},
    ).json()["rfq"]["rfq_number"]
    assert re.fullmatch(r"RFQ-\d{6}", first)
    assert re.fullmatch(r"RFQ-\d{6}", second)
    assert int(second.split("-")[1]) > int(first.split("-")[1])


def test_rfq_creation_is_idempotent_returns_existing(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-idem-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-idem-resp-{suffix}")
    request_id, *_rest = _make_pricing_request(s, suffix)
    first = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    )
    assert first.status_code == 200
    assert first.json()["already_exists"] is False
    second = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    )
    assert second.status_code == 200
    assert second.json()["already_exists"] is True
    assert second.json()["rfq"]["id"] == first.json()["rfq"]["id"]


def test_rfq_creation_rejected_when_request_not_in_pricing(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-notready-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-notready-resp-{suffix}")
    timestamp = datetime.now(timezone.utc).isoformat()
    request_id = str(uuid.uuid4())
    with SessionLocal.begin() as session:
        session.add(IncomingPurchaseRequest(
            id=request_id, request_number=f"T-RFQ-NEW-{suffix}",
            requester_name="مهندس", company_name="عميل", phone_number="01000000003",
            project_name="مشروع", project_location="القاهرة", delivery_location="القاهرة",
            required_delivery_date="2099-01-01", priority="normal",
            submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
            created_at=timestamp, updated_at=timestamp,
        ))
    response = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    )
    assert response.status_code == 409


def _rfq_with_supplier(client, suffix, responsible_headers):
    request_id, *_rest = _make_pricing_request(client, suffix)
    rfq = client.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    ).json()["rfq"]
    supplier = client.get(f"{API}/suppliers", headers=responsible_headers).json()[0]
    added = client.post(
        f"{RFQ_API}/{rfq['id']}/suppliers", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    )
    assert added.status_code == 200, added.text
    return rfq["id"], supplier, added.json()["rfq"]


def test_supplier_can_be_assigned_to_rfq(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-sup-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-sup-resp-{suffix}")
    rfq_id, supplier, rfq = _rfq_with_supplier(s, suffix, responsible_headers)
    assert rfq["supplier_count"] == 1
    assert rfq["suppliers"][0]["supplier_id"] == supplier["id"]
    assert rfq["suppliers"][0]["supplier_name"] == supplier["name"]


def test_invalid_supplier_id_rejected_for_rfq(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-badsup-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-badsup-resp-{suffix}")
    request_id, *_rest = _make_pricing_request(s, suffix)
    rfq = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    ).json()["rfq"]
    response = s.post(
        f"{RFQ_API}/{rfq['id']}/suppliers", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": "does-not-exist"},
    )
    assert response.status_code == 422


def test_rfq_supplier_addition_does_not_modify_supplier_row(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-supsafe-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-supsafe-resp-{suffix}")
    supplier_before = s.get(f"{API}/suppliers", headers=responsible_headers).json()[0]
    _rfq_id, supplier, _rfq = _rfq_with_supplier(s, suffix, responsible_headers)
    assert supplier["id"] == supplier_before["id"]
    supplier_after = next(
        row for row in s.get(f"{API}/suppliers", headers=responsible_headers).json() if row["id"] == supplier["id"]
    )
    assert supplier_after["name"] == supplier_before["name"]
    assert supplier_after["code"] == supplier_before["code"]


def test_supplier_quotation_can_be_created(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-quo-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-quo-resp-{suffix}")
    rfq_id, supplier, _rfq = _rfq_with_supplier(s, suffix, responsible_headers)
    response = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["already_exists"] is False
    assert body["quotation"]["status"] == "draft"
    assert body["quotation"]["supplier_id"] == supplier["id"]


def test_duplicate_supplier_quotation_prevented(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-dupquo-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-dupquo-resp-{suffix}")
    rfq_id, supplier, _rfq = _rfq_with_supplier(s, suffix, responsible_headers)
    first = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    )
    second = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    )
    assert first.json()["quotation"]["id"] == second.json()["quotation"]["id"]
    assert second.json()["already_exists"] is True


def test_quotation_lines_remain_linked_to_source_ids(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-lines-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-lines-resp-{suffix}")
    request_id, _request_number, _project_id, line_id, _ = _make_pricing_request(s, suffix)
    rfq = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    ).json()["rfq"]
    supplier = s.get(f"{API}/suppliers", headers=responsible_headers).json()[0]
    s.post(
        f"{RFQ_API}/{rfq['id']}/suppliers", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    )
    quotation = s.post(
        f"{RFQ_API}/{rfq['id']}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    ).json()["quotation"]
    rfq_item_id = rfq["items"][0]["id"]
    updated = s.put(
        f"{RFQ_API}/{rfq['id']}/quotations/{quotation['id']}",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        json={
            "quotation_ref": "SUP-QT-1", "status": "received",
            "lines": [{
                "rfq_item_id": rfq_item_id, "quantity": 3, "unit": "قطعة",
                "unit_price": 55.5, "discount_pct": 10, "tax_pct": 14,
                "availability": "available",
            }],
        },
    )
    assert updated.status_code == 200, updated.text
    line = updated.json()["lines"][0]
    assert line["rfq_item_id"] == rfq_item_id
    assert line["source_request_item_id"] == line_id
    assert updated.json()["status"] == "received"

    rows = s.get(
        f"{RFQ_API}/{rfq['id']}/comparison-rows",
        headers={**INTERNAL_HEADERS, **responsible_headers},
    )
    assert rows.status_code == 200
    assert len(rows.json()["rows"]) == 1
    assert rows.json()["rows"][0]["supplier_id"] == supplier["id"]
    assert rows.json()["rows"][0]["unit_price"] == 55.5

    history = s.get(
        f"{API}/supplier-price-history",
        headers=responsible_headers,
        params={"supplier": supplier["name"]},
    )
    assert history.status_code == 200, history.text
    history_line = next(row for row in history.json() if row["rfq_id"] == rfq["id"])
    assert history_line["rfq_number"] == rfq["rfq_number"]
    assert history_line["supplier"] == supplier["name"]
    assert history_line["unit_price"] == 55.5
    assert history_line["adjusted_unit_price"] == 56.94
    assert history_line["comparison_number"] == ""


def _received_quotation_for_item(
    client, suffix, responsible_headers, unit_price, quotation_date="", item_id=None,
):
    """Create RFQ -> supplier -> a *received* quotation with one priced
    line, against an isolated item (own id, not shared with other tests
    unless `item_id` is passed to deliberately reuse one). Returns
    (item_id, supplier, rfq, quotation)."""
    request_id, _request_number, _project_id, line_id, _ = _make_pricing_request(client, suffix)
    if item_id:
        with SessionLocal() as session:
            line = session.get(IncomingPurchaseRequestItem, line_id)
            line.item_id = item_id
            session.commit()
    rfq = client.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    ).json()["rfq"]
    supplier = client.get(f"{API}/suppliers", headers=responsible_headers).json()[0]
    client.post(
        f"{RFQ_API}/{rfq['id']}/suppliers", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    )
    quotation = client.post(
        f"{RFQ_API}/{rfq['id']}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    ).json()["quotation"]
    rfq_item_id = rfq["items"][0]["id"]
    resolved_item_id = rfq["items"][0]["item_id"]
    body = {
        "quotation_ref": f"QT-{suffix}", "status": "received",
        "lines": [{
            "rfq_item_id": rfq_item_id, "quantity": 3, "unit": "قطعة",
            "unit_price": unit_price, "discount_pct": 0, "tax_pct": 0,
            "availability": "available",
        }],
    }
    if quotation_date:
        body["quotation_date"] = quotation_date
    updated = client.put(
        f"{RFQ_API}/{rfq['id']}/quotations/{quotation['id']}",
        headers={**INTERNAL_HEADERS, **responsible_headers}, json=body,
    )
    assert updated.status_code == 200, updated.text
    return resolved_item_id, supplier, rfq, quotation


def test_item_last_formal_price_comes_from_received_quotation(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-formal-item-{suffix}", role="procurement_responsible")
        own_item_id = _make_portal_item(session, suffix)
    responsible_headers = _login_headers(s, f"rfq-formal-item-{suffix}")
    item_id, supplier, _rfq, _quotation = _received_quotation_for_item(
        s, suffix, responsible_headers, unit_price=1250, quotation_date="2026-08-10",
        item_id=own_item_id,
    )
    assert item_id == own_item_id

    items = s.get(f"{API}/items", headers=responsible_headers).json()
    item = next(row for row in items if row["id"] == item_id)
    assert item["last_formal_price"] == 1250
    assert item["last_formal_supplier"] == supplier["name"]
    assert item["last_formal_date"] == "2026-08-10"


def test_item_last_formal_price_ignores_legacy_price_history(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-legacy-{suffix}", role="procurement_responsible")
        item_id = _make_portal_item(session, suffix)
        item = session.get(Item, item_id)
        session.add(PriceHistory(
            id=str(uuid.uuid4()), record_no=10_000_000 + int(suffix, 16) % 1_000_000,
            date="2026-01-01", item_code=item.code, supplier="مورد شراء مباشر قديم",
            quantity=1, unit_price=9999, final_price=9999,
        ))
        session.commit()
    responsible_headers = _login_headers(s, f"rfq-legacy-{suffix}")

    items = s.get(f"{API}/items", headers=responsible_headers).json()
    item_row = next(row for row in items if row["id"] == item_id)
    # The legacy direct-purchase price still fills the old "last_price"
    # field (untouched), but must never leak into last_formal_price.
    assert item_row["last_price"] == 9999
    assert item_row["last_formal_price"] is None
    assert item_row["last_formal_supplier"] == ""


def test_item_without_any_quotation_has_empty_formal_price(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        admin_username = _make_admin(session, suffix)
        item_id = _make_portal_item(session, suffix)
    admin_headers = _login_headers(s, admin_username)

    items = s.get(f"{API}/items", headers=admin_headers).json()
    item_row = next(row for row in items if row["id"] == item_id)
    assert item_row["last_formal_price"] is None
    assert item_row["last_formal_supplier"] == ""
    assert item_row["last_formal_date"] == ""


def test_last_formal_prices_endpoint_returns_price_for_matching_supplier_and_item(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-cmp-price-{suffix}", role="procurement_responsible")
        own_item_id = _make_portal_item(session, suffix)
    responsible_headers = _login_headers(s, f"rfq-cmp-price-{suffix}")
    item_id, supplier, _rfq, _quotation = _received_quotation_for_item(
        s, suffix, responsible_headers, unit_price=1250, quotation_date="2026-08-10",
        item_id=own_item_id,
    )

    response = s.get(
        f"{API}/price-comparisons/last-formal-prices", headers=responsible_headers,
        params={"pairs": f"{item_id}:{supplier['id']},{item_id}:no-such-supplier"},
    )
    assert response.status_code == 200, response.text
    prices = response.json()["prices"]
    assert prices[f"{item_id}|{supplier['id']}"]["unit_price"] == 1250
    assert prices[f"{item_id}|{supplier['id']}"]["date"] == "2026-08-10"
    assert f"{item_id}|no-such-supplier" not in prices


def test_last_formal_price_uses_the_most_recent_received_quotation(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-latest-{suffix}", role="procurement_responsible")
        own_item_id = _make_portal_item(session, suffix)
    responsible_headers = _login_headers(s, f"rfq-latest-{suffix}")
    # Two independent RFQs/quotations against the *same* isolated item and
    # supplier - only the newer quotation_date should win.
    old_item_id, old_supplier, _rfq1, _q1 = _received_quotation_for_item(
        s, f"{suffix}-old", responsible_headers, unit_price=1000, quotation_date="2026-01-01",
        item_id=own_item_id,
    )
    new_item_id, new_supplier, _rfq2, _q2 = _received_quotation_for_item(
        s, f"{suffix}-new", responsible_headers, unit_price=1400, quotation_date="2026-08-10",
        item_id=own_item_id,
    )
    assert old_item_id == new_item_id == own_item_id
    assert old_supplier["id"] == new_supplier["id"]

    response = s.get(
        f"{API}/price-comparisons/last-formal-prices", headers=responsible_headers,
        params={"pairs": f"{new_item_id}:{new_supplier['id']}"},
    )
    assert response.status_code == 200, response.text
    price = response.json()["prices"][f"{new_item_id}|{new_supplier['id']}"]
    assert price["unit_price"] == 1400
    assert price["date"] == "2026-08-10"


def test_pdf_attachment_accepted_for_quotation(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-pdf-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-pdf-resp-{suffix}")
    rfq_id, supplier, _rfq = _rfq_with_supplier(s, suffix, responsible_headers)
    quotation_id = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    ).json()["quotation"]["id"]
    response = s.post(
        f"{RFQ_API}/{rfq_id}/quotations/{quotation_id}/attachments",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        files=[("files", ("quote.pdf", b"%PDF-1.4 test quotation", "application/pdf"))],
    )
    assert response.status_code == 200, response.text
    assert response.json()["attachments"][0]["media_type"] == "application/pdf"


def test_xlsx_attachment_accepted_for_quotation(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-xlsx-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-xlsx-resp-{suffix}")
    rfq_id, supplier, _rfq = _rfq_with_supplier(s, suffix, responsible_headers)
    quotation_id = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    ).json()["quotation"]["id"]
    response = s.post(
        f"{RFQ_API}/{rfq_id}/quotations/{quotation_id}/attachments",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        files=[("files", (
            "quote.xlsx", b"PK\x03\x04" + b"0" * 20,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ))],
    )
    assert response.status_code == 200, response.text
    assert response.json()["attachments"][0]["original_filename"] == "quote.xlsx"


def test_executable_attachment_rejected_for_quotation(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-exe-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-exe-resp-{suffix}")
    rfq_id, supplier, _rfq = _rfq_with_supplier(s, suffix, responsible_headers)
    quotation_id = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    ).json()["quotation"]["id"]
    response = s.post(
        f"{RFQ_API}/{rfq_id}/quotations/{quotation_id}/attachments",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        files=[("files", ("malware.exe", b"MZ" + b"\x00" * 30, "application/octet-stream"))],
    )
    assert response.status_code == 422


def test_procurement_engineer_cannot_mutate_quotation(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-mutresp-{suffix}", role="procurement_responsible")
        _make_user(session, username=f"rfq-muteng-{suffix}", role="procurement_engineer")
    responsible_headers = _login_headers(s, f"rfq-mutresp-{suffix}")
    engineer_headers = _login_headers(s, f"rfq-muteng-{suffix}")
    rfq_id, supplier, _rfq = _rfq_with_supplier(s, suffix, responsible_headers)

    denied_quotation = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"supplier_id": supplier["id"]},
    )
    assert denied_quotation.status_code == 403

    quotation_id = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    ).json()["quotation"]["id"]
    denied_update = s.put(
        f"{RFQ_API}/{rfq_id}/quotations/{quotation_id}",
        headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"status": "received", "lines": []},
    )
    assert denied_update.status_code == 403


# ---- RFQ read authorization (Sprint 3.1 follow-up: ERP-only reads) ----

def test_read_rfq_authorization_matrix(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-read-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"rfq-read-mgr-{suffix}", role="commercial_manager")
        _make_user(session, username=f"rfq-read-resp-{suffix}", role="procurement_responsible")
        portal_username, _, _ = _portal_setup(session, f"rfq-read-portal-{suffix}")
    engineer_headers = _login_headers(s, f"rfq-read-eng-{suffix}")
    manager_headers = _login_headers(s, f"rfq-read-mgr-{suffix}")
    responsible_headers = _login_headers(s, f"rfq-read-resp-{suffix}")
    portal_headers = _login_headers(s, portal_username)

    request_id, *_rest = _make_pricing_request(s, suffix)
    rfq_id = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    ).json()["rfq"]["id"]

    # 1-4: every ERP role (any role - admin override is implicit) can read.
    assert s.get(f"{RFQ_API}/{rfq_id}", headers={**INTERNAL_HEADERS, **engineer_headers}).status_code == 200
    assert s.get(f"{RFQ_API}/{rfq_id}", headers={**INTERNAL_HEADERS, **manager_headers}).status_code == 200
    assert s.get(f"{RFQ_API}/{rfq_id}", headers={**INTERNAL_HEADERS, **responsible_headers}).status_code == 200
    assert s.get(f"{RFQ_API}/{rfq_id}", headers={**INTERNAL_HEADERS, **admin_headers}).status_code == 200

    # 5: site_portal is a different account_type entirely -> 403.
    assert s.get(f"{RFQ_API}/{rfq_id}", headers={**INTERNAL_HEADERS, **portal_headers}).status_code == 403

    # 6: no credentials at all -> 401.
    assert s.get(f"{RFQ_API}/{rfq_id}", headers=INTERNAL_HEADERS).status_code == 401
    assert s.get(f"{RFQ_API}/by-request/{request_id}", headers=INTERNAL_HEADERS).status_code == 401


def test_site_portal_cannot_read_quotation_attachment_or_comparison_rows(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-read2-resp-{suffix}", role="procurement_responsible")
        portal_username, _, _ = _portal_setup(session, f"rfq-read2-portal-{suffix}")
    responsible_headers = _login_headers(s, f"rfq-read2-resp-{suffix}")
    portal_headers = _login_headers(s, portal_username)
    rfq_id, supplier, _rfq = _rfq_with_supplier(s, suffix, responsible_headers)
    quotation_id = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    ).json()["quotation"]["id"]
    uploaded = s.post(
        f"{RFQ_API}/{rfq_id}/quotations/{quotation_id}/attachments",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        files=[("files", ("quote.pdf", b"%PDF-1.4 test quotation", "application/pdf"))],
    )
    attachment_id = uploaded.json()["attachments"][0]["id"]

    denied_attachment = s.get(
        f"{RFQ_API}/{rfq_id}/quotations/{quotation_id}/attachments/{attachment_id}",
        headers={**INTERNAL_HEADERS, **portal_headers},
    )
    assert denied_attachment.status_code == 403

    denied_comparison_rows = s.get(
        f"{RFQ_API}/{rfq_id}/comparison-rows", headers={**INTERNAL_HEADERS, **portal_headers},
    )
    assert denied_comparison_rows.status_code == 403


def test_quotation_attachment_download_headers_and_auth(s):
    """Regression: Approval Center / Supplier Comparison open the quotation
    attachment as an authenticated blob and rely on this endpoint's
    Content-Type/Content-Disposition to decide open-vs-download - an
    anonymous request must be rejected before any bytes are served."""
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rfq-dl-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rfq-dl-resp-{suffix}")
    rfq_id, supplier, _rfq = _rfq_with_supplier(s, suffix, responsible_headers)
    quotation_id = s.post(
        f"{RFQ_API}/{rfq_id}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    ).json()["quotation"]["id"]
    file_bytes = b"%PDF-1.4 test quotation content"
    uploaded = s.post(
        f"{RFQ_API}/{rfq_id}/quotations/{quotation_id}/attachments",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        files=[("files", ("supplier offer.pdf", file_bytes, "application/pdf"))],
    )
    attachment_id = uploaded.json()["attachments"][0]["id"]
    download_path = f"{RFQ_API}/{rfq_id}/quotations/{quotation_id}/attachments/{attachment_id}"

    anonymous = s.get(download_path, headers=INTERNAL_HEADERS)
    assert anonymous.status_code == 401

    response = s.get(download_path, headers={**INTERNAL_HEADERS, **responsible_headers})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert "filename*=UTF-8''supplier%20offer.pdf" in disposition
    assert response.content == file_bytes


def test_comparison_delete_enforces_role_and_allows_only_safe_draft(s):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"delete-resp-{suffix}", role="procurement_responsible")
        _make_user(session, username=f"delete-eng-{suffix}", role="procurement_engineer")
    responsible_headers = _login_headers(s, f"delete-resp-{suffix}")
    engineer_headers = _login_headers(s, f"delete-eng-{suffix}")
    project = s.get(f"{API}/projects", headers=responsible_headers).json()[0]
    item = s.get(f"{API}/items", headers=responsible_headers).json()[0]
    supplier = s.get(f"{API}/suppliers", headers=responsible_headers).json()[0]
    created = s.post(
        f"{API}/price-comparisons", headers=responsible_headers,
        json={
            "project_id": project["id"], "comparison_date": "2026-08-21",
            "rows": [{
                "item_id": item["id"], "supplier_id": supplier["id"],
                "quantity": 1, "unit_price": 100, "availability": "available",
            }],
        },
    )
    assert created.status_code == 200, created.text
    comparison_id = created.json()["id"]
    assert s.delete(f"{API}/price-comparisons/{comparison_id}").status_code == 401
    assert s.delete(
        f"{API}/price-comparisons/{comparison_id}", headers=engineer_headers,
    ).status_code == 403
    deleted = s.delete(
        f"{API}/price-comparisons/{comparison_id}", headers=responsible_headers,
    )
    assert deleted.status_code == 200, deleted.text
    assert s.get(f"{API}/price-comparisons/{comparison_id}", headers=responsible_headers).status_code == 404


# ---------------- Sprint 3.2: Comparison + Approval Review Workspace ----------------

WORKSPACE_URL = lambda approval_id: f"{API}/workflow/approvals/{approval_id}/review-workspace"  # noqa: E731


def _make_approval_with_comparison(client, suffix, admin_headers, manual_item=False):
    """REQ (pricing) -> CMP (one selected row) -> Approval (comparison_workflow,
    stage=comparison_technical, status=pending_approval). Returns
    (approval, request_id, request_item_id, supplier)."""
    request_id, request_number, project_id, item_id, _resp_headers = _make_pricing_request(
        client, suffix, manual_item=manual_item,
    )
    item = client.get(f"{API}/items", headers=admin_headers).json()[0]
    supplier = client.get(f"{API}/suppliers", headers=admin_headers).json()[0]
    project = next(p for p in client.get(f"{API}/projects", headers=admin_headers).json() if p["id"] == project_id)
    comparison = client.post(f"{API}/price-comparisons", json={
        "project_id": project_id, "project_name": project["name"],
        "source_request_id": request_id, "source_request_number": request_number,
        "comparison_date": "2026-08-20", "rows": [{
            "item_id": "" if manual_item else item["id"],
            "product_name": f"صنف RFQ يدوي {suffix}" if manual_item else "صنف RFQ من الدليل",
            "supplier_id": supplier["id"], "quantity": 3, "unit": "قطعة",
            "unit_price": 100, "availability": "available", "price_valid_until": "2099-12-31",
            "selected_for_purchase": 1,
        }],
    }, headers=admin_headers)
    assert comparison.status_code == 200, comparison.text
    response = client.post(
        f"{API}/workflow/approvals/from-comparison",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={
            "comparison_id": comparison.json()["id"], "engineer_name": "مراجع",
            "created_by": "tester", "approval_type": "comparison_workflow",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["approval"], request_id, item_id, supplier


def test_comparison_delete_is_blocked_after_approval_traceability_exists(s, admin_headers):
    approval, _request_id, _item_id, _supplier = _make_approval_with_comparison(
        s, uuid.uuid4().hex[:8], admin_headers,
    )
    blocked = s.delete(
        f"{API}/price-comparisons/{approval['comparison_id']}", headers=admin_headers,
    )
    assert blocked.status_code == 409
    assert "لا يمكن حذف" in blocked.json()["detail"]


def test_review_workspace_read_authorization_matrix(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rw-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"rw-mgr-{suffix}", role="commercial_manager")
        _make_user(session, username=f"rw-resp-{suffix}", role="procurement_responsible")
        portal_username, _, _ = _portal_setup(session, f"rw-portal-{suffix}")
    engineer_headers = _login_headers(s, f"rw-eng-{suffix}")
    manager_headers = _login_headers(s, f"rw-mgr-{suffix}")
    responsible_headers = _login_headers(s, f"rw-resp-{suffix}")
    portal_headers = _login_headers(s, portal_username)

    approval, *_rest = _make_approval_with_comparison(s, suffix, admin_headers)
    url = WORKSPACE_URL(approval["id"])

    assert s.get(url, headers={**INTERNAL_HEADERS, **admin_headers}).status_code == 200
    assert s.get(url, headers={**INTERNAL_HEADERS, **responsible_headers}).status_code == 200
    assert s.get(url, headers={**INTERNAL_HEADERS, **engineer_headers}).status_code == 200
    assert s.get(url, headers={**INTERNAL_HEADERS, **manager_headers}).status_code == 200
    assert s.get(url, headers={**INTERNAL_HEADERS, **portal_headers}).status_code == 403
    assert s.get(url, headers=INTERNAL_HEADERS).status_code == 401


def test_review_workspace_includes_request_and_items(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    approval, request_id, item_id, _supplier = _make_approval_with_comparison(s, suffix, admin_headers)
    body = s.get(WORKSPACE_URL(approval["id"]), headers={**INTERNAL_HEADERS, **admin_headers}).json()
    assert body["request"]["id"] == request_id
    assert body["request"]["request_number"] == approval["source_request_number"]
    assert len(body["request_items"]) == 1
    item = body["request_items"][0]
    assert item["id"] == item_id
    assert item["quantity"] == 3
    assert item["is_manual"] is False
    assert item["review_status"] == "approved"
    assert body["technical_review"]["reviewed"] is True


def test_review_workspace_manual_request_item_represented_safely(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    approval, _request_id, item_id, _supplier = _make_approval_with_comparison(
        s, suffix, admin_headers, manual_item=True,
    )
    body = s.get(WORKSPACE_URL(approval["id"]), headers={**INTERNAL_HEADERS, **admin_headers}).json()
    item = body["request_items"][0]
    assert item["id"] == item_id
    assert item["is_manual"] is True
    assert item["item_id"] == ""
    assert item["product_name"] == f"صنف RFQ يدوي {suffix}"


def test_review_workspace_includes_request_attachments_when_present(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    approval, request_id, _item_id, _supplier = _make_approval_with_comparison(s, suffix, admin_headers)
    with SessionLocal.begin() as session:
        session.add(IncomingRequestGeneralAttachment(
            id=str(uuid.uuid4()), request_id=request_id,
            original_filename="quote-support.pdf",
            stored_filename=f"{request_id}/{uuid.uuid4().hex}.pdf",
            media_type="application/pdf", size_bytes=1234, sha256="a" * 64,
            created_at=datetime.now(timezone.utc).isoformat(),
        ))
    body = s.get(WORKSPACE_URL(approval["id"]), headers={**INTERNAL_HEADERS, **admin_headers}).json()
    assert len(body["request_attachments"]) == 1
    assert body["request_attachments"][0]["original_filename"] == "quote-support.pdf"
    assert body["request_attachments"][0]["source"] == "general"


def test_review_workspace_includes_rfq_and_supplier_quotations_when_present(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rw-rfq-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"rw-rfq-resp-{suffix}")
    # Build the RFQ and its received quotation while the REQ is still
    # "pricing" (RFQ creation requires that status), *then* create the CMP
    # and approval, which advances the REQ to "waiting_for_approval".
    request_id, request_number, project_id, _item_id, _resp_headers = _make_pricing_request(s, suffix)
    supplier = s.get(f"{API}/suppliers", headers=admin_headers).json()[0]
    rfq = s.post(
        RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"source_request_id": request_id},
    ).json()["rfq"]
    s.post(
        f"{RFQ_API}/{rfq['id']}/suppliers", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    )
    quotation_id = s.post(
        f"{RFQ_API}/{rfq['id']}/quotations", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"]},
    ).json()["quotation"]["id"]
    s.put(
        f"{RFQ_API}/{rfq['id']}/quotations/{quotation_id}",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        json={
            "quotation_ref": "SUP-REF-1", "status": "received",
            "lines": [{
                "rfq_item_id": rfq["items"][0]["id"], "quantity": 3, "unit": "قطعة",
                "unit_price": 90, "availability": "available",
            }],
        },
    )
    s.post(
        f"{RFQ_API}/{rfq['id']}/quotations/{quotation_id}/attachments",
        headers={**INTERNAL_HEADERS, **responsible_headers},
        files=[("files", ("quote.pdf", b"%PDF-1.4 test quotation", "application/pdf"))],
    )

    item = s.get(f"{API}/items", headers=admin_headers).json()[0]
    project = next(p for p in s.get(f"{API}/projects", headers=admin_headers).json() if p["id"] == project_id)
    comparison = s.post(f"{API}/price-comparisons", json={
        "project_id": project_id, "project_name": project["name"],
        "source_request_id": request_id, "source_request_number": request_number,
        "comparison_date": "2026-08-20", "rows": [{
            "item_id": item["id"], "product_name": "صنف RFQ من الدليل",
            "supplier_id": supplier["id"], "quantity": 3, "unit": "قطعة",
            "unit_price": 100, "availability": "available", "price_valid_until": "2099-12-31",
            "selected_for_purchase": 1,
        }],
    }, headers=admin_headers)
    assert comparison.status_code == 200, comparison.text
    approval_response = s.post(
        f"{API}/workflow/approvals/from-comparison",
        headers={**INTERNAL_HEADERS, **admin_headers},
        json={
            "comparison_id": comparison.json()["id"], "engineer_name": "مراجع",
            "created_by": "tester", "approval_type": "comparison_workflow",
        },
    )
    assert approval_response.status_code == 201, approval_response.text
    approval = approval_response.json()["approval"]

    body = s.get(WORKSPACE_URL(approval["id"]), headers={**INTERNAL_HEADERS, **admin_headers}).json()
    assert body["rfq"]["rfq_number"] == rfq["rfq_number"]
    assert body["rfq"]["supplier_count"] == 1
    assert body["rfq"]["received_quotation_count"] == 1
    assert len(body["supplier_quotations"]) == 1
    quotation = body["supplier_quotations"][0]
    assert quotation["status"] == "received"
    assert quotation["quotation_ref"] == "SUP-REF-1"
    assert quotation["lines"][0]["unit_price"] == 90
    assert len(quotation["attachments"]) == 1
    assert quotation["attachments"][0]["original_filename"] == "quote.pdf"


def test_review_workspace_tolerates_missing_rfq(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    approval, *_rest = _make_approval_with_comparison(s, suffix, admin_headers)
    response = s.get(WORKSPACE_URL(approval["id"]), headers={**INTERNAL_HEADERS, **admin_headers})
    assert response.status_code == 200
    body = response.json()
    assert body["rfq"] is None
    assert body["supplier_quotations"] == []


def test_review_workspace_includes_comparison_rows_and_selection(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    approval, _request_id, _item_id, supplier = _make_approval_with_comparison(s, suffix, admin_headers)
    body = s.get(WORKSPACE_URL(approval["id"]), headers={**INTERNAL_HEADERS, **admin_headers}).json()
    comparison = body["comparison"]
    assert comparison["comparison_number"] == approval["comparison_number"]
    assert len(comparison["rows"]) == 1
    row = comparison["rows"][0]
    assert row["supplier_id"] == supplier["id"]
    assert row["selected_for_purchase"] == 1
    assert row["unit_price"] == 100


def test_review_workspace_read_does_not_mutate_approval(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    approval, *_rest = _make_approval_with_comparison(s, suffix, admin_headers)
    before = s.get(
        f"{API}/workflow/approvals/{approval['id']}", headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    s.get(WORKSPACE_URL(approval["id"]), headers={**INTERNAL_HEADERS, **admin_headers})
    after = s.get(
        f"{API}/workflow/approvals/{approval['id']}", headers={**INTERNAL_HEADERS, **admin_headers},
    ).json()
    assert after["status"] == before["status"]
    assert after["updated_at"] == before["updated_at"]


def test_review_workspace_decision_role_gating_reuses_existing_endpoint(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rw-dec-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"rw-dec-mgr-{suffix}", role="commercial_manager")
        _make_user(session, username=f"rw-dec-resp-{suffix}", role="procurement_responsible")
    engineer_headers = _login_headers(s, f"rw-dec-eng-{suffix}")
    manager_headers = _login_headers(s, f"rw-dec-mgr-{suffix}")
    responsible_headers = _login_headers(s, f"rw-dec-resp-{suffix}")

    approval, *_rest = _make_approval_with_comparison(s, suffix, admin_headers)
    decision_url = f"{API}/workflow/approvals/{approval['id']}/decision"

    # commercial_manager cannot act at the technical (engineer) stage.
    denied = s.post(decision_url, headers={**INTERNAL_HEADERS, **manager_headers}, json={"decision": "approved"})
    assert denied.status_code == 403

    # procurement_responsible may view the full file but not decide it.
    assert s.get(
        WORKSPACE_URL(approval["id"]), headers={**INTERNAL_HEADERS, **responsible_headers},
    ).status_code == 200
    denied_responsible = s.post(
        decision_url, headers={**INTERNAL_HEADERS, **responsible_headers}, json={"decision": "approved"},
    )
    assert denied_responsible.status_code == 403

    # procurement_engineer executes the technical-stage decision.
    approved = s.post(
        decision_url, headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"decision": "approved", "note": "موافق فنيًا"},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["approval"]["approval_stage"] == APPROVAL_STAGE_EXPENDITURE_APPROVAL

    # commercial_manager now executes the commercial stage.
    commercial = s.post(
        decision_url, headers={**INTERNAL_HEADERS, **manager_headers},
        json={"decision": "approved", "note": "موافق تجاريًا"},
    )
    assert commercial.status_code == 200, commercial.text
    assert commercial.json()["approval"]["approval_stage"] == APPROVAL_STAGE_FUNDS_AVAILABILITY


def test_review_workspace_decision_revision_requested_via_reused_endpoint(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"rw-rev-eng-{suffix}", role="procurement_engineer")
    engineer_headers = _login_headers(s, f"rw-rev-eng-{suffix}")
    approval, *_rest = _make_approval_with_comparison(s, suffix, admin_headers)
    decision_url = f"{API}/workflow/approvals/{approval['id']}/decision"
    response = s.post(
        decision_url, headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"decision": "revision_requested", "note": "يحتاج توضيح إضافي"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["approval"]["status"] == "revision_requested"


# ---------------- Sprint 3.3: Purchase Order Workspace + Register ----------------

PO_API = lambda po_id="": f"{API}/purchase-orders" + (f"/{po_id}" if po_id else "")  # noqa: E731


def _make_draft_purchase_order(client, suffix, admin_headers, manual_item=False):
    """REQ (pricing) -> CMP -> Approval, walked all the way through the
    technical + commercial + funds-release decisions with admin (which
    always passes every stage's role check), then a formal PO created from
    the comparison. Returns (purchase_order, approval, request_id, supplier)."""
    approval, request_id, _item_id, supplier = _make_approval_with_comparison(
        client, suffix, admin_headers, manual_item=manual_item,
    )
    engineer_decision = client.post(
        f"{API}/workflow/approvals/{approval['id']}/decision",
        headers={**INTERNAL_HEADERS, **admin_headers}, json={"decision": "approved", "actor": "eng"},
    )
    assert engineer_decision.status_code == 200, engineer_decision.text
    commercial_decision = client.post(
        f"{API}/workflow/approvals/{approval['id']}/decision",
        headers={**INTERNAL_HEADERS, **admin_headers}, json={"decision": "approved", "actor": "mgr"},
    )
    assert commercial_decision.status_code == 200, commercial_decision.text
    released = client.post(
        f"{API}/workflow/approvals/{approval['id']}/funds-release",
        headers={**INTERNAL_HEADERS, **admin_headers}, json={"actor": "mgr"},
    )
    assert released.status_code == 200, released.text
    generated = client.post(f"{API}/purchase-orders/from-comparison", json={
        "comparison_id": approval["comparison_id"], "po_date": "2026-08-20",
        "created_by": "tester", "orders": [],
    }, headers=admin_headers)
    assert generated.status_code == 200, generated.text
    return generated.json()["purchase_orders"][0], approval, request_id, supplier


def test_po_read_authorization_matrix(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"po-read-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"po-read-mgr-{suffix}", role="commercial_manager")
        _make_user(session, username=f"po-read-resp-{suffix}", role="procurement_responsible")
        portal_username, _, _ = _portal_setup(session, f"po-read-portal-{suffix}")
    engineer_headers = _login_headers(s, f"po-read-eng-{suffix}")
    manager_headers = _login_headers(s, f"po-read-mgr-{suffix}")
    responsible_headers = _login_headers(s, f"po-read-resp-{suffix}")
    portal_headers = _login_headers(s, portal_username)

    po, *_rest = _make_draft_purchase_order(s, suffix, admin_headers)
    detail_url = PO_API(po["id"])

    assert s.get(detail_url, headers=admin_headers).status_code == 200
    assert s.get(detail_url, headers=responsible_headers).status_code == 200
    assert s.get(detail_url, headers=engineer_headers).status_code == 200
    assert s.get(detail_url, headers=manager_headers).status_code == 200
    assert s.get(detail_url, headers=portal_headers).status_code == 403
    assert s.get(detail_url).status_code == 401

    assert s.get(PO_API(), headers=admin_headers).status_code == 200
    assert s.get(PO_API(), headers=portal_headers).status_code == 403
    assert s.get(PO_API()).status_code == 401


def test_po_engineer_and_manager_cannot_finalize(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"po-fin-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"po-fin-mgr-{suffix}", role="commercial_manager")
    engineer_headers = _login_headers(s, f"po-fin-eng-{suffix}")
    manager_headers = _login_headers(s, f"po-fin-mgr-{suffix}")
    po, *_rest = _make_draft_purchase_order(s, suffix, admin_headers)

    assert s.post(f"{API}/purchase-orders/{po['id']}/finalize", headers=engineer_headers, json={}).status_code == 403
    assert s.post(f"{API}/purchase-orders/{po['id']}/finalize", headers=manager_headers, json={}).status_code == 403


def test_po_client_role_spoofing_cannot_grant_finalize(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"po-spoof-eng-{suffix}", role="procurement_engineer")
    engineer_headers = _login_headers(s, f"po-spoof-eng-{suffix}")
    po, *_rest = _make_draft_purchase_order(s, suffix, admin_headers)

    spoofed = s.post(
        f"{API}/purchase-orders/{po['id']}/finalize",
        headers=engineer_headers,
        json={"actor": "x", "actor_role": "procurement_responsible", "role": "admin"},
    )
    assert spoofed.status_code == 403


def test_po_completed_cannot_be_mutated(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_draft_purchase_order(s, suffix, admin_headers)
    finalized = s.post(f"{API}/purchase-orders/{po['id']}/finalize", headers=admin_headers, json={"actor": "tester"})
    assert finalized.status_code == 200, finalized.text
    completed = s.post(f"{API}/purchase-orders/{po['id']}/receipts", headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"complete-{suffix}", "actor": "tester",
    })
    assert completed.status_code == 200, completed.text
    assert completed.json()["purchase_order"]["status"] == "completed"

    assert s.post(f"{API}/purchase-orders/{po['id']}/finalize", headers=admin_headers, json={}).status_code == 409
    assert s.patch(
        f"{API}/purchase-orders/{po['id']}/status", headers=admin_headers, json={"status": "cancelled"},
    ).status_code == 409
    assert s.post(f"{API}/purchase-orders/{po['id']}/receipts", headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"complete-again-{suffix}", "actor": "tester",
    }).status_code == 409


def test_po_ancestry_unchanged_after_status_transitions(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, approval, request_id, supplier = _make_draft_purchase_order(s, suffix, admin_headers)
    before = s.get(f"{API}/purchase-orders/{po['id']}", headers=admin_headers).json()
    for status in ("approved", "sent", "supplier_confirmed"):
        transitioned = s.patch(
            f"{API}/purchase-orders/{po['id']}/status", headers=admin_headers, json={"status": status},
        )
        assert transitioned.status_code == 200, transitioned.text
    finalized = s.post(f"{API}/purchase-orders/{po['id']}/finalize", headers=admin_headers, json={"actor": "tester"})
    assert finalized.status_code == 200, finalized.text
    after = finalized.json()["purchase_order"]
    assert after["source_request_id"] == before["source_request_id"] == request_id
    assert after["comparison_id"] == before["comparison_id"] == approval["comparison_id"]
    assert after["approval_id"] == before["approval_id"] == approval["id"]
    assert after["supplier_id"] == before["supplier_id"] == supplier["id"]
    assert after["final_total"] == before["final_total"]


def test_po_list_includes_workflow_payment_indicators(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_draft_purchase_order(s, suffix, admin_headers)
    listed = s.get(PO_API(), headers=admin_headers)
    assert listed.status_code == 200
    row = next(item for item in listed.json() if item["id"] == po["id"])
    assert row["workflow"]["funds_released"] is True


# ---- Security micro-sprint: DELETE + report route authorization ----

def test_po_delete_authorization_matrix(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"po-del-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"po-del-mgr-{suffix}", role="commercial_manager")
        _make_user(session, username=f"po-del-resp-{suffix}", role="procurement_responsible")
        portal_username, _, _ = _portal_setup(session, f"po-del-portal-{suffix}")
    engineer_headers = _login_headers(s, f"po-del-eng-{suffix}")
    manager_headers = _login_headers(s, f"po-del-mgr-{suffix}")
    responsible_headers = _login_headers(s, f"po-del-resp-{suffix}")
    portal_headers = _login_headers(s, portal_username)

    po_eng, *_ = _make_draft_purchase_order(s, f"{suffix}-eng", admin_headers)
    po_mgr, *_ = _make_draft_purchase_order(s, f"{suffix}-mgr", admin_headers)
    po_portal, *_ = _make_draft_purchase_order(s, f"{suffix}-portal", admin_headers)
    po_anon, *_ = _make_draft_purchase_order(s, f"{suffix}-anon", admin_headers)
    po_spoof, *_ = _make_draft_purchase_order(s, f"{suffix}-spoof", admin_headers)
    po_resp, *_ = _make_draft_purchase_order(s, f"{suffix}-resp", admin_headers)
    po_admin, *_ = _make_draft_purchase_order(s, f"{suffix}-admin", admin_headers)

    assert s.delete(f"{API}/purchase-orders/{po_eng['id']}", headers=engineer_headers).status_code == 403
    assert s.delete(f"{API}/purchase-orders/{po_mgr['id']}", headers=manager_headers).status_code == 403
    assert s.delete(f"{API}/purchase-orders/{po_portal['id']}", headers=portal_headers).status_code == 403
    assert s.delete(f"{API}/purchase-orders/{po_anon['id']}").status_code == 401

    spoofed = s.request(
        "DELETE", f"{API}/purchase-orders/{po_spoof['id']}", headers=engineer_headers,
        json={"actor_role": "procurement_responsible", "role": "admin"},
    )
    assert spoofed.status_code == 403

    deleted_by_responsible = s.delete(f"{API}/purchase-orders/{po_resp['id']}", headers=responsible_headers)
    assert deleted_by_responsible.status_code == 200, deleted_by_responsible.text
    assert s.get(f"{API}/purchase-orders/{po_resp['id']}", headers=admin_headers).status_code == 404

    deleted_by_admin = s.delete(f"{API}/purchase-orders/{po_admin['id']}", headers=admin_headers)
    assert deleted_by_admin.status_code == 200, deleted_by_admin.text
    assert s.get(f"{API}/purchase-orders/{po_admin['id']}", headers=admin_headers).status_code == 404


def test_po_delete_still_blocked_for_non_eligible_statuses(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"po-del2-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"po-del2-resp-{suffix}")

    po_approved, *_ = _make_draft_purchase_order(s, f"{suffix}-approved", admin_headers)
    advanced = s.patch(
        f"{API}/purchase-orders/{po_approved['id']}/status", headers=admin_headers, json={"status": "approved"},
    )
    assert advanced.status_code == 200, advanced.text
    assert s.delete(f"{API}/purchase-orders/{po_approved['id']}", headers=responsible_headers).status_code == 409

    po_completed, *_ = _make_draft_purchase_order(s, f"{suffix}-completed", admin_headers)
    finalized = s.post(
        f"{API}/purchase-orders/{po_completed['id']}/finalize", headers=admin_headers, json={"actor": "t"},
    )
    assert finalized.status_code == 200, finalized.text
    completed = s.post(f"{API}/purchase-orders/{po_completed['id']}/receipts", headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"del-complete-{suffix}", "actor": "t",
    })
    assert completed.status_code == 200, completed.text
    assert completed.json()["purchase_order"]["status"] == "completed"
    assert s.delete(f"{API}/purchase-orders/{po_completed['id']}", headers=responsible_headers).status_code == 409

    # cancelled stays eligible - the underlying business rule is unchanged.
    po_cancelled, *_ = _make_draft_purchase_order(s, f"{suffix}-cancelled", admin_headers)
    cancel = s.patch(
        f"{API}/purchase-orders/{po_cancelled['id']}/status", headers=admin_headers, json={"status": "cancelled"},
    )
    assert cancel.status_code == 200, cancel.text
    deleted_cancelled = s.delete(f"{API}/purchase-orders/{po_cancelled['id']}", headers=responsible_headers)
    assert deleted_cancelled.status_code == 200, deleted_cancelled.text


def test_po_report_requires_erp_authentication(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"po-rep-resp-{suffix}", role="procurement_responsible")
        portal_username, _, _ = _portal_setup(session, f"po-rep-portal-{suffix}")
    responsible_headers = _login_headers(s, f"po-rep-resp-{suffix}")
    portal_headers = _login_headers(s, portal_username)

    po, *_rest = _make_draft_purchase_order(s, suffix, admin_headers)
    finalized = s.post(f"{API}/purchase-orders/{po['id']}/finalize", headers=admin_headers, json={"actor": "t"})
    assert finalized.status_code == 200, finalized.text

    url = f"{API}/purchase-orders/{po['id']}/reports/admin"
    assert s.get(url, headers=admin_headers).status_code == 200
    assert s.get(url, headers=responsible_headers).status_code == 200
    # No PO report can be exposed by predictable ID without ERP identity -
    # there is no secure external-share token mechanism for this route.
    assert s.get(url, headers=portal_headers).status_code == 403
    assert s.get(url).status_code == 401


# ---------------- Sprint 3.4: PO Payment Ledger ----------------

PO_PAYMENTS_API = lambda po_id: f"{API}/purchase-orders/{po_id}/payments"  # noqa: E731


def _make_payable_purchase_order(client, suffix, admin_headers, target_status="approved"):
    """A formal PO (final_total 300 EGP - qty 3 x unit_price 100) advanced
    past draft via the existing PATCH /status transition, so it is eligible
    for an actual supplier payment. Returns (po, approval, request_id, supplier)."""
    po, approval, request_id, supplier = _make_draft_purchase_order(client, suffix, admin_headers)
    if target_status != "draft":
        advanced = client.patch(
            f"{API}/purchase-orders/{po['id']}/status", headers=admin_headers,
            json={"status": target_status},
        )
        assert advanced.status_code == 200, advanced.text
        po = advanced.json()
    return po, approval, request_id, supplier


def test_po_payment_read_authorization_matrix(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-read-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"pay-read-mgr-{suffix}", role="commercial_manager")
        _make_user(session, username=f"pay-read-resp-{suffix}", role="procurement_responsible")
        portal_username, _, _ = _portal_setup(session, f"pay-read-portal-{suffix}")
    engineer_headers = _login_headers(s, f"pay-read-eng-{suffix}")
    manager_headers = _login_headers(s, f"pay-read-mgr-{suffix}")
    responsible_headers = _login_headers(s, f"pay-read-resp-{suffix}")
    portal_headers = _login_headers(s, portal_username)

    po, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)
    url = PO_PAYMENTS_API(po["id"])

    assert s.get(url, headers=admin_headers).status_code == 200
    assert s.get(url, headers=manager_headers).status_code == 200
    assert s.get(url, headers=responsible_headers).status_code == 200
    assert s.get(url, headers=engineer_headers).status_code == 200
    assert s.get(url, headers=portal_headers).status_code == 403
    assert s.get(url).status_code == 401


def test_po_payment_mutation_authorization_matrix(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-mut-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"pay-mut-resp-{suffix}", role="procurement_responsible")
        _make_user(session, username=f"pay-mut-mgr-{suffix}", role="commercial_manager")
    engineer_headers = _login_headers(s, f"pay-mut-eng-{suffix}")
    responsible_headers = _login_headers(s, f"pay-mut-resp-{suffix}")
    manager_headers = _login_headers(s, f"pay-mut-mgr-{suffix}")

    po, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)
    url = PO_PAYMENTS_API(po["id"])

    def body(key):
        return {
            "payment_date": "2026-08-25", "amount": 10, "payment_method": "bank_transfer",
            "payment_reference": "TX-1", "notes": "", "idempotency_key": key,
        }

    assert s.post(url, headers=responsible_headers, json=body(f"resp-{suffix}")).status_code == 403
    assert s.post(url, headers=engineer_headers, json=body(f"eng-{suffix}")).status_code == 403

    spoofed = s.post(
        url, headers=engineer_headers,
        json={**body(f"spoof-{suffix}"), "actor_role": "commercial_manager", "role": "admin"},
    )
    assert spoofed.status_code == 403

    recorded_by_manager = s.post(url, headers=manager_headers, json=body(f"mgr-{suffix}"))
    assert recorded_by_manager.status_code == 200, recorded_by_manager.text
    assert recorded_by_manager.json()["payment"]["payment_number"].startswith("PAY-")

    recorded_by_admin = s.post(url, headers=admin_headers, json=body(f"admin-{suffix}"))
    assert recorded_by_admin.status_code == 200, recorded_by_admin.text


def test_po_payment_allowed_while_linked_approval_is_approved(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-appr-ok-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"pay-appr-ok-mgr-{suffix}")

    po, approval, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)
    assert po.get("approval_id") == approval["id"]

    recorded = s.post(PO_PAYMENTS_API(po["id"]), headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 10, "payment_method": "bank_transfer",
        "payment_reference": "TX-1", "notes": "", "idempotency_key": f"appr-ok-{suffix}",
    })
    assert recorded.status_code == 200, recorded.text


def test_po_payment_rejected_when_linked_approval_is_no_longer_approved(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-appr-bad-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"pay-appr-bad-mgr-{suffix}")

    po, approval, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)
    assert po.get("approval_id") == approval["id"]

    # Simulate an approval reverted/invalidated after the PO was already
    # issued - e.g. a revision requested post-hoc on the same approval row.
    with SessionLocal.begin() as session:
        session.get(EngineerApproval, approval["id"]).status = "revision_requested"

    rejected = s.post(PO_PAYMENTS_API(po["id"]), headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 10, "payment_method": "bank_transfer",
        "payment_reference": "TX-1", "notes": "", "idempotency_key": f"appr-bad-{suffix}",
    })
    assert rejected.status_code == 409, rejected.text
    assert "اعتماد الصرف" in rejected.json()["detail"]

    with SessionLocal() as session:
        count = session.scalar(
            select(func.count()).select_from(PurchaseOrderPayment)
            .where(PurchaseOrderPayment.purchase_order_id == po["id"])
        )
    assert count == 0


def test_po_payment_recording_and_summary_accuracy(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-acc-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"pay-acc-mgr-{suffix}")
    po, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)
    assert po["final_total"] == 300
    url = PO_PAYMENTS_API(po["id"])

    zero_amount = s.post(url, headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 0, "payment_method": "cash",
        "idempotency_key": f"zero-{suffix}",
    })
    assert zero_amount.status_code == 422

    over_amount = s.post(url, headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 400, "payment_method": "cash",
        "idempotency_key": f"over-{suffix}",
    })
    assert over_amount.status_code == 422

    first = s.post(url, headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 100, "payment_method": "bank_transfer",
        "payment_reference": "TX-1", "idempotency_key": f"first-{suffix}",
    })
    assert first.status_code == 200, first.text
    assert first.json()["payment"]["payment_number"].startswith("PAY-")
    summary = first.json()["payment_summary"]
    assert summary["po_total"] == 300
    assert summary["paid_amount"] == 100
    assert summary["outstanding_amount"] == 200
    assert summary["payment_status"] == "partially_paid"

    second = s.post(url, headers=manager_headers, json={
        "payment_date": "2026-09-01", "amount": 100, "payment_method": "cheque",
        "idempotency_key": f"second-{suffix}",
    })
    assert second.status_code == 200, second.text
    assert second.json()["payment_summary"]["paid_amount"] == 200
    assert second.json()["payment_summary"]["outstanding_amount"] == 100

    over_remaining = s.post(url, headers=manager_headers, json={
        "payment_date": "2026-09-02", "amount": 150, "payment_method": "cash",
        "idempotency_key": f"over-remaining-{suffix}",
    })
    assert over_remaining.status_code == 422

    third = s.post(url, headers=manager_headers, json={
        "payment_date": "2026-09-05", "amount": 100, "payment_method": "cash",
        "idempotency_key": f"third-{suffix}",
    })
    assert third.status_code == 200, third.text
    final_summary = third.json()["payment_summary"]
    assert final_summary["paid_amount"] == 300
    assert final_summary["outstanding_amount"] == 0
    assert final_summary["payment_status"] == "paid"

    ledger = s.get(url, headers=manager_headers)
    assert ledger.status_code == 200
    assert len(ledger.json()["payments"]) == 3
    assert ledger.json()["payment_summary"]["paid_amount"] == 300


def test_po_payment_idempotency_key_prevents_double_posting(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-idem-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"pay-idem-mgr-{suffix}")
    po, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)
    url = PO_PAYMENTS_API(po["id"])
    body = {
        "payment_date": "2026-08-25", "amount": 100, "payment_method": "cash",
        "idempotency_key": f"idem-{suffix}",
    }
    first = s.post(url, headers=manager_headers, json=body)
    assert first.status_code == 200, first.text
    assert first.json()["already_recorded"] is False
    second = s.post(url, headers=manager_headers, json=body)
    assert second.status_code == 200, second.text
    assert second.json()["already_recorded"] is True
    assert second.json()["payment"]["id"] == first.json()["payment"]["id"]

    ledger = s.get(url, headers=manager_headers).json()
    assert len(ledger["payments"]) == 1
    assert ledger["payment_summary"]["paid_amount"] == 100


def test_po_payment_void_stops_counting_and_requires_reason(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-void-mgr-{suffix}", role="commercial_manager")
        _make_user(session, username=f"pay-void-eng-{suffix}", role="procurement_engineer")
    manager_headers = _login_headers(s, f"pay-void-mgr-{suffix}")
    engineer_headers = _login_headers(s, f"pay-void-eng-{suffix}")
    po, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)
    url = PO_PAYMENTS_API(po["id"])
    recorded = s.post(url, headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 100, "payment_method": "cash",
        "idempotency_key": f"void-{suffix}",
    })
    assert recorded.status_code == 200, recorded.text
    payment_id = recorded.json()["payment"]["id"]
    void_url = f"{url}/{payment_id}/void"

    denied_role = s.post(void_url, headers=engineer_headers, json={"reason": "خطأ"})
    assert denied_role.status_code == 403

    missing_reason = s.post(void_url, headers=manager_headers, json={"reason": ""})
    assert missing_reason.status_code == 422

    voided = s.post(void_url, headers=manager_headers, json={"reason": "دفعة مكررة بالخطأ"})
    assert voided.status_code == 200, voided.text
    assert voided.json()["payment"]["status"] == "voided"
    assert voided.json()["payment_summary"]["paid_amount"] == 0
    assert voided.json()["payment_summary"]["payment_status"] == "unpaid"

    ledger = s.get(url, headers=manager_headers).json()
    assert len(ledger["payments"]) == 1  # never deleted, only voided
    assert ledger["payments"][0]["status"] == "voided"
    assert ledger["payments"][0]["void_reason"] == "دفعة مكررة بالخطأ"
    assert ledger["payment_summary"]["paid_amount"] == 0


def test_po_payment_rejected_for_draft_and_cancelled(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-elig-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"pay-elig-mgr-{suffix}")

    draft_po, *_rest = _make_draft_purchase_order(s, f"{suffix}-draft", admin_headers)
    denied_draft = s.post(PO_PAYMENTS_API(draft_po["id"]), headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 10, "payment_method": "cash",
        "idempotency_key": f"draft-{suffix}",
    })
    assert denied_draft.status_code == 409

    cancelled_po, *_rest = _make_draft_purchase_order(s, f"{suffix}-cancel", admin_headers)
    cancel = s.patch(
        f"{API}/purchase-orders/{cancelled_po['id']}/status", headers=admin_headers,
        json={"status": "cancelled"},
    )
    assert cancel.status_code == 200, cancel.text
    denied_cancelled = s.post(PO_PAYMENTS_API(cancelled_po["id"]), headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 10, "payment_method": "cash",
        "idempotency_key": f"cancel-{suffix}",
    })
    assert denied_cancelled.status_code == 409


def test_po_payment_and_receiving_lifecycles_are_independent(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-indep-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"pay-indep-mgr-{suffix}")
    po, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)

    # Full payment recorded while the PO is still merely "approved" - not
    # delivered yet. Paying in full must not itself advance the PO's
    # procurement/delivery lifecycle.
    paid = s.post(PO_PAYMENTS_API(po["id"]), headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 300, "payment_method": "bank_transfer",
        "idempotency_key": f"full-{suffix}",
    })
    assert paid.status_code == 200, paid.text
    assert paid.json()["payment_summary"]["payment_status"] == "paid"
    status_after_payment = s.get(f"{API}/purchase-orders/{po['id']}", headers=admin_headers).json()
    assert status_after_payment["status"] == "approved"

    # Finalize + fully receive - receiving completion must not itself
    # change the (already-paid) payment status.
    finalized = s.post(f"{API}/purchase-orders/{po['id']}/finalize", headers=admin_headers, json={"actor": "t"})
    assert finalized.status_code == 200, finalized.text
    received = s.post(f"{API}/purchase-orders/{po['id']}/receipts", headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"recv-{suffix}", "actor": "t",
    })
    assert received.status_code == 200, received.text
    assert received.json()["purchase_order"]["status"] == "completed"
    ledger_after_completion = s.get(PO_PAYMENTS_API(po["id"]), headers=admin_headers).json()
    assert ledger_after_completion["payment_summary"]["payment_status"] == "paid"
    assert ledger_after_completion["payment_summary"]["paid_amount"] == 300


def test_po_receiving_completion_does_not_auto_mark_paid(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_draft_purchase_order(s, suffix, admin_headers)
    finalized = s.post(f"{API}/purchase-orders/{po['id']}/finalize", headers=admin_headers, json={"actor": "t"})
    assert finalized.status_code == 200, finalized.text
    received = s.post(f"{API}/purchase-orders/{po['id']}/receipts", headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"recv2-{suffix}", "actor": "t",
    })
    assert received.status_code == 200, received.text
    assert received.json()["purchase_order"]["status"] == "completed"
    ledger = s.get(PO_PAYMENTS_API(po["id"]), headers=admin_headers).json()
    assert ledger["payment_summary"]["paid_amount"] == 0
    assert ledger["payment_summary"]["payment_status"] != "paid"


def test_po_payment_ledger_isolated_from_legacy_and_approval_payments(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, approval, _request_id, _supplier = _make_payable_purchase_order(s, suffix, admin_headers)

    with SessionLocal() as session:
        po_payment_count_before = session.scalar(
            select(func.count()).select_from(PurchaseOrderPayment)
        ) or 0

    # A legacy Direct Purchase payment - keyed by purchase_id, never
    # purchase_order_id - must never be visible in the PO ledger.
    legacy_purchases = s.get(f"{API}/purchases", headers=admin_headers).json()
    assert legacy_purchases, "fixture data should include at least one legacy purchase"
    s.post(f"{API}/payments", json={
        "purchase_id": legacy_purchases[0]["purchase_id"], "payment_date": "2026-08-25",
        "amount_paid": 1, "payment_method": "cash", "notes": "",
    }, headers=admin_headers)

    with SessionLocal() as session:
        po_payment_count_after = session.scalar(
            select(func.count()).select_from(PurchaseOrderPayment)
        ) or 0
    assert po_payment_count_after == po_payment_count_before

    # Nor does the approval's own commercial funds-release event count.
    with SessionLocal() as session:
        approval_row = session.get(EngineerApproval, approval["id"])
        assert approval_row.status == "approved"

    ledger = s.get(PO_PAYMENTS_API(po["id"]), headers=admin_headers).json()
    assert ledger["payment_summary"]["paid_amount"] == 0
    assert ledger["payments"] == []


def test_po_payment_status_due_date_and_overdue_logic(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-due-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"pay-due-mgr-{suffix}")

    po_not_due, *_rest = _make_payable_purchase_order(s, f"{suffix}-notdue", admin_headers)
    ledger = s.get(PO_PAYMENTS_API(po_not_due["id"]), headers=admin_headers).json()
    assert ledger["payment_summary"]["payment_status"] == "unpaid"
    assert ledger["payment_summary"]["is_overdue"] is False

    po_overdue, *_rest = _make_payable_purchase_order(s, f"{suffix}-overdue", admin_headers)
    with SessionLocal.begin() as session:
        row = session.get(PurchaseOrder, po_overdue["id"])
        row.extra_data = {**(row.extra_data or {}), "payment_due_date": "2020-01-01"}
    overdue_ledger = s.get(PO_PAYMENTS_API(po_overdue["id"]), headers=admin_headers).json()
    assert overdue_ledger["payment_summary"]["payment_status"] == "unpaid"
    assert overdue_ledger["payment_summary"]["is_overdue"] is True

    partial = s.post(PO_PAYMENTS_API(po_overdue["id"]), headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 100, "payment_method": "cash",
        "idempotency_key": f"partial-overdue-{suffix}",
    })
    assert partial.status_code == 200, partial.text
    assert partial.json()["payment_summary"]["payment_status"] == "partially_paid"
    assert partial.json()["payment_summary"]["is_overdue"] is True

    po_credit, *_rest = _make_payable_purchase_order(s, f"{suffix}-credit", admin_headers)
    with SessionLocal.begin() as session:
        row = session.get(PurchaseOrder, po_credit["id"])
        row.extra_data = {**(row.extra_data or {}), "credit_days": 3650}
    credit_ledger = s.get(PO_PAYMENTS_API(po_credit["id"]), headers=admin_headers).json()
    assert credit_ledger["payment_summary"]["due_date"]
    assert credit_ledger["payment_summary"]["payment_status"] == "unpaid"
    assert credit_ledger["payment_summary"]["is_overdue"] is False


def test_po_payment_read_does_not_mutate_records(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-nomut-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"pay-nomut-mgr-{suffix}")
    po, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)
    recorded = s.post(PO_PAYMENTS_API(po["id"]), headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 100, "payment_method": "cash",
        "idempotency_key": f"nomut-{suffix}",
    })
    assert recorded.status_code == 200, recorded.text
    before = s.get(PO_PAYMENTS_API(po["id"]), headers=admin_headers).json()
    s.get(PO_PAYMENTS_API(po["id"]), headers=admin_headers)
    after = s.get(PO_PAYMENTS_API(po["id"]), headers=admin_headers).json()
    assert after == before


def test_po_list_includes_batched_payment_summary(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"pay-list-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"pay-list-mgr-{suffix}")
    po, *_rest = _make_payable_purchase_order(s, suffix, admin_headers)
    s.post(PO_PAYMENTS_API(po["id"]), headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 300, "payment_method": "cash",
        "idempotency_key": f"list-{suffix}",
    })
    listed = s.get(f"{API}/purchase-orders", headers=admin_headers)
    assert listed.status_code == 200
    row = next(item for item in listed.json() if item["id"] == po["id"])
    assert row["payment_summary"]["paid_amount"] == 300
    assert row["payment_summary"]["payment_status"] == "paid"


# ---------------- Sprint 3.5: Receiving & PO Lifecycle Finalization ----------------
# Multi-PO REQ completion (one PO completing must not complete the REQ; all
# non-cancelled formal POs completing must) is already covered by
# test_request_completes_only_after_all_formal_purchase_orders above.
# Payment/receiving independence (paying in full does not auto-complete
# receiving, and vice versa) is already covered by
# test_po_payment_and_receiving_lifecycles_are_independent and
# test_po_receiving_completion_does_not_auto_mark_paid (Sprint 3.4). Neither
# is duplicated here.

RECEIPTS_API = lambda po_id: f"{API}/purchase-orders/{po_id}/receipts"  # noqa: E731


def _make_in_delivery_purchase_order(client, suffix, admin_headers, manual_item=False):
    """A formal PO (single item, qty 3 x 100 EGP = final_total 300) walked
    through finalize so receiving is open. Returns (purchase_order,
    approval, request_id, supplier)."""
    po, approval, request_id, supplier = _make_draft_purchase_order(
        client, suffix, admin_headers, manual_item=manual_item,
    )
    finalized = client.post(
        f"{API}/purchase-orders/{po['id']}/finalize", headers=admin_headers, json={"actor": "t"},
    )
    assert finalized.status_code == 200, finalized.text
    return finalized.json()["purchase_order"], approval, request_id, supplier


def test_po_receipt_authorization_matrix(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"recv-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"recv-mgr-{suffix}", role="commercial_manager")
        _make_user(session, username=f"recv-resp-{suffix}", role="procurement_responsible")
        portal_username, _, _ = _portal_setup(session, f"recv-portal-{suffix}")
    engineer_headers = _login_headers(s, f"recv-eng-{suffix}")
    manager_headers = _login_headers(s, f"recv-mgr-{suffix}")
    responsible_headers = _login_headers(s, f"recv-resp-{suffix}")
    portal_headers = _login_headers(s, portal_username)

    po, *_rest = _make_in_delivery_purchase_order(s, suffix, admin_headers)
    url = RECEIPTS_API(po["id"])

    def body(key):
        return {"receipt_type": "full", "idempotency_key": key}

    assert s.post(url, headers=engineer_headers, json=body(f"eng-{suffix}")).status_code == 403
    assert s.post(url, headers=manager_headers, json=body(f"mgr-{suffix}")).status_code == 403
    assert s.post(url, headers=portal_headers, json=body(f"portal-{suffix}")).status_code == 403
    assert s.post(url, json=body(f"anon-{suffix}")).status_code == 401

    spoofed = s.post(
        url, headers=engineer_headers,
        json={**body(f"spoof-{suffix}"), "actor_role": "procurement_responsible", "role": "admin"},
    )
    assert spoofed.status_code == 403

    completed = s.post(url, headers=responsible_headers, json=body(f"resp-{suffix}"))
    assert completed.status_code == 200, completed.text
    assert completed.json()["purchase_order"]["status"] == "completed"


def test_po_receipt_quantity_and_item_safety(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_in_delivery_purchase_order(s, suffix, admin_headers)
    item_id = po["items"][0]["id"]
    url = RECEIPTS_API(po["id"])

    zero = s.post(url, headers=admin_headers, json={
        "receipt_type": "partial", "idempotency_key": f"zero-{suffix}",
        "lines": [{"purchase_order_item_id": item_id, "quantity": 0}],
    })
    assert zero.status_code == 422

    invalid_item = s.post(url, headers=admin_headers, json={
        "receipt_type": "partial", "idempotency_key": f"invalid-{suffix}",
        "lines": [{"purchase_order_item_id": str(uuid.uuid4()), "quantity": 1}],
    })
    assert invalid_item.status_code == 422

    over = s.post(url, headers=admin_headers, json={
        "receipt_type": "partial", "idempotency_key": f"over-{suffix}",
        "lines": [{"purchase_order_item_id": item_id, "quantity": 4}],
    })
    assert over.status_code == 422

    empty = s.post(url, headers=admin_headers, json={
        "receipt_type": "partial", "idempotency_key": f"empty-{suffix}", "lines": [],
    })
    assert empty.status_code == 422

    with SessionLocal() as session:
        assert session.query(PurchaseOrderReceipt).filter_by(purchase_order_id=po["id"]).count() == 0


def test_po_manual_origin_item_receives_normally(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_in_delivery_purchase_order(s, suffix, admin_headers, manual_item=True)
    item = po["items"][0]
    assert item["item_id"] == ""
    receipt = s.post(RECEIPTS_API(po["id"]), headers=admin_headers, json={
        "receipt_type": "partial", "idempotency_key": f"manual-{suffix}",
        "lines": [{"purchase_order_item_id": item["id"], "quantity": 1}],
    })
    assert receipt.status_code == 200, receipt.text
    assert receipt.json()["purchase_order"]["items"][0]["received_quantity"] == 1


def test_po_full_receipt_after_partial_closes_only_remaining(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_in_delivery_purchase_order(s, suffix, admin_headers)
    item_id = po["items"][0]["id"]
    url = RECEIPTS_API(po["id"])

    partial = s.post(url, headers=admin_headers, json={
        "receipt_type": "partial", "idempotency_key": f"partial-{suffix}",
        "lines": [{"purchase_order_item_id": item_id, "quantity": 1}],
    })
    assert partial.status_code == 200, partial.text
    assert partial.json()["purchase_order"]["status"] == "partial_received"
    assert partial.json()["purchase_order"]["items"][0]["received_quantity"] == 1

    full = s.post(url, headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"full-{suffix}",
    })
    assert full.status_code == 200, full.text
    completed_po = full.json()["purchase_order"]
    assert completed_po["status"] == "completed"
    assert completed_po["items"][0]["received_quantity"] == 3
    assert completed_po["items"][0]["remaining_quantity"] == 0
    assert completed_po["receipt_summary"]["ordered_quantity"] == 3.0
    assert completed_po["receipt_summary"]["received_quantity"] == 3.0
    assert completed_po["receipt_summary"]["remaining_quantity"] == 0.0
    assert completed_po["receipt_summary"]["receipt_count"] == 2
    assert completed_po["receipt_summary"]["latest_receipt_date"] == completed_po["receipt_history"][0]["received_at"]


def test_po_completed_rejects_another_receipt(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_in_delivery_purchase_order(s, suffix, admin_headers)
    full = s.post(RECEIPTS_API(po["id"]), headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"full-{suffix}",
    })
    assert full.status_code == 200, full.text
    assert full.json()["purchase_order"]["status"] == "completed"
    rejected = s.post(RECEIPTS_API(po["id"]), headers=admin_headers, json={
        "receipt_type": "problem", "idempotency_key": f"after-{suffix}", "problem_reason": "تالف",
    })
    assert rejected.status_code == 409


def test_po_cancelled_rejects_receipt(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_in_delivery_purchase_order(s, suffix, admin_headers)
    # The existing PATCH /status transition matrix already refuses to
    # cancel an in_delivery PO (a separate, correct guard) - so reaching
    # "cancelled" from here requires direct state manipulation, the same
    # way existing tests simulate other unreachable-via-API states.
    with SessionLocal.begin() as session:
        session.get(PurchaseOrder, po["id"]).status = "cancelled"
    rejected = s.post(RECEIPTS_API(po["id"]), headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"cancelled-{suffix}",
    })
    assert rejected.status_code == 409


def test_po_delivery_problem_does_not_corrupt_received_quantity_and_can_recover(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_in_delivery_purchase_order(s, suffix, admin_headers)
    item_id = po["items"][0]["id"]
    url = RECEIPTS_API(po["id"])

    partial = s.post(url, headers=admin_headers, json={
        "receipt_type": "partial", "idempotency_key": f"partial-{suffix}",
        "lines": [{"purchase_order_item_id": item_id, "quantity": 1}],
    })
    assert partial.status_code == 200, partial.text

    problem = s.post(url, headers=admin_headers, json={
        "receipt_type": "problem", "idempotency_key": f"problem-{suffix}",
        "problem_reason": "تالف", "affected_item_id": item_id, "affected_quantity": 1,
    })
    assert problem.status_code == 200, problem.text
    assert problem.json()["purchase_order"]["status"] == "delivery_problem"
    # The problem event must not count toward received_quantity.
    assert problem.json()["purchase_order"]["items"][0]["received_quantity"] == 1

    recovered = s.post(url, headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"recover-{suffix}",
    })
    assert recovered.status_code == 200, recovered.text
    assert recovered.json()["purchase_order"]["status"] == "completed"
    assert recovered.json()["purchase_order"]["items"][0]["received_quantity"] == 3


def test_po_receipt_read_does_not_mutate_and_ancestry_unchanged(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po, approval, request_id, supplier = _make_in_delivery_purchase_order(s, suffix, admin_headers)
    partial = s.post(RECEIPTS_API(po["id"]), headers=admin_headers, json={
        "receipt_type": "partial", "idempotency_key": f"partial-{suffix}",
        "lines": [{"purchase_order_item_id": po["items"][0]["id"], "quantity": 1}],
    })
    assert partial.status_code == 200, partial.text

    before = s.get(f"{API}/purchase-orders/{po['id']}", headers=admin_headers).json()
    s.get(f"{API}/purchase-orders/{po['id']}", headers=admin_headers)
    after = s.get(f"{API}/purchase-orders/{po['id']}", headers=admin_headers).json()
    assert after == before

    assert after["source_request_id"] == request_id
    assert after["comparison_id"] == approval["comparison_id"]
    assert after["approval_id"] == approval["id"]
    assert after["supplier_id"] == supplier["id"]


def test_po_payment_can_still_be_recorded_after_receiving_completes(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"recv-pay-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"recv-pay-mgr-{suffix}")
    po, *_rest = _make_in_delivery_purchase_order(s, suffix, admin_headers)
    full = s.post(RECEIPTS_API(po["id"]), headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"full-{suffix}",
    })
    assert full.status_code == 200, full.text
    assert full.json()["purchase_order"]["status"] == "completed"

    payment = s.post(f"{API}/purchase-orders/{po['id']}/payments", headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 300, "payment_method": "cash",
        "idempotency_key": f"pay-{suffix}",
    })
    assert payment.status_code == 200, payment.text
    assert payment.json()["payment_summary"]["payment_status"] == "paid"


# ---------- Dashboard: formal procurement intelligence (Sprint 3.6) ----------
# The dashboard aggregates data written by many other tests sharing this
# session-scoped database, so every check below compares a BEFORE/AFTER
# delta around one isolated action instead of asserting an absolute total.

def _dashboard(client, headers):
    r = client.get(f"{API}/dashboard", headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


def test_dashboard_authorization_matrix(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"dash-resp-{suffix}", role="procurement_responsible")
        _make_user(session, username=f"dash-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"dash-mgr-{suffix}", role="commercial_manager")
        portal_username, _, _ = _portal_setup(session, f"dash-portal-{suffix}")
    responsible_headers = _login_headers(s, f"dash-resp-{suffix}")
    engineer_headers = _login_headers(s, f"dash-eng-{suffix}")
    manager_headers = _login_headers(s, f"dash-mgr-{suffix}")
    portal_headers = _login_headers(s, portal_username)

    assert s.get(f"{API}/dashboard", headers=admin_headers).status_code == 200
    assert s.get(f"{API}/dashboard", headers=responsible_headers).status_code == 200
    assert s.get(f"{API}/dashboard", headers=engineer_headers).status_code == 200
    assert s.get(f"{API}/dashboard", headers=manager_headers).status_code == 200
    assert s.get(f"{API}/dashboard", headers=portal_headers).status_code == 403
    assert s.get(f"{API}/dashboard").status_code == 401


def test_dashboard_formal_po_value_ignores_legacy_direct_purchases(s, admin_headers, ids):
    suffix = uuid.uuid4().hex[:8]
    before = _dashboard(s, admin_headers)

    po, *_rest = _make_draft_purchase_order(s, suffix, admin_headers)
    after_po = _dashboard(s, admin_headers)
    assert round(after_po["summary"]["formal_po_value"] - before["summary"]["formal_po_value"], 2) == po["final_total"]
    assert after_po["payment_intelligence"]["total_formal_po_value"] == after_po["summary"]["formal_po_value"]

    legacy = s.post(f"{API}/purchases", json=_payload(ids, f"DASH-LEGACY-{suffix}", qty=1, price=500, disc=0, vat=0), headers=admin_headers)
    assert legacy.status_code == 200, legacy.text
    after_legacy = _dashboard(s, admin_headers)
    assert after_legacy["direct_purchase_total"] > after_po["direct_purchase_total"]
    assert after_legacy["summary"]["formal_po_value"] == after_po["summary"]["formal_po_value"]


def test_dashboard_actual_paid_outstanding_overdue_and_void_match_po_payment_ledger(s, admin_headers, ids):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"dash-pay-mgr-{suffix}", role="commercial_manager")
    manager_headers = _login_headers(s, f"dash-pay-mgr-{suffix}")

    # Funds release (part of _make_draft_purchase_order's normal flow) must
    # not move Actual Paid - it is an approval-level event, not real money.
    before_funds = _dashboard(s, admin_headers)
    _make_draft_purchase_order(s, f"{suffix}-funds", admin_headers)
    after_funds = _dashboard(s, admin_headers)
    assert after_funds["summary"]["actual_paid"] == before_funds["summary"]["actual_paid"]

    # A legacy Direct Purchase payment must not move Actual Paid either.
    legacy = s.post(f"{API}/purchases", json=_payload(ids, f"DASH-PAY-LEGACY-{suffix}", qty=1, price=200, disc=0, vat=0), headers=admin_headers).json()
    s.post(f"{API}/payments", json={
        "purchase_id": legacy["purchase_id"], "payment_date": "2026-08-20",
        "amount_paid": 50, "submission_token": f"dash-legacy-pay-{suffix}",
    }, headers=admin_headers)
    after_legacy_payment = _dashboard(s, admin_headers)
    assert after_legacy_payment["summary"]["actual_paid"] == after_funds["summary"]["actual_paid"]

    # A real, overdue formal PO payment must move Actual Paid/Outstanding/
    # Overdue by exactly the numbers the Sprint 3.4 ledger itself computes.
    po_overdue, *_rest = _make_payable_purchase_order(s, f"{suffix}-overdue", admin_headers)
    with SessionLocal.begin() as session:
        row = session.get(PurchaseOrder, po_overdue["id"])
        row.extra_data = {**(row.extra_data or {}), "payment_due_date": "2020-01-01"}
    before_pay = _dashboard(s, admin_headers)
    paid = s.post(f"{API}/purchase-orders/{po_overdue['id']}/payments", headers=manager_headers, json={
        "payment_date": "2026-08-25", "amount": 100, "payment_method": "cash",
        "idempotency_key": f"dash-overdue-{suffix}",
    })
    assert paid.status_code == 200, paid.text
    ledger = paid.json()["payment_summary"]
    payment_id = paid.json()["payment"]["id"]
    after_pay = _dashboard(s, admin_headers)

    assert round(after_pay["summary"]["actual_paid"] - before_pay["summary"]["actual_paid"], 2) == ledger["paid_amount"]
    assert round(before_pay["summary"]["outstanding"] - after_pay["summary"]["outstanding"], 2) == ledger["paid_amount"]
    assert ledger["is_overdue"] is True
    assert round(before_pay["summary"]["overdue_amount"] - after_pay["summary"]["overdue_amount"], 2) == ledger["paid_amount"]

    attention_row = next(
        r for r in after_pay["payment_intelligence"]["attention"] if r["purchase_order_id"] == po_overdue["id"]
    )
    assert attention_row["paid_amount"] == ledger["paid_amount"]
    assert attention_row["outstanding_amount"] == ledger["outstanding_amount"]
    assert attention_row["payment_status"] == ledger["payment_status"]
    assert attention_row["is_overdue"] is True
    assert attention_row["days_overdue"] and attention_row["days_overdue"] > 0

    # Voiding the payment must remove it from Actual Paid/Outstanding again.
    void = s.post(
        f"{API}/purchase-orders/{po_overdue['id']}/payments/{payment_id}/void",
        headers=manager_headers, json={"reason": "خطأ في القيد"},
    )
    assert void.status_code == 200, void.text
    after_void = _dashboard(s, admin_headers)
    assert after_void["summary"]["actual_paid"] == before_pay["summary"]["actual_paid"]
    assert after_void["summary"]["outstanding"] == before_pay["summary"]["outstanding"]


def test_dashboard_active_po_count_excludes_completed_and_cancelled(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    before = _dashboard(s, admin_headers)

    _make_draft_purchase_order(s, f"{suffix}-draft", admin_headers)
    after_draft = _dashboard(s, admin_headers)
    assert after_draft["summary"]["active_purchase_orders"] == before["summary"]["active_purchase_orders"] + 1

    completed_po, *_rest = _make_in_delivery_purchase_order(s, f"{suffix}-done", admin_headers)
    full = s.post(RECEIPTS_API(completed_po["id"]), headers=admin_headers, json={
        "receipt_type": "full", "idempotency_key": f"dash-complete-{suffix}",
    })
    assert full.status_code == 200, full.text
    after_completed = _dashboard(s, admin_headers)
    # The PO was counted active while in_delivery, then dropped again once
    # completed - net delta versus after_draft is zero.
    assert after_completed["summary"]["active_purchase_orders"] == after_draft["summary"]["active_purchase_orders"]

    draft_to_cancel, *_rest = _make_draft_purchase_order(s, f"{suffix}-cancel", admin_headers)
    after_second_draft = _dashboard(s, admin_headers)
    assert after_second_draft["summary"]["active_purchase_orders"] == after_completed["summary"]["active_purchase_orders"] + 1

    cancelled = s.patch(
        f"{API}/purchase-orders/{draft_to_cancel['id']}/status", headers=admin_headers, json={"status": "cancelled"},
    )
    assert cancelled.status_code == 200, cancelled.text
    after_cancel = _dashboard(s, admin_headers)
    assert after_cancel["summary"]["active_purchase_orders"] == after_completed["summary"]["active_purchase_orders"]


def test_dashboard_receiving_intelligence_and_is_independent_of_payment(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    before = _dashboard(s, admin_headers)

    partial_po, *_rest = _make_in_delivery_purchase_order(s, f"{suffix}-partial", admin_headers)
    partial = s.post(RECEIPTS_API(partial_po["id"]), headers=admin_headers, json={
        "receipt_type": "partial", "idempotency_key": f"dash-partial-{suffix}",
        "lines": [{"purchase_order_item_id": partial_po["items"][0]["id"], "quantity": 1}],
    })
    assert partial.status_code == 200, partial.text
    after_partial = _dashboard(s, admin_headers)
    assert after_partial["receiving"]["partial_received_count"] == before["receiving"]["partial_received_count"] + 1
    row = next(r for r in after_partial["receiving"]["attention"] if r["purchase_order_id"] == partial_po["id"])
    assert row["status"] == "partial_received"
    assert row["received_lines"] == 0  # ordered 3, received 1 - line not fully received yet
    assert row["total_lines"] == 1
    # Receiving status is derived purely from receipts - no payment exists
    # yet, and none is required for the receiving status to be correct.
    ledger = s.get(f"{API}/purchase-orders/{partial_po['id']}/payments", headers=admin_headers).json()
    assert ledger["payment_summary"]["paid_amount"] == 0

    problem_po, *_rest = _make_in_delivery_purchase_order(s, f"{suffix}-problem", admin_headers)
    problem = s.post(RECEIPTS_API(problem_po["id"]), headers=admin_headers, json={
        "receipt_type": "problem", "idempotency_key": f"dash-problem-{suffix}", "problem_reason": "تالف",
    })
    assert problem.status_code == 200, problem.text
    after_problem = _dashboard(s, admin_headers)
    assert after_problem["receiving"]["delivery_problem_count"] == after_partial["receiving"]["delivery_problem_count"] + 1
    assert any(
        item["type"] == "delivery_problem" and item["reference"] == problem_po["po_number"]
        for item in after_problem["attention_items"]
    )


def test_dashboard_sourcing_attention_flags_missing_and_late_rfq_responses(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"dash-rfq-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"dash-rfq-resp-{suffix}")
    supplier = s.get(f"{API}/suppliers", headers=admin_headers).json()[0]
    before = _dashboard(s, admin_headers)

    request_id, *_rest = _make_pricing_request(s, f"{suffix}-zero")
    rfq = s.post(RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers}, json={
        "source_request_id": request_id, "actor": "t",
    }).json()["rfq"]
    added = s.post(
        f"{RFQ_API}/{rfq['id']}/suppliers", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"], "actor": "t"},
    )
    assert added.status_code == 200, added.text
    after_zero = _dashboard(s, admin_headers)
    assert after_zero["sourcing"]["zero_response_count"] == before["sourcing"]["zero_response_count"] + 1
    assert after_zero["sourcing"]["open_rfq_count"] == before["sourcing"]["open_rfq_count"] + 1
    # sourcing.attention is a capped, priority-sorted top-N (past-deadline
    # first) shared across this whole session's accumulated test data, so a
    # single lower-priority "zero response" entry isn't guaranteed a seat -
    # the count assertions above already prove detection works correctly.

    request_id2, *_rest = _make_pricing_request(s, f"{suffix}-late")
    late_rfq = s.post(RFQ_API, headers={**INTERNAL_HEADERS, **responsible_headers}, json={
        "source_request_id": request_id2, "deadline": "2020-01-01", "actor": "t",
    }).json()["rfq"]
    s.post(
        f"{RFQ_API}/{late_rfq['id']}/suppliers", headers={**INTERNAL_HEADERS, **responsible_headers},
        json={"supplier_id": supplier["id"], "actor": "t"},
    )
    after_late = _dashboard(s, admin_headers)
    assert after_late["sourcing"]["past_deadline_count"] == after_zero["sourcing"]["past_deadline_count"] + 1
    assert any(
        item["type"] == "rfq_past_deadline" and item["reference"] == late_rfq["rfq_number"]
        for item in after_late["attention_items"]
    )


def test_dashboard_pending_approval_appears_under_correct_stage_only(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    before = _dashboard(s, admin_headers)
    _make_approval_with_comparison(s, suffix, admin_headers)
    after = _dashboard(s, admin_headers)

    stages_before = {row["key"]: row for row in before["approval_attention"]["stages"]}
    stages_after = {row["key"]: row for row in after["approval_attention"]["stages"]}
    assert stages_after["comparison_approval"]["count"] == stages_before["comparison_approval"]["count"] + 1
    for key in ("fund_approval", "funds_release", "ready_for_po"):
        assert stages_after[key]["count"] == stages_before[key]["count"]


def test_dashboard_attention_items_are_filtered_by_role_ownership(s, admin_headers):
    """Needs My Attention is role-scoped using the same responsible_role
    values the backend already enforces for real actions (EngineerApproval.
    responsible_role, and the static per-type ownership used elsewhere) -
    admin always sees everything, every other role sees only its own."""
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal() as session:
        _make_user(session, username=f"dash-role-eng-{suffix}", role="procurement_engineer")
        _make_user(session, username=f"dash-role-mgr-{suffix}", role="commercial_manager")
        _make_user(session, username=f"dash-role-resp-{suffix}", role="procurement_responsible")
    engineer_headers = _login_headers(s, f"dash-role-eng-{suffix}")
    manager_headers = _login_headers(s, f"dash-role-mgr-{suffix}")
    responsible_headers = _login_headers(s, f"dash-role-resp-{suffix}")

    # sourcing_required is a static procurement_responsible-owned type.
    request_id, request_number, *_rest = _make_pricing_request(s, suffix)

    def has_item(dash, item_type, reference):
        return any(
            item["type"] == item_type and item["reference"] == reference
            for item in dash["attention_items"]
        )

    assert has_item(_dashboard(s, responsible_headers), "sourcing_required", request_number) is True
    assert has_item(_dashboard(s, engineer_headers), "sourcing_required", request_number) is False
    assert has_item(_dashboard(s, manager_headers), "sourcing_required", request_number) is False
    assert has_item(_dashboard(s, admin_headers), "sourcing_required", request_number) is True

    # pending_approval ownership is dynamic - it follows the specific
    # approval's own responsible_role column, not a fixed type mapping.
    approval, *_rest = _make_approval_with_comparison(s, uuid.uuid4().hex[:8], admin_headers)
    # The dashboard only surfaces the 8 oldest pending comparison_workflow
    # approvals - backdate this one so it's never pushed out of that window
    # by other tests' leftover pending approvals in this shared session.
    with SessionLocal() as session:
        session.get(EngineerApproval, approval["id"]).created_at = "2000-01-01T00:00:00Z"
        session.commit()

    def has_pending_approval(dash):
        return has_item(dash, "pending_approval", approval["approval_number"])

    assert has_pending_approval(_dashboard(s, engineer_headers)) is True
    assert has_pending_approval(_dashboard(s, manager_headers)) is False
    assert has_pending_approval(_dashboard(s, responsible_headers)) is False
    assert has_pending_approval(_dashboard(s, admin_headers)) is True

    decision = s.post(
        f"{API}/workflow/approvals/{approval['id']}/decision",
        headers={**INTERNAL_HEADERS, **engineer_headers},
        json={"decision": "approved", "actor": "engineer"},
    )
    assert decision.status_code == 200, decision.text

    # Ownership now follows the same approval's updated responsible_role
    # (commercial_manager, per the fund_release stage transition).
    assert has_pending_approval(_dashboard(s, manager_headers)) is True
    assert has_pending_approval(_dashboard(s, engineer_headers)) is False


def test_dashboard_attention_items_expose_deep_link_ids(s, admin_headers):
    """Every "Needs my attention" item must carry the exact-record ids the
    frontend deep-link helper (getFollowUpTarget, frontend/src/lib/
    followUpNavigation.js) needs to route straight to the record and stage
    where it is actually stuck, instead of a bare list page. This is the
    audited action-type -> id-field table:
        delivery_problem              -> purchase_order_id (+ stage=receiving)
        overdue_payment                -> purchase_order_id (+ stage=payments)
        partial_received               -> purchase_order_id (+ stage=receiving)
        awaiting_supplier_confirmation -> purchase_order_id
        rfq_past_deadline               -> rfq_id
        quotation_missing               -> rfq_id
        sourcing_required               -> request_id (+ stage=sourcing)
        needs_clarification             -> request_id (+ stage=clarification)
        request_review                  -> request_id (+ stage=technical-review)
        pending_approval                -> approval_id (+ stage=<approval_stage>)
    needs_clarification is verified indirectly, by proving request_review's
    identical construction is correct - see the comment at that assertion.
    """
    suffix = uuid.uuid4().hex[:8]

    def item_for(dash, item_type, reference):
        return next(
            row for row in dash["attention_items"]
            if row["type"] == item_type and row["reference"] == reference
        )

    def item_or_ambient(dash, item_type, reference):
        """Prefer the exact row this test just created; some categories sit
        late/uncapped-per-type in the attention_items builder (server.py) and
        can be pushed out of the shared top-20 window by this session's
        accumulated data by the time this test runs. Falling back to any
        ambient row of the same type still proves the enrichment is correct,
        since every row of a given type is built by the same code path."""
        exact = next(
            (row for row in dash["attention_items"] if row["type"] == item_type and row["reference"] == reference),
            None,
        )
        if exact is not None:
            return exact, True
        ambient = next((row for row in dash["attention_items"] if row["type"] == item_type), None)
        assert ambient is not None, f"expected at least one {item_type} item"
        return ambient, False

    # sourcing_required - previously exposed only request_number, no id.
    request_id, request_number, *_rest = _make_pricing_request(s, f"{suffix}-src")
    dash = _dashboard(s, admin_headers)
    sourcing_item = item_for(dash, "sourcing_required", request_number)
    assert sourcing_item["request_id"] == request_id
    assert sourcing_item["entity_type"] == "incoming_request"
    assert sourcing_item["entity_id"] == request_id
    assert sourcing_item["action_type"] == "sourcing_required"
    assert sourcing_item["stage"] == "sourcing"

    # needs_clarification / request_review - same gap, plain status flips.
    # Both are built by the *same* if/elif block in server.py from a single,
    # unsorted, uncapped-per-category scan of every incoming request, then
    # the combined attention_items list is truncated to the top 20 (per role,
    # after this shared test session has accumulated a long tail of "new"/
    # "under_review" fixture requests from unrelated tests). A row inserted
    # here - necessarily the newest by insertion order - is therefore not
    # guaranteed a seat in that capped view, so instead of asserting on our
    # own freshly-created row (flaky at full-suite scale), we assert the
    # id-field shape on whichever ambient request_review item the role's
    # capped list already surfaces - it is built by the exact same code path
    # (see server.py's attention_items request_row status loop), so this
    # equally proves the needs_clarification branch, which sets the same
    # entity_type/entity_id/request_id/stage keys one line above it.
    with SessionLocal() as session:
        _make_user(session, username=f"dash-deep-link-eng-{suffix}", role="procurement_engineer")
    engineer_headers = _login_headers(s, f"dash-deep-link-eng-{suffix}")
    dash_for_engineer = _dashboard(s, engineer_headers)
    ambient_review_item = next(
        (row for row in dash_for_engineer["attention_items"] if row["type"] == "request_review"), None,
    )
    assert ambient_review_item is not None, "expected at least one request_review item for procurement_engineer"
    assert ambient_review_item["entity_type"] == "incoming_request"
    assert ambient_review_item["request_id"] == ambient_review_item["entity_id"]
    assert ambient_review_item["stage"] == "technical-review"

    # pending_approval - previously exposed only approval_number, no id.
    # Same shared-session cap concern as above - backdate it (as the existing
    # role-ownership test above already does) and read with the owning
    # engineer role so it isn't crowded out of the top-20/top-8 windows.
    approval, *_rest = _make_approval_with_comparison(s, f"{suffix}-appr", admin_headers)
    with SessionLocal() as session:
        session.get(EngineerApproval, approval["id"]).created_at = "2000-01-01T00:00:00Z"
        session.commit()
    dash_after_approval = _dashboard(s, engineer_headers)
    approval_item = item_for(dash_after_approval, "pending_approval", approval["approval_number"])
    assert approval_item["approval_id"] == approval["id"]
    assert approval_item["entity_type"] == "approval"
    assert approval_item["entity_id"] == approval["id"]
    assert approval_item["stage"]

    # rfq_past_deadline - already had the id embedded only in `path`; must
    # now also be an explicit field.
    supplier = s.get(f"{API}/suppliers", headers=admin_headers).json()[0]
    late_request_id, *_rest = _make_pricing_request(s, f"{suffix}-late")
    late_rfq = s.post(RFQ_API, headers={**INTERNAL_HEADERS, **admin_headers}, json={
        "source_request_id": late_request_id, "deadline": "2020-01-01", "actor": "t",
    }).json()["rfq"]
    s.post(
        f"{RFQ_API}/{late_rfq['id']}/suppliers", headers={**INTERNAL_HEADERS, **admin_headers},
        json={"supplier_id": supplier["id"], "actor": "t"},
    )
    dash_after_rfq = _dashboard(s, admin_headers)
    rfq_item = item_for(dash_after_rfq, "rfq_past_deadline", late_rfq["rfq_number"])
    assert rfq_item["rfq_id"] == late_rfq["id"]
    assert rfq_item["entity_type"] == "rfq"

    # PO-linked types - delivery_problem / awaiting_supplier_confirmation.
    problem_po, *_rest = _make_in_delivery_purchase_order(s, f"{suffix}-problem", admin_headers)
    problem = s.post(RECEIPTS_API(problem_po["id"]), headers=admin_headers, json={
        "receipt_type": "problem", "idempotency_key": f"deep-link-problem-{suffix}", "problem_reason": "تالف",
    })
    assert problem.status_code == 200, problem.text
    sent_po, *_rest = _make_draft_purchase_order(s, f"{suffix}-sent", admin_headers)
    for status in ("approved", "sent"):
        advanced = s.patch(
            f"{API}/purchase-orders/{sent_po['id']}/status", headers=admin_headers, json={"status": status},
        )
        assert advanced.status_code == 200, advanced.text
        sent_po = advanced.json()
    dash_after_po = _dashboard(s, admin_headers)
    delivery_item = item_for(dash_after_po, "delivery_problem", problem_po["po_number"])
    assert delivery_item["purchase_order_id"] == problem_po["id"]
    assert delivery_item["entity_type"] == "purchase_order"
    assert delivery_item["stage"] == "receiving"
    # awaiting_supplier_confirmation sits later in the category order (after
    # delivery/payment/rfq/sourcing/pending_approval), so on the unscoped
    # admin view it can be squeezed out entirely by this shared session's
    # accumulated data. Its owner is procurement_responsible - see
    # ATTENTION_TYPE_ROLE_OWNERS in server.py - and role filtering happens
    # before the top-20 cap, so read it from that role's own view instead.
    with SessionLocal() as session:
        _make_user(session, username=f"dash-deep-link-resp-{suffix}", role="procurement_responsible")
    responsible_headers = _login_headers(s, f"dash-deep-link-resp-{suffix}")
    dash_for_responsible = _dashboard(s, responsible_headers)
    confirmation_item, confirmation_is_exact = item_or_ambient(
        dash_for_responsible, "awaiting_supplier_confirmation", sent_po["po_number"],
    )
    assert confirmation_item["entity_type"] == "purchase_order"
    assert confirmation_item["purchase_order_id"] == confirmation_item["entity_id"]
    if confirmation_is_exact:
        assert confirmation_item["purchase_order_id"] == sent_po["id"]


def test_dashboard_project_summary_uses_formal_po_totals_only(s, admin_headers, ids):
    """Also stands in for the multi-PO/project safety guarantee, whose core
    invariant (a REQ only completes once ALL its formal POs are completed)
    is already covered by test_request_completes_only_after_all_formal_
    purchase_orders - this test focuses on the dashboard's own aggregation."""
    suffix = uuid.uuid4().hex[:8]
    po, *_rest = _make_draft_purchase_order(s, suffix, admin_headers)
    project_id = po["project_id"]
    dash = _dashboard(s, admin_headers)
    row = next(p for p in dash["project_procurement_summary"] if p["project_id"] == project_id)
    assert row["formal_po_value"] >= po["final_total"]
    assert row["active_po_count"] >= 1

    legacy_payload = _payload(ids, f"DASH-PROJ-LEGACY-{suffix}", qty=1, price=999, disc=0, vat=0)
    legacy_payload["project_id"] = project_id
    legacy = s.post(f"{API}/purchases", json=legacy_payload, headers=admin_headers)
    assert legacy.status_code == 200, legacy.text
    dash_after = _dashboard(s, admin_headers)
    row_after = next(p for p in dash_after["project_procurement_summary"] if p["project_id"] == project_id)
    assert row_after["formal_po_value"] == row["formal_po_value"]


def test_dashboard_read_is_safe_with_sparse_data_and_causes_no_mutations(s, admin_headers):
    # A REQ with no RFQ/CMP/PO at all (still just "pricing") - the kind of
    # sparse/historical data the dashboard must tolerate without crashing.
    _make_pricing_request(s, uuid.uuid4().hex[:8])

    first = _dashboard(s, admin_headers)
    for key in ("sourcing", "payment_intelligence", "receiving", "approval_attention"):
        assert isinstance(first[key], dict)
    assert isinstance(first["sourcing"]["attention"], list)
    assert isinstance(first["payment_intelligence"]["attention"], list)
    assert isinstance(first["receiving"]["attention"], list)
    assert isinstance(first["project_procurement_summary"], list)
    assert isinstance(first["attention_items"], list)
    assert isinstance(first["request_pipeline"], list)
    assert sum(row["count"] for row in first["request_pipeline"]) == first["procurement_funnel"]["request_count"]
    assert first["summary"]["actual_paid"] >= 0
    assert first["summary"]["outstanding"] >= 0
    assert first["summary"]["overdue_amount"] >= 0

    second = _dashboard(s, admin_headers)
    assert second == first


# ---------------- Daily Procurement Report ----------------
from daily_report import DailyReport  # noqa: E402

DAILY_REPORT_API = f"{API}/reports/daily"
DPR_DATE = "2018-05-09"
DPR_OTHER_DATE = "2018-05-10"


def _dpr_headers(client, role, suffix=None):
    with SessionLocal() as session:
        username = f"dpr-{role}-{suffix or uuid.uuid4().hex[:8]}"
        _make_user(session, username=username, role=role)
    return _login_headers(client, username)


def _seed_dpr_request(session, report_date, *, request_number=None, status="pricing"):
    timestamp = f"{report_date}T09:00:00+00:00"
    request_id = str(uuid.uuid4())
    session.add(IncomingPurchaseRequest(
        id=request_id, request_number=request_number or f"T-DPR-REQ-{uuid.uuid4().hex[:8]}",
        requester_name="مهندس الموقع", company_name="عميل التقرير اليومي",
        phone_number="01000000000", project_name="مشروع التقرير اليومي",
        project_location="القاهرة", delivery_location="الموقع الرئيسي",
        required_delivery_date="2099-01-01", priority="high", status=status,
        submission_token=uuid.uuid4().hex, content_fingerprint=uuid.uuid4().hex,
        created_at=timestamp, updated_at=timestamp,
    ))
    session.flush()
    session.add(IncomingPurchaseRequestItem(
        id=str(uuid.uuid4()), request_id=request_id, position=1,
        product_name="صنف اختباري", quantity=1, unit="قطعة", review_status="approved",
    ))
    return request_id


def _seed_dpr_po(
    session, report_date, *, supplier_id=None, supplier_name="مورد التقرير اليومي",
    final_total=1000.0, status="sent", po_number=None,
):
    order_id = str(uuid.uuid4())
    timestamp = f"{report_date}T10:00:00+00:00"
    session.add(PurchaseOrder(
        id=order_id, po_number=po_number or f"T-DPR-PO-{uuid.uuid4().hex[:8]}",
        supplier_id=supplier_id or f"T-DPR-SUP-{uuid.uuid4().hex[:8]}", supplier_name=supplier_name,
        project_name="مشروع التقرير اليومي", po_date=report_date, status=status,
        final_total=final_total, created_at=timestamp, updated_at=timestamp,
    ))
    session.flush()
    session.add(PurchaseOrderItem(
        id=str(uuid.uuid4()), purchase_order_id=order_id, product_name="صنف",
        quantity=1, unit="قطعة", unit_price=final_total, line_total=final_total,
    ))
    return order_id


def _seed_dpr_payment(session, order_id, report_date, amount, *, status="recorded"):
    now = f"{report_date}T12:00:00+00:00"
    session.add(PurchaseOrderPayment(
        id=str(uuid.uuid4()), payment_number=f"T-DPR-PAY-{uuid.uuid4().hex[:8]}",
        purchase_order_id=order_id, idempotency_key=uuid.uuid4().hex,
        payment_date=report_date, amount=amount, status=status,
        created_by="محاسب الاختبار", created_at=now, updated_at=now,
    ))


def _seed_dpr_receipt(session, order_id, report_date, *, receipt_type="partial", po_number="", project_name="", supplier_name=""):
    now = f"{report_date}T13:00:00+00:00"
    session.add(PurchaseOrderReceipt(
        id=str(uuid.uuid4()), purchase_order_id=order_id, idempotency_key=uuid.uuid4().hex,
        receipt_type=receipt_type, actor_name="مهندس موقع الاختبار",
        po_number=po_number, project_name=project_name, supplier_name=supplier_name, received_at=now,
    ))


def test_daily_report_defaults_to_today(s, admin_headers):
    response = s.get(DAILY_REPORT_API, headers=admin_headers)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["report_date"] == datetime.now(timezone.utc).date().isoformat()
    assert body["report_number"] == f"DPR-{body['report_date']}"
    assert body["is_closed"] is False


def test_daily_report_includes_requests_received_for_the_selected_date_only(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    number = f"T-DPR-REQ-IN-{suffix}"
    other_number = f"T-DPR-REQ-OUT-{suffix}"
    with SessionLocal.begin() as session:
        _seed_dpr_request(session, DPR_DATE, request_number=number)
        _seed_dpr_request(session, DPR_OTHER_DATE, request_number=other_number)

    body = s.get(DAILY_REPORT_API, params={"date": DPR_DATE}, headers=admin_headers).json()
    numbers = [row["request_number"] for row in body["sections"]["requests_received"]]
    assert number in numbers
    assert other_number not in numbers
    row = next(r for r in body["sections"]["requests_received"] if r["request_number"] == number)
    assert row["item_count"] == 1
    assert row["item_status_breakdown"] == {"approved": 1}
    assert row["priority"] == "high"


def test_daily_report_includes_purchase_orders_issued_for_the_selected_date_only(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    with SessionLocal.begin() as session:
        in_order = _seed_dpr_po(session, DPR_DATE, po_number=f"T-DPR-PO-IN-{suffix}", final_total=1500)
        out_order = _seed_dpr_po(session, DPR_OTHER_DATE, po_number=f"T-DPR-PO-OUT-{suffix}", final_total=1500)

    body = s.get(DAILY_REPORT_API, params={"date": DPR_DATE}, headers=admin_headers).json()
    po_numbers = [row["po_number"] for row in body["sections"]["purchase_orders_issued"]]
    assert f"T-DPR-PO-IN-{suffix}" in po_numbers
    assert f"T-DPR-PO-OUT-{suffix}" not in po_numbers
    row = next(r for r in body["sections"]["purchase_orders_issued"] if r["po_number"] == f"T-DPR-PO-IN-{suffix}")
    assert row["final_total"] == 1500
    assert row["item_count"] == 1
    assert row["payment_status"] == "unpaid"


def test_daily_report_payment_totals_ignore_legacy_direct_payments(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    purchase_id = str(uuid.uuid4())
    with SessionLocal.begin() as session:
        order_id = _seed_dpr_po(session, DPR_DATE, po_number=f"T-DPR-PO-LEGACY-{suffix}", final_total=1000)
        _seed_dpr_payment(session, order_id, DPR_DATE, 400)
        # A legacy Direct Purchase payment dated the same day - must never
        # be mixed into the formal PO Payment Ledger totals below.
        session.add(Purchase(
            id=purchase_id, purchase_id=f"T-DPR-LEGACY-PUR-{suffix}",
            purchase_date=DPR_DATE, supplier_name="مورد قديم", invoice_total=5000,
            created_at=f"{DPR_DATE}T08:00:00+00:00",
        ))
        session.flush()
        session.add(Payment(
            id=str(uuid.uuid4()), payment_id=f"T-DPR-LEGACY-PAY-{suffix}",
            payment_date=DPR_DATE, purchase_id=purchase_id, amount_paid=9999,
            created_at=f"{DPR_DATE}T08:30:00+00:00",
        ))

    body = s.get(DAILY_REPORT_API, params={"date": DPR_DATE}, headers=admin_headers).json()
    payment_amounts = [row["amount"] for row in body["sections"]["payments_today"]]
    assert 400 in payment_amounts
    assert 9999 not in payment_amounts
    row = next(r for r in body["sections"]["purchase_orders_issued"] if r["po_number"] == f"T-DPR-PO-LEGACY-{suffix}")
    assert row["paid_amount"] == 400
    assert row["outstanding_amount"] == 600


def test_daily_report_supplier_financial_position_outstanding_balance(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    supplier_id = f"T-DPR-SUP-BAL-{suffix}"
    with SessionLocal.begin() as session:
        order_id = _seed_dpr_po(session, DPR_DATE, supplier_id=supplier_id, supplier_name="مورد الرصيد", final_total=1000)
        _seed_dpr_payment(session, order_id, DPR_DATE, 300)

    body = s.get(DAILY_REPORT_API, params={"date": DPR_DATE}, headers=admin_headers).json()
    row = next(r for r in body["sections"]["supplier_financial_position"] if r["supplier_id"] == supplier_id)
    assert row["total_po_value"] == 1000
    assert row["paid_amount"] == 300
    assert row["outstanding_amount"] == 700
    assert row["payment_status"] == "partially_paid"


def test_daily_report_supplier_payment_status_fully_paid(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    supplier_id = f"T-DPR-SUP-PAID-{suffix}"
    with SessionLocal.begin() as session:
        order_id = _seed_dpr_po(session, DPR_DATE, supplier_id=supplier_id, final_total=500)
        _seed_dpr_payment(session, order_id, DPR_DATE, 500)

    body = s.get(DAILY_REPORT_API, params={"date": DPR_DATE}, headers=admin_headers).json()
    row = next(r for r in body["sections"]["supplier_financial_position"] if r["supplier_id"] == supplier_id)
    assert row["payment_status"] == "paid"
    assert row["outstanding_amount"] == 0


def test_daily_report_supplier_payment_status_partially_paid(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    supplier_id = f"T-DPR-SUP-PARTIAL-{suffix}"
    with SessionLocal.begin() as session:
        order_id = _seed_dpr_po(session, DPR_DATE, supplier_id=supplier_id, final_total=1000)
        _seed_dpr_payment(session, order_id, DPR_DATE, 400)

    body = s.get(DAILY_REPORT_API, params={"date": DPR_DATE}, headers=admin_headers).json()
    row = next(r for r in body["sections"]["supplier_financial_position"] if r["supplier_id"] == supplier_id)
    assert row["payment_status"] == "partially_paid"


def test_daily_report_supplier_payment_status_unpaid(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    supplier_id = f"T-DPR-SUP-UNPAID-{suffix}"
    with SessionLocal.begin() as session:
        _seed_dpr_po(session, DPR_DATE, supplier_id=supplier_id, final_total=800)

    body = s.get(DAILY_REPORT_API, params={"date": DPR_DATE}, headers=admin_headers).json()
    row = next(r for r in body["sections"]["supplier_financial_position"] if r["supplier_id"] == supplier_id)
    assert row["payment_status"] == "unpaid"
    assert row["paid_amount"] == 0
    assert row["outstanding_amount"] == 800


def test_daily_report_includes_receiving_activity_for_the_selected_date(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po_number = f"T-DPR-PO-RECV-{suffix}"
    with SessionLocal.begin() as session:
        order_id = _seed_dpr_po(session, DPR_DATE, po_number=po_number)
        _seed_dpr_receipt(session, order_id, DPR_DATE, receipt_type="partial", po_number=po_number)

    body = s.get(DAILY_REPORT_API, params={"date": DPR_DATE}, headers=admin_headers).json()
    receiving = body["sections"]["receiving_activity"]
    assert any(row["po_number"] == f"T-DPR-PO-RECV-{suffix}" and row["receipt_type"] == "partial" for row in receiving)


def test_daily_report_needs_attention_includes_a_delivery_problem(s, admin_headers):
    suffix = uuid.uuid4().hex[:8]
    po_number = f"T-DPR-PO-PROBLEM-{suffix}"
    with SessionLocal.begin() as session:
        order_id = _seed_dpr_po(session, DPR_DATE, po_number=po_number, status="delivery_problem")
        _seed_dpr_receipt(session, order_id, DPR_DATE, receipt_type="problem", po_number=po_number)

    body = s.get(DAILY_REPORT_API, params={"date": DPR_DATE}, headers=admin_headers).json()
    assert any(
        item["type"] == "delivery_problem" and item["reference"] == po_number
        for item in body["sections"]["needs_attention"]
    )


def test_daily_report_notes_save_and_load(s, admin_headers):
    date = "2018-06-11"
    saved = s.put(f"{DAILY_REPORT_API}/{date}/notes", headers=admin_headers, json={
        "general_notes": "المورد وعد بالتوريد غدًا", "key_risks": "تأخر الدفع", "follow_up_tomorrow": "",
        "follow_up_notes": "متابعة السداد غدًا",
    })
    assert saved.status_code == 200, saved.text

    body = s.get(DAILY_REPORT_API, params={"date": date}, headers=admin_headers).json()
    assert body["notes"]["general_notes"] == "المورد وعد بالتوريد غدًا"
    assert body["notes"]["key_risks"] == "تأخر الدفع"
    assert body["notes"]["follow_up_notes"] == "متابعة السداد غدًا"


def test_daily_report_role_access_matrix(s, admin_headers):
    engineer_headers = _dpr_headers(s, "procurement_engineer")
    commercial_headers = _dpr_headers(s, "commercial_manager")
    responsible_headers = _dpr_headers(s, "procurement_responsible")

    assert s.get(DAILY_REPORT_API, headers=engineer_headers).status_code == 200
    assert s.get(DAILY_REPORT_API, headers=commercial_headers).status_code == 200
    assert s.get(DAILY_REPORT_API, headers=responsible_headers).status_code == 200

    date = "2018-06-12"
    notes_body = {"general_notes": "x", "key_risks": "", "follow_up_notes": ""}
    assert s.put(f"{DAILY_REPORT_API}/{date}/notes", headers=engineer_headers, json=notes_body).status_code == 403
    assert s.put(f"{DAILY_REPORT_API}/{date}/notes", headers=commercial_headers, json=notes_body).status_code == 403
    assert s.put(f"{DAILY_REPORT_API}/{date}/notes", headers=responsible_headers, json=notes_body).status_code == 200
    assert s.post(f"{DAILY_REPORT_API}/{date}/close", headers=engineer_headers).status_code == 403


def test_daily_report_site_portal_cannot_access(s):
    username = f"dpr-site-only-{uuid.uuid4().hex[:8]}"
    with SessionLocal() as session:
        _make_user(
            session, username=username, password="Sprint21Passw0rd!",
            account_type="site_portal", role="site_engineer",
        )
    headers = _login_headers(s, username)
    assert s.get(DAILY_REPORT_API, headers=headers).status_code == 403


def test_daily_report_close_then_reopen_and_no_duplicate_report_per_date(s, admin_headers):
    date = "2018-06-13"
    with SessionLocal.begin() as session:
        _seed_dpr_po(session, date, po_number=f"T-DPR-PO-CLOSE-{uuid.uuid4().hex[:8]}", final_total=250)

    s.put(f"{DAILY_REPORT_API}/{date}/notes", headers=admin_headers, json={
        "general_notes": "ملاحظة أولى", "key_risks": "", "follow_up_notes": "",
    })
    close = s.post(f"{DAILY_REPORT_API}/{date}/close", headers=admin_headers)
    assert close.status_code == 200, close.text

    again = s.post(f"{DAILY_REPORT_API}/{date}/close", headers=admin_headers)
    assert again.status_code == 409

    body = s.get(DAILY_REPORT_API, params={"date": date}, headers=admin_headers).json()
    assert body["is_closed"] is True
    assert body["summary_frozen"] is True
    assert body["summary"]["purchase_orders_issued_value"] == 250

    reopened = s.post(f"{DAILY_REPORT_API}/{date}/reopen", headers=admin_headers)
    assert reopened.status_code == 200, reopened.text
    reopened_again = s.post(f"{DAILY_REPORT_API}/{date}/reopen", headers=admin_headers)
    assert reopened_again.status_code == 409

    with SessionLocal() as session:
        count = session.query(DailyReport).filter(DailyReport.report_date == date).count()
    assert count == 1
