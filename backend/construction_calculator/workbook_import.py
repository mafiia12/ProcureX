"""Controlled parser for the authoritative construction-rate workbook."""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .domain import normalize_arabic_search


HEADER_MARKERS = {
    "B3": "ملاحظات",
    "H3": "نسبة زيادة الكمية المنفذة",
    "J3": "معدل هوالك الخامات",
    "L3": "معدل استهلاك الخامات / وحدة من البند",
    "P3": "انتاجية الطاقم",
    "R3": "العمالة والمعدات",
    "U3": "الوحدة",
    "V3": "البند",
}
EQUIPMENT_TERMS = {
    "حفار",
    "قلاب",
    "شاكوش",
    "لودر",
    "هراس",
    "دكاك",
    "هلتى",
    "خلاطة",
    "مضخة",
    "طلمبة",
}
MATERIAL_UNITS = (
    "متر طولى",
    "متر طولي",
    "م.ط",
    "كجم",
    "لتر",
    "م3",
    "م2",
    "طوبة",
    "قطعة",
    "روول",
    "رول",
)
DASH_ONLY = re.compile(r"^[\sـ\-–—]+$")
NUMBER = r"\d+(?:[.,]\d+)?"
RANGE_RATE = re.compile(
    rf"^\(?\s*(?P<minimum>{NUMBER})\s*[-–—]\s*(?P<maximum>{NUMBER})\s*"
    rf"(?P<unit>{'|'.join(map(re.escape, MATERIAL_UNITS))})\s*\)?\s*(?P<name>.+)$"
)
FIXED_RATE = re.compile(
    rf"^(?P<value>{NUMBER})\s*(?P<unit>{'|'.join(map(re.escape, MATERIAL_UNITS))})\s*(?P<name>.+)$"
)
COUNT_ONLY_RATE = re.compile(rf"^(?P<value>{NUMBER})\s*(?P<unit>طوبة|قطعة)\s*$")
UNITLESS_RATE = re.compile(rf"^(?P<value>{NUMBER})\s*(?P<name>[^%]+)$")
COVERAGE_NOTE = re.compile(rf"^\(?\s*تعمل\s+بؤج.*حوالى\s+{NUMBER}\s*م2.*\)?$")
WASTE_RATE = re.compile(rf"^(?P<name>.*?)\s*(?P<value>{NUMBER})\s*%\s*$")
ADJUSTMENT_RANGE = re.compile(
    rf"(?P<minimum>{NUMBER})\s*%?\s*[-–—]\s*(?P<maximum>{NUMBER})\s*%"
)


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()


def _decimal(value: str | int | float | Decimal) -> Decimal:
    return Decimal(str(value).replace(",", ".").strip())


def _source_identity(entity: str, sheet: str, row: int, suffix: str = "") -> str:
    tail = f"|{suffix}" if suffix else ""
    return f"{sheet}|{entity}|{row:04d}{tail}"


@dataclass
class ImportIssue:
    severity: str
    entity_type: str
    sheet: str
    row: int
    source_text: str
    reason: str


@dataclass
class ParsedParticipant:
    kind: str
    name: str
    quantity: int | None
    source_row: int
    source_text: str
    source_identity: str


@dataclass
class ParsedMaterialRate:
    name: str
    material_unit: str
    rate_kind: str
    fixed_rate: Decimal | None
    minimum_rate: Decimal | None
    maximum_rate: Decimal | None
    basis_quantity: Decimal
    basis_unit: str
    waste_percentage: Decimal | None
    conditions: str
    source_row: int
    source_text: str
    source_identity: str


@dataclass
class ParsedCrew:
    source_row: int
    source_identity: str
    description_parts: list[str] = field(default_factory=list)
    source_text_parts: list[str] = field(default_factory=list)
    productivity: Decimal | None = None
    productivity_unit: str = ""
    productivity_conditions: str = ""
    participants: list[ParsedParticipant] = field(default_factory=list)
    rates: list[ParsedMaterialRate] = field(default_factory=list)
    waste_specs: list[tuple[str, Decimal, int, str]] = field(default_factory=list)


@dataclass
class ParsedAdjustment:
    kind: str
    fixed_percentage: Decimal | None
    minimum_percentage: Decimal | None
    maximum_percentage: Decimal | None
    documented_meaning: str
    source_row: int
    source_text: str
    source_identity: str


@dataclass
class ParsedWorkItem:
    category_identity: str
    name_parts: list[str]
    measurement_unit: str
    source_row: int
    source_identity: str
    source_text_parts: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    crews: list[ParsedCrew] = field(default_factory=list)
    adjustments: list[ParsedAdjustment] = field(default_factory=list)


@dataclass
class ParsedCategory:
    name: str
    source_row: int
    source_text: str
    source_identity: str
    notes: list[str] = field(default_factory=list)


@dataclass
class ImportBundle:
    workbook_path: Path
    workbook_name: str
    workbook_sha256: str
    import_version: str
    worksheet: str
    categories: list[ParsedCategory]
    work_items: list[ParsedWorkItem]
    issues: list[ImportIssue]


def _header_matches(actual: str, expected: str) -> bool:
    normalized = normalize_arabic_search(actual)
    return normalize_arabic_search(expected) in normalized


def _split_plus(value: str) -> list[str]:
    return [
        part.strip(" ()") for part in re.split(r"\s*\+\s*", value) if part.strip(" ()")
    ]


def _material_display_name(raw: str) -> tuple[str, str]:
    name = _clean(raw).strip("+")
    conditions = ""
    if " / " in name:
        name, conditions = (_clean(part) for part in name.split(" / ", 1))
    coverage = re.search(r"\([^)]*\)\s*يفرد\s+مسطح\s+([\d.,]+)\s*(م2)", name)
    if coverage:
        name = _clean(name[: coverage.start()])
        conditions = _clean(raw[coverage.start() :])
    return name, conditions


def _material_reference_key(name: str) -> str:
    tokens = normalize_arabic_search(name).split()
    return " ".join(
        token[2:] if token.startswith("ال") and len(token) > 3 else token
        for token in tokens
    )


def _material_unit_references(
    worksheet, last_row: int
) -> dict[str, list[tuple[str, str, int]]]:
    references: dict[str, list[tuple[str, str, int]]] = {}
    for row in range(5, last_row + 1):
        cleaned = _clean(worksheet[f"L{row}"].value)
        if not cleaned or DASH_ONLY.fullmatch(cleaned):
            continue
        for segment in _split_plus(cleaned):
            match = RANGE_RATE.match(segment) or FIXED_RATE.match(segment)
            if not match:
                continue
            name, _ = _material_display_name(match.group("name"))
            key = _material_reference_key(name)
            reference = (name, match.group("unit"), row)
            references.setdefault(key, [])
            if reference not in references[key]:
                references[key].append(reference)
    return references


def _parse_material_segments(
    text: str,
    sheet: str,
    row: int,
    basis_unit: str,
    issues: list[ImportIssue],
    unit_references: dict[str, list[tuple[str, str, int]]],
) -> list[ParsedMaterialRate]:
    cleaned = _clean(text)
    if not cleaned or DASH_ONLY.fullmatch(cleaned):
        return []
    parsed: list[ParsedMaterialRate] = []
    for index, segment in enumerate(_split_plus(cleaned), 1):
        range_match = RANGE_RATE.match(segment)
        fixed_match = FIXED_RATE.match(segment)
        count_only_match = COUNT_ONLY_RATE.match(segment)
        if range_match:
            name, conditions = _material_display_name(range_match.group("name"))
            parsed.append(
                ParsedMaterialRate(
                    name=name,
                    material_unit=range_match.group("unit"),
                    rate_kind="range",
                    fixed_rate=None,
                    minimum_rate=_decimal(range_match.group("minimum")),
                    maximum_rate=_decimal(range_match.group("maximum")),
                    basis_quantity=Decimal("1"),
                    basis_unit=basis_unit,
                    waste_percentage=None,
                    conditions=conditions,
                    source_row=row,
                    source_text=segment,
                    source_identity=_source_identity("rate", sheet, row, str(index)),
                )
            )
            continue
        if fixed_match:
            name, conditions = _material_display_name(fixed_match.group("name"))
            basis_quantity = Decimal("1")
            coverage = re.search(r"يفرد\s+مسطح\s+([\d.,]+)\s*(م2)", segment)
            if coverage:
                basis_quantity = _decimal(coverage.group(1))
            parsed.append(
                ParsedMaterialRate(
                    name=name,
                    material_unit=fixed_match.group("unit"),
                    rate_kind="fixed",
                    fixed_rate=_decimal(fixed_match.group("value")),
                    minimum_rate=None,
                    maximum_rate=None,
                    basis_quantity=basis_quantity,
                    basis_unit=basis_unit,
                    waste_percentage=None,
                    conditions=conditions,
                    source_row=row,
                    source_text=segment,
                    source_identity=_source_identity("rate", sheet, row, str(index)),
                )
            )
            continue
        if count_only_match:
            unit = count_only_match.group("unit")
            parsed.append(
                ParsedMaterialRate(
                    name=unit,
                    material_unit=unit,
                    rate_kind="fixed",
                    fixed_rate=_decimal(count_only_match.group("value")),
                    minimum_rate=None,
                    maximum_rate=None,
                    basis_quantity=Decimal("1"),
                    basis_unit=basis_unit,
                    waste_percentage=None,
                    conditions="",
                    source_row=row,
                    source_text=segment,
                    source_identity=_source_identity("rate", sheet, row, str(index)),
                )
            )
            continue
        unitless_match = UNITLESS_RATE.match(segment)
        if unitless_match:
            name, conditions = _material_display_name(unitless_match.group("name"))
            references = unit_references.get(_material_reference_key(name), [])
            units = {unit for _, unit, _ in references}
            if len(units) == 1:
                unit = next(iter(units))
                reference_name = references[0][0]
                reference_rows = ", ".join(
                    str(reference_row) for _, _, reference_row in references
                )
                audit_note = (
                    f"تم استكمال الوحدة ({unit}) من اسم خامة مطابق بعد توحيد "
                    f"أداة التعريف داخل نفس الشيت، صف {reference_rows}؛ السطر "
                    "الأصلي لا يذكر الوحدة."
                )
                conditions = " | ".join(
                    part for part in (conditions, audit_note) if part
                )
                parsed.append(
                    ParsedMaterialRate(
                        name=reference_name,
                        material_unit=unit,
                        rate_kind="fixed",
                        fixed_rate=_decimal(unitless_match.group("value")),
                        minimum_rate=None,
                        maximum_rate=None,
                        basis_quantity=Decimal("1"),
                        basis_unit=basis_unit,
                        waste_percentage=None,
                        conditions=conditions,
                        source_row=row,
                        source_text=segment,
                        source_identity=_source_identity(
                            "rate", sheet, row, str(index)
                        ),
                    )
                )
                continue
        if COVERAGE_NOTE.match(segment):
            # This is a source annotation about the setup coverage, not a
            # material rate. The caller still retains it in crew source text.
            continue
        if not re.match(rf"^\(?\s*{NUMBER}", segment):
            issues.append(
                ImportIssue(
                    severity="warning",
                    entity_type="material_source_note",
                    sheet=sheet,
                    row=row,
                    source_text=segment,
                    reason="Source text is descriptive and was retained without creating a numerical material rate",
                )
            )
            continue
        issues.append(
            ImportIssue(
                severity="rejected",
                entity_type="material_rate",
                sheet=sheet,
                row=row,
                source_text=segment,
                reason="A numeric rate, explicit material unit, and material name could not all be identified",
            )
        )
    return parsed


def _parse_waste(text: str, sheet: str, row: int, issues: list[ImportIssue]):
    cleaned = _clean(text)
    if not cleaned or DASH_ONLY.fullmatch(cleaned):
        return []
    values = []
    for segment in _split_plus(cleaned):
        match = WASTE_RATE.match(segment)
        if not match:
            issues.append(
                ImportIssue(
                    severity="rejected",
                    entity_type="waste_rate",
                    sheet=sheet,
                    row=row,
                    source_text=segment,
                    reason="Waste percentage could not be parsed without inference",
                )
            )
            continue
        values.append(
            (
                _clean(match.group("name")),
                _decimal(match.group("value")) / 100,
                row,
                segment,
            )
        )
    return values


def _name_tokens(value: str) -> set[str]:
    return {
        token for token in normalize_arabic_search(value).split() if len(token) >= 2
    }


def _apply_waste(crew: ParsedCrew, sheet: str, issues: list[ImportIssue]) -> None:
    for waste_name, percentage, row, source in crew.waste_specs:
        normalized = normalize_arabic_search(waste_name)
        if "كل خامه" in normalized:
            for rate in crew.rates:
                if rate.waste_percentage is None:
                    rate.waste_percentage = percentage
            continue
        waste_tokens = _name_tokens(waste_name)
        matches = [
            rate
            for rate in crew.rates
            if waste_tokens & _name_tokens(rate.name)
            or any(
                left.startswith(right[:3]) or right.startswith(left[:3])
                for left in waste_tokens
                for right in _name_tokens(rate.name)
            )
        ]
        if not matches:
            issues.append(
                ImportIssue(
                    severity="warning",
                    entity_type="waste_rate",
                    sheet=sheet,
                    row=row,
                    source_text=source,
                    reason="Waste percentage was preserved but no unambiguous material-rate match was found",
                )
            )
            continue
        for rate in matches:
            rate.waste_percentage = percentage


def _participant_kind(name: str) -> str:
    normalized = normalize_arabic_search(name)
    if any(normalize_arabic_search(term) in normalized for term in EQUIPMENT_TERMS):
        return "equipment"
    return "labor"


def _labor_kind(name: str) -> str:
    normalized = normalize_arabic_search(name)
    if any(term in normalized for term in ("مساعد", "موان", "صبي")):
        return "helper"
    if "عامل" in normalized:
        return "general_worker"
    return "craftsman"


def _parse_participants(text: str, sheet: str, row: int, crew_suffix: str):
    participants = []
    for index, segment in enumerate(_split_plus(_clean(text)), 1):
        match = re.match(r"^(?P<count>\d+)\s*(?P<name>.+)$", segment)
        count = int(match.group("count")) if match else None
        name = _clean(match.group("name") if match else segment)
        kind = _participant_kind(name)
        participants.append(
            ParsedParticipant(
                kind=kind,
                name=name,
                quantity=count,
                source_row=row,
                source_text=segment,
                source_identity=_source_identity(
                    "crew-equipment" if kind == "equipment" else "crew-labor",
                    sheet,
                    row,
                    f"{crew_suffix}-{index}",
                ),
            )
        )
    return participants


def _parse_productivity(
    value: object, basis_unit: str
) -> tuple[Decimal | None, str, str]:
    if value is None or _clean(value) == "":
        return None, "", ""
    if isinstance(value, (int, float, Decimal)):
        return _decimal(value), f"{basis_unit}/day", ""
    text = _clean(value)
    match = re.match(rf"^(?P<value>{NUMBER})(?P<conditions>.*)$", text)
    if not match:
        return None, "", text
    return (
        _decimal(match.group("value")),
        f"{basis_unit}/day",
        _clean(match.group("conditions")),
    )


def _parse_adjustment(
    value: object, sheet: str, row: int, meaning: str
) -> ParsedAdjustment | None:
    if value is None or _clean(value) == "":
        return None
    source = _clean(value)
    if isinstance(value, (int, float, Decimal)):
        fixed = _decimal(value)
        return ParsedAdjustment(
            kind="fixed",
            fixed_percentage=fixed,
            minimum_percentage=None,
            maximum_percentage=None,
            documented_meaning=meaning,
            source_row=row,
            source_text=source,
            source_identity=_source_identity("adjustment", sheet, row),
        )
    match = ADJUSTMENT_RANGE.search(source)
    if match:
        return ParsedAdjustment(
            kind="range",
            fixed_percentage=None,
            minimum_percentage=_decimal(match.group("minimum")) / 100,
            maximum_percentage=_decimal(match.group("maximum")) / 100,
            documented_meaning=meaning,
            source_row=row,
            source_text=source,
            source_identity=_source_identity("adjustment", sheet, row),
        )
    return None


def parse_workbook(path: Path, import_version: str | None = None) -> ImportBundle:
    path = Path(path).resolve()
    workbook_bytes = path.read_bytes()
    workbook_sha = hashlib.sha256(workbook_bytes).hexdigest()
    version = import_version or f"construction-v1-{workbook_sha[:12]}"
    # Read through an in-memory stream so a validation failure cannot leave the
    # uploaded workbook locked on Windows while its temporary directory cleans up.
    workbook = load_workbook(
        io.BytesIO(workbook_bytes), data_only=True, read_only=False
    )
    if len(workbook.worksheets) != 1:
        raise ValueError(
            "The authoritative workbook must contain exactly one relevant worksheet"
        )
    worksheet = workbook.worksheets[0]
    for address, expected in HEADER_MARKERS.items():
        if not _header_matches(_clean(worksheet[address].value), expected):
            raise ValueError(
                f"Unexpected workbook layout at {worksheet.title}!{address}"
            )

    last_row = max(
        row
        for row in range(1, worksheet.max_row + 1)
        if any(worksheet[f"{column}{row}"].value is not None for column in "BHJLPRUV")
    )
    material_unit_references = _material_unit_references(worksheet, last_row)
    categories: list[ParsedCategory] = []
    work_items: list[ParsedWorkItem] = []
    issues: list[ImportIssue] = []
    current_category: ParsedCategory | None = None
    current_item: ParsedWorkItem | None = None
    current_crew: ParsedCrew | None = None

    for row in range(5, last_row + 1):
        values = {
            column: _clean(worksheet[f"{column}{row}"].value) for column in "BHJLPRUV"
        }
        if not any(values.values()):
            continue
        category_cell = worksheet[f"B{row}"]
        if (
            values["B"]
            and category_cell.font.bold
            and (category_cell.font.sz or 0) >= 20
        ):
            current_category = ParsedCategory(
                name=values["B"],
                source_row=row,
                source_text=values["B"],
                source_identity=_source_identity("category", worksheet.title, row),
            )
            categories.append(current_category)
            current_item = None
            current_crew = None
            continue
        if current_category is None:
            issues.append(
                ImportIssue(
                    "rejected",
                    "row",
                    worksheet.title,
                    row,
                    " | ".join(values.values()),
                    "No category was active",
                )
            )
            continue
        if values["U"] and values["V"]:
            current_item = ParsedWorkItem(
                category_identity=current_category.source_identity,
                name_parts=[values["V"]],
                measurement_unit=values["U"],
                source_row=row,
                source_identity=_source_identity("work-item", worksheet.title, row),
                source_text_parts=[values["V"]],
            )
            work_items.append(current_item)
            current_crew = None
        elif values["V"] and current_item is not None:
            current_item.name_parts.append(values["V"])
            current_item.source_text_parts.append(values["V"])
        if current_item is None:
            if values["B"]:
                current_category.notes.append(values["B"])
            continue
        if values["B"]:
            current_item.notes.append(values["B"])
            current_item.source_text_parts.append(values["B"])

        productivity, productivity_unit, productivity_conditions = _parse_productivity(
            worksheet[f"P{row}"].value, current_item.measurement_unit
        )
        if values["R"] or productivity is not None:
            if values["R"] and productivity is None and current_crew is not None:
                current_crew.description_parts.append(values["R"])
                current_crew.source_text_parts.append(values["R"])
                current_crew.participants.extend(
                    _parse_participants(
                        values["R"], worksheet.title, row, str(current_crew.source_row)
                    )
                )
            else:
                current_crew = ParsedCrew(
                    source_row=row,
                    source_identity=_source_identity("crew", worksheet.title, row),
                    description_parts=[values["R"]] if values["R"] else [],
                    source_text_parts=[values["R"]] if values["R"] else [],
                    productivity=productivity,
                    productivity_unit=productivity_unit,
                    productivity_conditions=productivity_conditions,
                )
                if values["R"]:
                    current_crew.participants.extend(
                        _parse_participants(values["R"], worksheet.title, row, str(row))
                    )
                current_item.crews.append(current_crew)
        if values["L"]:
            if current_crew is None:
                current_crew = ParsedCrew(
                    source_row=row,
                    source_identity=_source_identity("crew", worksheet.title, row),
                )
                current_item.crews.append(current_crew)
            current_crew.rates.extend(
                _parse_material_segments(
                    values["L"],
                    worksheet.title,
                    row,
                    current_item.measurement_unit,
                    issues,
                    material_unit_references,
                )
            )
            current_crew.source_text_parts.append(values["L"])
        if values["J"]:
            if current_crew is None:
                current_crew = ParsedCrew(
                    source_row=row,
                    source_identity=_source_identity("crew", worksheet.title, row),
                )
                current_item.crews.append(current_crew)
            current_crew.waste_specs.extend(
                _parse_waste(values["J"], worksheet.title, row, issues)
            )
            current_crew.source_text_parts.append(values["J"])
        if values["H"]:
            adjustment = _parse_adjustment(
                worksheet[f"H{row}"].value,
                worksheet.title,
                row,
                values["B"],
            )
            if adjustment:
                current_item.adjustments.append(adjustment)
            else:
                issues.append(
                    ImportIssue(
                        "rejected",
                        "adjustment",
                        worksheet.title,
                        row,
                        values["H"],
                        "Adjustment percentage could not be parsed",
                    )
                )

    unique_work_items: list[ParsedWorkItem] = []
    seen_work_item_names: dict[tuple[str, str], int] = {}
    for item in work_items:
        duplicate_key = (
            item.category_identity,
            _clean(" ".join(item.name_parts)).casefold(),
        )
        original_row = seen_work_item_names.get(duplicate_key)
        if original_row is not None:
            issues.append(
                ImportIssue(
                    "rejected",
                    "duplicate_work_item_name",
                    worksheet.title,
                    item.source_row,
                    " | ".join(item.source_text_parts),
                    f"Duplicate work-item name in the same category; first source row is {original_row}",
                )
            )
            continue
        seen_work_item_names[duplicate_key] = item.source_row
        unique_work_items.append(item)

    for item in unique_work_items:
        for crew in item.crews:
            _apply_waste(crew, worksheet.title, issues)
        if not item.crews:
            issues.append(
                ImportIssue(
                    "warning",
                    "work_item",
                    worksheet.title,
                    item.source_row,
                    " | ".join(item.source_text_parts),
                    "Work item has no crew, material, or productivity scenario",
                )
            )
    bundle = ImportBundle(
        workbook_path=path,
        workbook_name=path.name,
        workbook_sha256=workbook_sha,
        import_version=version,
        worksheet=worksheet.title,
        categories=categories,
        work_items=unique_work_items,
        issues=issues,
    )
    workbook.close()
    return bundle


def labor_role_kind(name: str) -> str:
    return _labor_kind(name)
