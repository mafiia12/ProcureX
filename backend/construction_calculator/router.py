"""FastAPI adapter for the construction calculator application service."""

from __future__ import annotations

import tempfile
from decimal import Decimal
from pathlib import Path
from typing import Annotated

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field

try:
    from ..database import SessionLocal
    from ..incoming_requests import require_internal_access
except ImportError:  # pragma: no cover
    from database import SessionLocal
    from incoming_requests import require_internal_access

from .permissions import require_permission
from .repository import ConstructionRepository
from .service import ConstructionCalculatorService


router = APIRouter(
    prefix="/construction-calculator",
    tags=["construction-calculator"],
    dependencies=[Depends(require_internal_access)],
)


def get_service():
    with SessionLocal() as session:
        yield ConstructionCalculatorService(ConstructionRepository(session))


Service = Annotated[ConstructionCalculatorService, Depends(get_service)]


class CalculationIn(BaseModel):
    work_item_id: str
    project_id: str | None = None
    measured_quantity: Decimal = Field(gt=0)
    measurement_unit: str = Field(min_length=1, max_length=80)
    selected_crew_ids: list[str] = Field(min_length=1, max_length=20)
    primary_crew_id: str
    crew_count: int | None = Field(default=None, gt=0, le=10000)
    required_days: Decimal | None = Field(default=None, gt=0)
    adjustment_id: str | None = None
    adjustment_percentage: Decimal | None = Field(default=None, ge=0, le=1)
    selected_rates: dict[str, Decimal] = Field(default_factory=dict)
    package_capacities: dict[str, Decimal | None] = Field(default_factory=dict)
    labor_count_overrides: dict[str, int] = Field(default_factory=dict)
    equipment_quantity_overrides: dict[str, int] = Field(default_factory=dict)


class SaveCalculationIn(CalculationIn):
    created_by: str = Field(default="", max_length=200)


class PurchaseRequestConversionIn(BaseModel):
    required_delivery_date: str = Field(pattern=r"^\d{4}-\d{2}-\d{2}$")
    project_location: str = Field(default="", max_length=500)
    delivery_location: str = Field(min_length=2, max_length=500)
    priority: str = Field(default="normal", pattern="^(low|normal|high|urgent)$")
    created_by: str = Field(default="", max_length=200)
    create_revision: bool = False


class PackagingIn(BaseModel):
    package_capacity: Decimal | None = Field(default=None, gt=0)
    package_unit: str = Field(default="", max_length=80)
    package_purchase_unit: str = Field(default="", max_length=80)
    package_notes: str = Field(default="", max_length=1000)


class RecalculateIn(BaseModel):
    created_by: str = Field(default="", max_length=200)


class ReferenceRateStatusIn(BaseModel):
    is_active: bool


@router.get(
    "/categories", dependencies=[Depends(require_permission("view_construction_rates"))]
)
def categories(service: Service):
    return service.categories()


@router.get(
    "/options", dependencies=[Depends(require_permission("view_construction_rates"))]
)
def options(service: Service):
    return service.options()


@router.get(
    "/work-items", dependencies=[Depends(require_permission("view_construction_rates"))]
)
def search_work_items(
    service: Service,
    q: str = Query(default="", max_length=200),
    category_id: str | None = None,
    material_id: str | None = None,
    labor_role_id: str | None = None,
    equipment_id: str | None = None,
    measurement_unit: str | None = None,
    import_version: str | None = None,
    rate_kind: str | None = Query(default=None, pattern="^(fixed|range)$"),
    crew_status: str | None = Query(default=None, pattern="^(specified|unspecified)$"),
    active: bool = True,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
):
    return service.search(
        query=q,
        category_id=category_id,
        material_id=material_id,
        labor_role_id=labor_role_id,
        equipment_id=equipment_id,
        measurement_unit=measurement_unit,
        import_version=import_version,
        rate_kind=rate_kind,
        crew_status=crew_status,
        active=active,
        page=page,
        page_size=page_size,
    )


@router.get(
    "/work-items/{work_item_id}",
    dependencies=[
        Depends(require_permission("view_construction_rates")),
        Depends(require_permission("view_source_audit")),
    ],
)
def work_item(work_item_id: str, service: Service):
    try:
        return service.work_item(work_item_id)
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc


@router.post("/calculate", dependencies=[Depends(require_permission("use_calculator"))])
def calculate_preview(body: CalculationIn, service: Service):
    try:
        return service.preview(body.model_dump())
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.post(
    "/calculations", dependencies=[Depends(require_permission("save_calculations"))]
)
def save_calculation(body: SaveCalculationIn, service: Service):
    try:
        payload = body.model_dump(exclude={"created_by"})
        return service.save(payload, body.created_by.strip())
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.get(
    "/calculations", dependencies=[Depends(require_permission("save_calculations"))]
)
def saved_calculations(service: Service, limit: int = Query(default=100, ge=1, le=500)):
    return service.list_saved(limit)


@router.get(
    "/calculations/{session_id}",
    dependencies=[Depends(require_permission("save_calculations"))],
)
def saved_calculation(session_id: str, service: Service):
    try:
        return service.saved_detail(session_id)
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc


@router.post(
    "/calculations/{session_id}/recalculate-latest",
    dependencies=[
        Depends(require_permission("use_calculator")),
        Depends(require_permission("save_calculations")),
    ],
)
def recalculate_latest(session_id: str, body: RecalculateIn, service: Service):
    try:
        return service.recalculate_latest(session_id, body.created_by.strip())
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.post(
    "/calculations/{session_id}/purchase-request",
    dependencies=[Depends(require_permission("create_purchase_requests"))],
)
def convert_to_purchase_request(
    session_id: str, body: PurchaseRequestConversionIn, service: Service
):
    try:
        payload = body.model_dump(exclude={"created_by"})
        return service.convert_to_purchase_request(
            session_id, payload, body.created_by.strip()
        )
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except FileExistsError as exc:
        raise HTTPException(409, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.put(
    "/materials/{material_id}/packaging",
    dependencies=[Depends(require_permission("manage_packaging", administrator=True))],
)
def update_packaging(material_id: str, body: PackagingIn, service: Service):
    try:
        return service.update_packaging(material_id, body.model_dump())
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.put(
    "/admin/material-rates/{rate_id}/status",
    dependencies=[
        Depends(require_permission("edit_approved_reference_rates", administrator=True))
    ],
)
def set_material_rate_status(
    rate_id: str, body: ReferenceRateStatusIn, service: Service
):
    try:
        return service.set_material_rate_status(rate_id, body.is_active)
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc


@router.post(
    "/admin/import-workbook",
    dependencies=[
        Depends(require_permission("import_workbook_updates", administrator=True))
    ],
)
async def import_workbook(
    service: Service,
    file: UploadFile = File(...),
    import_version: str = Form(..., min_length=3, max_length=100),
):
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "Workbook exceeds the 10 MB limit")
    if not content.startswith(b"PK") or not (file.filename or "").lower().endswith(
        ".xlsx"
    ):
        raise HTTPException(422, "A valid XLSX workbook is required")
    with tempfile.TemporaryDirectory(prefix="procurex-construction-import-") as folder:
        path = Path(folder) / "construction-rates.xlsx"
        path.write_bytes(content)
        try:
            return service.import_workbook(path, import_version.strip())
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc


@router.post(
    "/admin/import-market-aliases",
    dependencies=[
        Depends(require_permission("manage_market_aliases", administrator=True))
    ],
)
async def import_market_aliases(service: Service, file: UploadFile = File(...)):
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Alias workbook exceeds the 5 MB limit")
    if not content.startswith(b"PK") or not (file.filename or "").lower().endswith(
        ".xlsx"
    ):
        raise HTTPException(422, "A valid XLSX workbook is required")
    with tempfile.TemporaryDirectory(prefix="procurex-construction-alias-") as folder:
        path = Path(folder) / "aliases.xlsx"
        path.write_bytes(content)
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            worksheet = workbook.active
            headers = [
                str(cell.value or "").strip().casefold()
                for cell in next(worksheet.iter_rows())
            ]
            required = ["entity_type", "code", "alias"]
            if any(name not in headers for name in required):
                raise HTTPException(
                    422, "Alias workbook headers must be entity_type, code, alias"
                )
            rows = []
            for values in worksheet.iter_rows(min_row=2, values_only=True):
                row = dict(zip(headers, values))
                if any(row.get(name) not in (None, "") for name in required):
                    rows.append(row)
        finally:
            workbook.close()
        return service.import_market_aliases(rows, file.filename or "aliases.xlsx")
