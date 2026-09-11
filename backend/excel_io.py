"""Legacy workbook.xlsm import/export - test-fixture infrastructure only.

Not reachable from the live application: server.py has no route that calls
parse_workbook, import_data, or build_export_workbook (the former
POST /api/import/excel and GET /api/export/excel were removed). The only
remaining callers are tests/backend_test.py's module-level seed-data
bootstrap and the standalone backend/migrate.py script (itself only
referenced, commented-out, from start_backend.bat). Suppliers, customers,
projects, and items are added through the app's own UI now - this module
predates that and is kept solely so the existing test suite's seed data
keeps working; it is not an active data-entry path.
"""

import uuid
import warnings
from io import BytesIO
from datetime import datetime, date, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill

try:
    from .business_codes import next_business_code
except ImportError:
    from business_codes import next_business_code

SETTINGS_LISTS = [
    ("purchase_status", "حالة الطلب", 0),
    ("purchase_type", "نوع الشراء", 1),
    ("currency", "العملة", 2),
    ("payment_methods", "طرق الدفع", 4),
    ("vat_rates", "نسب الضريبة", 5),
    ("units", "الوحدات", 6),
    ("supplier_categories", "تصنيفات الموردين", 7),
    ("project_status", "حالة المشروع", 8),
    ("delivery_status", "حالة التسليم", 9),
    ("priority", "الأولوية", 10),
    ("governorates", "المحافظات", 11),
    ("record_status", "حالة السجل", 12),
]


def _s(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _n(v):
    try:
        if v is None or v == "":
            return 0.0
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _pct(v):
    n = _n(v)
    return n * 100 if 0 < n <= 1 else n


def parse_workbook(content: bytes) -> dict:
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Print area cannot be set to Defined name.*")
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    out = {"suppliers": [], "customers": [], "projects": [], "items": [],
           "purchases": [], "price_history": [], "payments": [], "settings": {}}

    if "Suppliers" in wb.sheetnames:
        ws = wb["Suppliers"]
        group, started = "", False
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=15, values_only=True):
            row = list(row) + [None] * (15 - len(row))
            first, name = row[0], row[1]
            if _s(first) == "كود المورد":
                started = True
                continue
            if not started or _s(first).startswith("PUR-"):
                continue
            if _s(name) == "":
                texts = [_s(c) for c in row if _s(c) and not _s(c).replace(".", "").isdigit()]
                if texts:
                    group = texts[0]
                continue
            out["suppliers"].append({
                "name": _s(name), "specialty": _s(row[2]), "governorate": _s(row[3]),
                "city": _s(row[4]), "address": _s(row[5]), "contact_person": _s(row[6]),
                "phone": _s(row[7]), "whatsapp": _s(row[8]), "email": _s(row[9]),
                "payment_terms": _s(row[10]), "lead_time_days": _s(row[11]),
                "rating": _s(row[12]), "status": _s(row[13]) or "نشط",
                "notes": _s(row[14]), "group_name": group,
            })

    if "Customers" in wb.sheetnames:
        started = False
        for row in wb["Customers"].iter_rows(min_row=1, max_col=11, values_only=True):
            row = list(row) + [None] * (11 - len(row))
            if _s(row[0]) == "كود العميل":
                started = True
                continue
            if not started or _s(row[1]) == "":
                continue
            out["customers"].append({
                "code": _s(row[0]), "name": _s(row[1]), "contact_person": _s(row[2]),
                "phone": _s(row[3]), "whatsapp": _s(row[4]), "email": _s(row[5]),
                "governorate": _s(row[6]), "city": _s(row[7]), "address": _s(row[8]),
                "status": _s(row[9]), "notes": _s(row[10]),
            })

    if "Projects" in wb.sheetnames:
        started = False
        for row in wb["Projects"].iter_rows(min_row=1, max_col=12, values_only=True):
            row = list(row) + [None] * (12 - len(row))
            if _s(row[0]) == "كود المشروع":
                started = True
                continue
            if not started or _s(row[1]) == "":
                continue
            out["projects"].append({
                "code": _s(row[0]), "name": _s(row[1]), "customer_name": _s(row[2]),
                "governorate": _s(row[3]), "city": _s(row[4]), "address": _s(row[5]),
                "engineer": _s(row[6]), "start_date": _s(row[7]), "end_date": _s(row[8]),
                "budget": _s(row[9]), "status": _s(row[10]), "notes": _s(row[11]),
            })

    if "Items" in wb.sheetnames:
        started = False
        has_category_hierarchy = False
        has_product_identity = False
        for row in wb["Items"].iter_rows(min_row=1, max_col=12, values_only=True):
            row = list(row) + [None] * (12 - len(row))
            if _s(row[0]) == "كود الصنف":
                started = True
                has_product_identity = _s(row[2]) == "اسم المنتج"
                has_category_hierarchy = _s(row[3]) == "التصنيف الفرعي"
                continue
            if not started or _s(row[1]) == "":
                continue
            if has_product_identity:
                legacy_name, product_name, brand = _s(row[1]), _s(row[2]), _s(row[3])
                main_category, subcategory = _s(row[4]), _s(row[5])
                unit, specifications, pref, notes = (
                    _s(row[6]), _s(row[7]), _s(row[8]), _s(row[9])
                )
            elif has_category_hierarchy:
                legacy_name, product_name, brand = _s(row[1]), _s(row[1]), ""
                main_category, subcategory = _s(row[2]), _s(row[3])
                unit, specifications, pref = _s(row[4]), _s(row[5]), _s(row[6])
                notes = _s(row[7])
            else:
                legacy_name, product_name, brand = _s(row[1]), _s(row[1]), ""
                main_category, subcategory = _s(row[2]), ""
                unit, specifications, pref = _s(row[3]), _s(row[4]), _s(row[5])
                notes = _s(row[11]) or _s(row[6])
            out["items"].append({
                "code": _s(row[0]), "name": legacy_name,
                "product_name": product_name or legacy_name, "brand": brand,
                "category": main_category, "main_category": main_category,
                "subcategory": subcategory, "unit": unit,
                "specs": specifications, "specifications": specifications,
                "preferred_supplier": "" if pref == "0" else pref,
                "notes": notes,
            })

    if "Purchase Register" in wb.sheetnames:
        for row in wb["Purchase Register"].iter_rows(min_row=2, max_col=10, values_only=True):
            row = list(row) + [None] * (10 - len(row))
            if not _s(row[0]).startswith("PUR-"):
                continue
            out["purchases"].append({
                "purchase_id": _s(row[0]), "purchase_date": _s(row[1]),
                "invoice_number": _s(row[2]), "supplier": _s(row[3]),
                "project": _s(row[4]), "customer": _s(row[5]),
                "invoice_total": _n(row[6]), "payment_method": _s(row[7]),
                "payment_status": _s(row[8]), "notes": _s(row[9]),
            })

    if "Price History" in wb.sheetnames:
        headers = [_s(cell.value) for cell in wb["Price History"][1]]
        has_product_identity = "اسم المنتج" in headers
        max_col = 20 if has_product_identity else 15
        for row in wb["Price History"].iter_rows(min_row=2, max_col=max_col, values_only=True):
            row = list(row) + [None] * (max_col - len(row))
            if _s(row[4]) == "" and _s(row[5]) == "":
                continue
            if has_product_identity:
                product = {
                    "item_name": _s(row[5]), "product_name": _s(row[6]) or _s(row[5]),
                    "brand": _s(row[7]), "main_category": _s(row[8]),
                    "subcategory": _s(row[9]), "specifications": _s(row[10]),
                }
                supplier_idx, project_idx, quantity_idx, unit_idx = 11, 12, 13, 14
                price_idx, discount_idx, vat_idx, final_idx, employee_idx = 15, 16, 17, 18, 19
            else:
                product = {
                    "item_name": _s(row[5]), "product_name": _s(row[5]),
                    "brand": "", "main_category": "", "subcategory": "",
                    "specifications": "",
                }
                supplier_idx, project_idx, quantity_idx, unit_idx = 6, 7, 8, 9
                price_idx, discount_idx, vat_idx, final_idx, employee_idx = 10, 11, 12, 13, 14
            out["price_history"].append({
                "date": _s(row[1]), "po_number": _s(row[2]), "invoice_number": _s(row[3]),
                "item_code": _s(row[4]), **product, "supplier": _s(row[supplier_idx]),
                "project": _s(row[project_idx]), "quantity": _n(row[quantity_idx]),
                "unit": _s(row[unit_idx]), "unit_price": _n(row[price_idx]),
                "discount_pct": _pct(row[discount_idx]), "vat_pct": _pct(row[vat_idx]),
                "final_price": _n(row[final_idx]), "employee": _s(row[employee_idx]),
            })

    if "Payments" in wb.sheetnames:
        for row in wb["Payments"].iter_rows(min_row=2, max_col=11, values_only=True):
            row = list(row) + [None] * (11 - len(row))
            if not _s(row[0]).startswith("PAY-") or _s(row[2]) == "" or row[6] is None:
                continue
            out["payments"].append({
                "payment_id": _s(row[0]), "payment_date": _s(row[1]),
                "purchase_id": _s(row[2]), "amount_paid": _n(row[6]),
                "payment_method": _s(row[8]), "notes": _s(row[10]),
            })

    if "Settings" in wb.sheetnames:
        rows = list(wb["Settings"].iter_rows(min_row=2, values_only=True))
        for key, label, col in SETTINGS_LISTS:
            values = []
            for r in rows:
                if col < len(r) and r[col] is not None and _s(r[col]) != "":
                    v = r[col]
                    if key == "vat_rates":
                        values.append(_pct(v))
                    else:
                        values.append(_s(v))
            if key == "vat_rates" and 0 not in values and _n(rows[0][col] if rows else None) == 0:
                pass
            out["settings"][key] = {"label": label, "values": values}
        if "vat_rates" in out["settings"]:
            vr = out["settings"]["vat_rates"]["values"]
            if 0 not in vr:
                vr.insert(0, 0)
    return out


async def next_seq_id(coll, field: str, prefix: str, width: int):
    docs = await coll.find({}, {field: 1}).to_list(100000)
    mx = 0
    for d in docs:
        c = str(d.get(field, ""))
        if c.startswith(prefix):
            try:
                mx = max(mx, int(c[len(prefix):]))
            except ValueError:
                pass
    return f"{prefix}{mx + 1:0{width}d}"


async def next_record_no(coll):
    docs = await coll.find({}, {"record_no": 1}).to_list(100000)
    return max((int(d.get("record_no", 0) or 0) for d in docs), default=0) + 1


async def recompute_payment_status(db, purchase_id: str):
    pur = await db.purchases.find_one({"purchase_id": purchase_id})
    if not pur:
        return
    pays = await db.payments.find({"purchase_id": purchase_id}).to_list(10000)
    paid = sum(p.get("amount_paid", 0) for p in pays)
    total = pur.get("invoice_total", 0)
    if paid <= 0:
        status = "غير مدفوع"
    elif paid >= total - 0.001:
        status = "مدفوع"
    else:
        status = "مدفوع جزئي"
    await db.purchases.update_one({"purchase_id": purchase_id}, {"$set": {"payment_status": status}})


def _norm_status(s):
    if s in ("تم الدفع", "مدفوع"):
        return "مدفوع"
    if s in ("مدفوع جزئي", "جزئي"):
        return "مدفوع جزئي"
    return "غير مدفوع"


async def import_data(db, parsed: dict) -> dict:
    counts = {}

    added = 0
    for s in parsed["suppliers"]:
        if await db.suppliers.find_one({"name": s["name"]}):
            continue
        s = dict(s)
        s["id"] = str(uuid.uuid4())
        s["code"] = next_business_code("suppliers")
        await db.suppliers.insert_one(s)
        added += 1
    counts["suppliers"] = added

    added = 0
    for c in parsed["customers"]:
        if await db.customers.find_one({"$or": [{"code": c["code"]}, {"name": c["name"]}]}):
            continue
        c = dict(c)
        c["id"] = str(uuid.uuid4())
        if not c.get("code"):
            c["code"] = next_business_code("customers")
        await db.customers.insert_one(c)
        added += 1
    counts["customers"] = added

    added = 0
    for p in parsed["projects"]:
        if await db.projects.find_one({"$or": [{"code": p["code"]}, {"name": p["name"]}]}):
            continue
        p = dict(p)
        p["id"] = str(uuid.uuid4())
        if not p.get("code"):
            p["code"] = next_business_code("projects")
        await db.projects.insert_one(p)
        added += 1
    counts["projects"] = added

    added = 0
    for it in parsed["items"]:
        if await db.items.find_one({"$or": [{"code": it["code"]}, {"name": it["name"]}]}):
            continue
        it = dict(it)
        it["id"] = str(uuid.uuid4())
        if not it.get("code"):
            it["code"] = next_business_code("items")
        await db.items.insert_one(it)
        added += 1
    counts["items"] = added

    added = 0
    for p in parsed["purchases"]:
        if await db.purchases.find_one({"purchase_id": p["purchase_id"]}):
            continue
        if await db.purchases.find_one({"invoice_number": p["invoice_number"], "supplier_name": p["supplier"]}):
            continue
        sup = await db.suppliers.find_one({"name": p["supplier"]})
        prj = await db.projects.find_one({"name": p["project"]})
        cus = await db.customers.find_one({"name": p["customer"]})
        related = [h for h in parsed["price_history"]
                   if h["invoice_number"] == p["invoice_number"] and h["supplier"] == p["supplier"]]
        subtotal = sum(h["quantity"] * h["unit_price"] for h in related)
        discount_total = sum(h["quantity"] * h["unit_price"] * h["discount_pct"] / 100 for h in related)
        vat_total = sum(h["quantity"] * h["unit_price"] * (1 - h["discount_pct"] / 100) * h["vat_pct"] / 100 for h in related)
        doc = {
            "id": str(uuid.uuid4()), "purchase_id": p["purchase_id"],
            "po_number": related[0]["po_number"] if related else "",
            "purchase_date": p["purchase_date"], "invoice_number": p["invoice_number"],
            "invoice_date": p["purchase_date"],
            "supplier_id": sup["id"] if sup else "", "supplier_name": p["supplier"],
            "project_id": prj["id"] if prj else "", "project_name": p["project"],
            "customer_id": cus["id"] if cus else "", "customer_name": p["customer"],
            "purchase_type": "محلي", "currency": "EGP",
            "created_by": related[0]["employee"] if related else "",
            "payment_method": p["payment_method"], "notes": p["notes"],
            "shipping_cost": 0, "other_costs": 0,
            "subtotal": round(subtotal, 2), "discount_total": round(discount_total, 2),
            "after_discount": round(subtotal - discount_total, 2), "vat_total": round(vat_total, 2),
            "invoice_total": p["invoice_total"],
            "payment_status": _norm_status(p["payment_status"]),
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.purchases.insert_one(doc)
        for h in related:
            item = await db.items.find_one({"code": h["item_code"]})
            await db.purchase_items.insert_one({
                "id": str(uuid.uuid4()), "purchase_id": p["purchase_id"],
                "item_id": item["id"] if item else "", "item_code": h["item_code"],
                "item_name": h["item_name"],
                "product_name": h.get("product_name") or (item or {}).get("product_name") or h["item_name"],
                "brand": h.get("brand") or (item or {}).get("brand", ""),
                "main_category": h.get("main_category") or (item or {}).get("main_category", ""),
                "subcategory": h.get("subcategory") or (item or {}).get("subcategory", ""),
                "specifications": h.get("specifications") or (item or {}).get("specifications", ""),
                "unit": h["unit"], "quantity": h["quantity"],
                "unit_price": h["unit_price"], "discount_pct": h["discount_pct"],
                "vat_pct": h["vat_pct"], "line_total": h["final_price"],
            })
        added += 1
    counts["purchases"] = added

    added = 0
    for h in parsed["price_history"]:
        dup = await db.price_history.find_one({
            "invoice_number": h["invoice_number"], "item_code": h["item_code"],
            "quantity": h["quantity"], "unit_price": h["unit_price"], "date": h["date"],
        })
        if dup:
            continue
        rec = dict(h)
        item = await db.items.find_one({"code": h["item_code"]})
        if item:
            rec["product_name"] = rec.get("product_name") or item.get("product_name") or item.get("name", "")
            rec["brand"] = rec.get("brand") or item.get("brand", "")
            rec["main_category"] = rec.get("main_category") or item.get("main_category", "")
            rec["subcategory"] = rec.get("subcategory") or item.get("subcategory", "")
            rec["specifications"] = rec.get("specifications") or item.get("specifications") or item.get("specs", "")
        rec["id"] = str(uuid.uuid4())
        rec["record_no"] = await next_record_no(db.price_history)
        await db.price_history.insert_one(rec)
        added += 1
    counts["price_history"] = added

    added = 0
    touched = set()
    for pay in parsed["payments"]:
        if await db.payments.find_one({"payment_id": pay["payment_id"]}):
            continue
        pay = dict(pay)
        pay["id"] = str(uuid.uuid4())
        pay["created_at"] = datetime.now(timezone.utc).isoformat()
        await db.payments.insert_one(pay)
        touched.add(pay["purchase_id"])
        added += 1
    counts["payments"] = added
    for pid in touched:
        await recompute_payment_status(db, pid)

    added = 0
    for key, label, _ in SETTINGS_LISTS:
        if await db.settings.find_one({"key": key}):
            continue
        data = parsed["settings"].get(key, {"label": label, "values": []})
        await db.settings.insert_one({"id": str(uuid.uuid4()), "key": key,
                                      "label": data["label"], "values": data["values"]})
        added += 1
    counts["settings"] = added
    return counts


HEADER_FILL = PatternFill(start_color="1E4E9C", end_color="1E4E9C", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}1"
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = openpyxl.styles.Alignment(horizontal="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for index, header in enumerate(headers, 1):
        values = [header, *(ws.cell(row=row, column=index).value for row in range(2, ws.max_row + 1))]
        content_width = max((len(str(value)) for value in values if value is not None), default=0)
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = min(45, max(12, content_width + 3))
    return ws


def _format_export_columns(ws, *, dates=(), money=(), numbers=(), percentages=()):
    for column in dates:
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2):
            for value_cell in cell:
                if isinstance(value_cell.value, str):
                    try:
                        value_cell.value = datetime.strptime(value_cell.value[:10], "%Y-%m-%d").date()
                    except ValueError:
                        pass
                value_cell.number_format = "yyyy-mm-dd"
    for column in money:
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2):
            for value_cell in cell:
                value_cell.number_format = '#,##0.00 "EGP"'
    for column in numbers:
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2):
            for value_cell in cell:
                value_cell.number_format = "#,##0.00"
    for column in percentages:
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2):
            for value_cell in cell:
                value_cell.number_format = '0.00"%"'


async def build_export_workbook(db) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    purchases = await db.purchases.find({}, {"_id": 0}).sort("purchase_id", 1).to_list(100000)
    payments = await db.payments.find({}, {"_id": 0}).sort("payment_id", 1).to_list(100000)
    pur_map = {p["purchase_id"]: p for p in purchases}
    paid_map = {}
    for p in payments:
        paid_map[p["purchase_id"]] = paid_map.get(p["purchase_id"], 0) + p.get("amount_paid", 0)

    purchase_register = _sheet(wb, "Purchase Register",
           ["رقم العملية", "تاريخ الشراء", "رقم الفاتورة", "المورد", "المشروع", "العميل",
            "إجمالي الفاتورة", "طريقة الدفع", "حالة الدفع", "ملاحظات"],
           [[p["purchase_id"], p["purchase_date"], p["invoice_number"], p["supplier_name"],
             p["project_name"], p["customer_name"], p["invoice_total"], p.get("payment_method", ""),
             p.get("payment_status", ""), p.get("notes", "")] for p in purchases])
    _format_export_columns(purchase_register, dates=(2,), money=(7,))

    items_rows = await db.purchase_items.find({}, {"_id": 0}).sort("purchase_id", 1).to_list(100000)
    purchase_items = _sheet(wb, "Purchase Items",
           ["رقم العملية", "كود الصنف", "اسم الصنف السابق", "اسم المنتج", "العلامة التجارية",
            "التصنيف الرئيسي", "التصنيف الفرعي", "المواصفات", "الوحدة", "الكمية",
            "سعر الوحدة", "الخصم %", "الضريبة %", "الإجمالي"],
           [[r["purchase_id"], r["item_code"], r["item_name"],
             r.get("product_name") or r["item_name"], r.get("brand", ""),
             r.get("main_category", ""), r.get("subcategory", ""),
             r.get("specifications", ""), r["unit"], r["quantity"], r["unit_price"],
             r["discount_pct"], r["vat_pct"], r["line_total"]] for r in items_rows])
    _format_export_columns(
        purchase_items, money=(11, 14), numbers=(10,), percentages=(12, 13)
    )

    hist = await db.price_history.find({}, {"_id": 0}).sort("record_no", 1).to_list(100000)
    price_history = _sheet(wb, "Price History",
           ["رقم السجل", "التاريخ", "رقم أمر الشراء", "رقم الفاتورة", "كود الصنف",
            "اسم الصنف السابق", "اسم المنتج", "العلامة التجارية", "التصنيف الرئيسي",
            "التصنيف الفرعي", "المواصفات", "المورد", "المشروع", "الكمية", "الوحدة",
            "سعر الوحدة", "الخصم %", "الضريبة %", "السعر النهائي", "الموظف"],
           [[h["record_no"], h["date"], h["po_number"], h["invoice_number"], h["item_code"],
             h["item_name"], h.get("product_name") or h["item_name"], h.get("brand", ""),
             h.get("main_category", ""), h.get("subcategory", ""), h.get("specifications", ""),
             h["supplier"], h["project"], h["quantity"], h["unit"], h["unit_price"],
             h["discount_pct"], h["vat_pct"], h["final_price"], h["employee"]]
            for h in hist])
    _format_export_columns(
        price_history, dates=(2,), money=(16, 19), numbers=(14,), percentages=(17, 18)
    )

    pay_rows = []
    for p in payments:
        pur = pur_map.get(p["purchase_id"], {})
        total = pur.get("invoice_total", 0)
        pay_rows.append([p["payment_id"], p["payment_date"], p["purchase_id"],
                         pur.get("invoice_number", ""), pur.get("supplier_name", ""), total,
                         p["amount_paid"], max(0, total - paid_map.get(p["purchase_id"], 0)),
                         p.get("payment_method", ""), pur.get("payment_status", ""), p.get("notes", "")])
    payments_sheet = _sheet(wb, "Payments",
           ["رقم الدفعة", "تاريخ الدفع", "رقم العملية", "رقم الفاتورة", "المورد", "إجمالي الفاتورة",
            "المبلغ المدفوع", "المتبقي", "طريقة الدفع", "حالة الدفع", "ملاحظات"], pay_rows)
    _format_export_columns(payments_sheet, dates=(2,), money=(6, 7, 8))

    sups = await db.suppliers.find({}, {"_id": 0}).sort("code", 1).to_list(100000)
    _sheet(wb, "Suppliers",
           ["كود المورد", "اسم المورد", "التخصص", "المجموعة", "المحافظة", "المدينة", "العنوان",
            "مسؤول التواصل", "رقم الهاتف", "واتساب", "البريد الإلكتروني", "شروط الدفع",
            "مدة التوريد بالأيام", "تقييم المورد", "الحالة", "ملاحظات"],
           [[s.get("code", ""), s["name"], s.get("specialty", ""), s.get("group_name", ""),
             s.get("governorate", ""), s.get("city", ""), s.get("address", ""),
             s.get("contact_person", ""), s.get("phone", ""), s.get("whatsapp", ""),
             s.get("email", ""), s.get("payment_terms", ""), s.get("lead_time_days", ""),
             s.get("rating", ""), s.get("status", ""), s.get("notes", "")] for s in sups])

    prjs = await db.projects.find({}, {"_id": 0}).sort("code", 1).to_list(100000)
    projects_sheet = _sheet(wb, "Projects",
           ["كود المشروع", "اسم المشروع", "العميل", "المحافظة", "المدينة", "العنوان",
            "المهندس المسؤول", "تاريخ البدء", "تاريخ الانتهاء المتوقع", "الميزانية",
            "حالة المشروع", "ملاحظات"],
           [[p.get("code", ""), p["name"], p.get("customer_name", ""), p.get("governorate", ""),
             p.get("city", ""), p.get("address", ""), p.get("engineer", ""), p.get("start_date", ""),
             p.get("end_date", ""), p.get("budget", ""), p.get("status", ""), p.get("notes", "")]
            for p in prjs])
    _format_export_columns(projects_sheet, dates=(8, 9), money=(10,))

    items = await db.items.find({}, {"_id": 0}).sort("code", 1).to_list(100000)
    _sheet(wb, "Items",
           ["كود الصنف", "اسم الصنف السابق", "اسم المنتج", "العلامة التجارية",
            "التصنيف الرئيسي", "التصنيف الفرعي", "الوحدة", "المواصفات",
            "المورد المفضل", "ملاحظات"],
           [[i.get("code", ""), i["name"], i.get("product_name") or i["name"],
             i.get("brand", ""), i.get("main_category") or i.get("category", ""),
             i.get("subcategory", ""), i.get("unit", ""),
             i.get("specifications") or i.get("specs", ""),
             i.get("preferred_supplier", ""), i.get("notes", "")] for i in items])

    cus = await db.customers.find({}, {"_id": 0}).sort("code", 1).to_list(100000)
    _sheet(wb, "Customers",
           ["كود العميل", "اسم العميل", "مسؤول التواصل", "رقم الهاتف", "واتساب",
            "البريد الإلكتروني", "المحافظة", "المدينة", "العنوان", "الحالة", "ملاحظات"],
           [[c.get("code", ""), c["name"], c.get("contact_person", ""), c.get("phone", ""),
             c.get("whatsapp", ""), c.get("email", ""), c.get("governorate", ""),
             c.get("city", ""), c.get("address", ""), c.get("status", ""), c.get("notes", "")]
            for c in cus])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
